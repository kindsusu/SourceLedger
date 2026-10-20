"""Publish one evidence-preserving workbook and a standalone report dashboard."""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from html import escape
import os
from pathlib import Path
import shutil
import tempfile
from urllib.parse import quote, urlparse

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from .report import _validate_xlsx


def _header_row(ws):
    if ws.title == "Collection Changes":
        return 5
    return 2 if str(ws.cell(1, 1).value or "").startswith("Demo data") else 1


def _headers(ws):
    row = _header_row(ws)
    return tuple(ws.cell(row, col).value for col in range(1, ws.max_column + 1))


def _safe_link(link):
    if link is None:
        return None
    if link.target:
        parsed = urlparse(link.target)
        if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
            return None
    elif link.location and not link.location.startswith("#"):
        return None
    return copy(link)


def _copy_cell(source: Cell, target: Cell):
    # A supplied workbook is data. Formula cells must remain literal after merging.
    target.value = str(source.value) if source.data_type == "f" else source.value
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)
    if source.comment:
        target.comment = copy(source.comment)
    target.hyperlink = _safe_link(source.hyperlink)
    if isinstance(target.value, str):
        target.data_type = "s"


def _copy_sheet(source, destination, *, source_run=None):
    for row in source:
        for cell in row:
            _copy_cell(cell, destination.cell(cell.row, cell.column))
        if source.row_dimensions[row[0].row].height is not None:
            destination.row_dimensions[row[0].row].height = source.row_dimensions[row[0].row].height
    for key, dim in source.column_dimensions.items():
        destination.column_dimensions[key].width = dim.width
        destination.column_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        destination.merge_cells(str(merged))
    destination.freeze_panes = source.freeze_panes
    destination.auto_filter.ref = source.auto_filter.ref
    if source_run is not None:
        col = len(_headers(source)) + 1
        destination.cell(_header_row(source), col, "Source Run")
        destination.column_dimensions[destination.cell(1, col).column_letter].width = 24
        for row in range(_header_row(source) + 1, destination.max_row + 1):
            destination.cell(row, col, source_run).data_type = "s"
        if destination.auto_filter.ref:
            from openpyxl.utils.cell import get_column_letter
            destination.auto_filter.ref = f"A{_header_row(source)}:{get_column_letter(col)}{destination.max_row}"


def _run_info(wb):
    if "Run Summary" not in wb:
        return {}
    ws = wb["Run Summary"]
    return {str(ws.cell(row, 1).value): ws.cell(row, 2).value
            for row in range(_header_row(ws) + 1, ws.max_row + 1)
            if ws.cell(row, 1).value is not None}


def _merge(paths, destination):
    from openpyxl import Workbook

    combined = Workbook()
    combined.remove(combined.active)
    summaries = []
    run_details = []
    change_details = []
    sheet_state = {}  # name -> (original headers, header row, provenance column, last data row)
    try:
        for index, path in enumerate(paths, 1):
            wb = load_workbook(path, data_only=False)
            try:
                info = _run_info(wb)
                run_id = str(info.get("Run ID") or f"Run {index}")
                summaries.append((run_id, path.name, info.get("Status"), info.get("Demo data")))
                if "Run Summary" in wb:
                    ws = wb["Run Summary"]
                    for row in range(_header_row(ws) + 1, ws.max_row + 1):
                        run_details.append((run_id, ws.cell(row, 1), ws.cell(row, 2)))
                if "Collection Changes" in wb:
                    ws = wb["Collection Changes"]
                    for row in range(1, 5):
                        change_details.append((run_id, row, tuple(ws[row])))
                for source in wb:
                    if source.title in {"Run Summary", "Source Runs"}:
                        continue
                    headers = _headers(source)
                    if source.title not in sheet_state:
                        target = combined.create_sheet(source.title)
                        _copy_sheet(source, target, source_run=run_id)
                        sheet_state[source.title] = (headers, _header_row(source), len(headers) + 1,
                                                     target.max_row)
                        continue
                    target = combined[source.title]
                    expected_headers, header_row, provenance_col, last_row = sheet_state[source.title]
                    if headers != expected_headers:
                        raise ValueError(f"incompatible sheet headers: {source.title}")
                    start = last_row + 1
                    for source_row in source.iter_rows(min_row=_header_row(source) + 1):
                        for cell in source_row:
                            _copy_cell(cell, target.cell(start, cell.column))
                        target.cell(start, provenance_col, run_id).data_type = "s"
                        start += 1
                    last_row = start - 1
                    sheet_state[source.title] = (expected_headers, header_row, provenance_col, last_row)
                    if target.auto_filter.ref:
                        from openpyxl.utils.cell import get_column_letter
                        target.auto_filter.ref = f"A{header_row}:{get_column_letter(provenance_col)}{last_row}"
            finally:
                wb.close()
        summary = combined.create_sheet("Run Summary", 0)
        summary.append(["Field", "Value"])
        summary.append(["Merged source runs", len(paths)])
        summary.append(["Demo data", "Yes" if any(str(item[3]).lower() == "yes" for item in summaries) else "No"])
        summary.append(["Note", "Rows retain their source run; source workbooks remain unchanged."])
        if "Collection Changes" in combined:
            changes = combined["Collection Changes"]
            for row in range(1, 5):
                for cell in changes[row]:
                    cell.value = None
            changes.cell(1, 1, "Per-run change metadata")
            changes.cell(1, 2, "See Change Summaries; rows below combine source runs.")
        provenance = combined.create_sheet("Source Runs", 1)
        provenance.append(["Run ID", "Source workbook", "Run status", "Demo data"])
        for item in summaries:
            provenance.append(item)
            for cell in provenance[provenance.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        details = combined.create_sheet("Run Details", 2)
        details.append(["Source Run", "Field", "Value"])
        for run_id, field, value in run_details:
            row = details.max_row + 1
            details.cell(row, 1, run_id).data_type = "s"
            _copy_cell(field, details.cell(row, 2))
            _copy_cell(value, details.cell(row, 3))
        if change_details:
            changes = combined.create_sheet("Change Summaries", 3)
            changes.append(["Source Run", "Original row", *[f"Column {n}" for n in range(1, 10)]])
            for run_id, original_row, cells in change_details:
                row = changes.max_row + 1
                changes.cell(row, 1, run_id).data_type = "s"
                changes.cell(row, 2, original_row)
                for col, cell in enumerate(cells, 3):
                    _copy_cell(cell, changes.cell(row, col))
            changes.freeze_panes = "C2"
        for ws in (summary, provenance, details):
            for cell in ws[1]:
                cell.font = copy(combined["Run Summary"]["A1"].font)
            ws.freeze_panes = "A2"
        combined.save(destination)
    finally:
        combined.close()


_PRICE = "Recorded Price (not a final quote)"


def _stream_records(wb, sheet, *, keep_rows=True):
    """Read one sheet once, in streaming order: (header, data rows or count)."""
    if sheet not in wb.sheetnames:
        return [], [] if keep_rows else 0
    rows = wb[sheet].iter_rows(values_only=True)
    first = next(rows, None)
    if first is None:
        return [], [] if keep_rows else 0
    header_row = 5 if sheet == "Collection Changes" else (
        2 if str((first[0] if first else None) or "").startswith("Demo data") else 1)
    header = first
    for _ in range(header_row - 1):
        header = next(rows, ())
    header = [str(value or "") for value in header]
    if not keep_rows:
        return header, sum(1 for _ in rows)
    return header, [list(row[:len(header)]) for row in rows]


def _column(header, row, name):
    return row[header.index(name)] if name in header and header.index(name) < len(row) else None


def workbook_summary(path) -> dict:
    """Dashboard inputs read from a finished workbook without loading every sheet.

    ``report._export`` returns the same structure for the workbook it wrote.
    """
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        result_h, results = _stream_records(wb, "Results")
        _, task_count = _stream_records(wb, "Collection Status", keep_rows=False)
        _, comparison_count = _stream_records(wb, "Price Comparison", keep_rows=False)
        evidence_urls = []
        if not results:
            evidence_h, evidence = _stream_records(wb, "Source Evidence")
            evidence_urls = [_column(evidence_h, row, "Source URL") for row in evidence]
        categories = None
        if "Category" in result_h:
            categories = Counter()
            for row in results:
                value = _column(result_h, row, "Category")
                categories[str(value) if value not in (None, "") else None] += 1
        demo = False
        if "Run Summary" in wb.sheetnames:
            _, run_rows = _stream_records(wb, "Run Summary")
            info = {str(row[0]): row[1] if len(row) > 1 else None for row in run_rows if row and row[0] is not None}
            demo = str(info.get("Demo data") or "").casefold() == "yes"
        return {"result_rows": [[_column(result_h, row, "Source URL"), _column(result_h, row, _PRICE) is not None]
                                for row in results],
                "price_column": _PRICE in result_h, "evidence_urls": evidence_urls,
                "categories": dict(categories) if categories is not None else None,
                "task_count": task_count, "comparison_count": comparison_count,
                "demo": demo, "sheet_names": list(wb.sheetnames)}
    finally:
        wb.close()


def _dashboard(summary, xlsx_name, metadata):
    results = summary["result_rows"]
    price_count = sum(1 for _, priced in results if priced) if summary["price_column"] else 0
    sites = defaultdict(lambda: [0, 0])
    for url, priced in results:
        site = urlparse(str(url or "")).hostname
        sites[site][0] += 1
        if priced:
            sites[site][1] += 1
    for url in ([] if results else summary["evidence_urls"]):
        sites[urlparse(str(url or "")).hostname][0] += 1
    category_available = summary["categories"] is not None
    categories = Counter(summary["categories"] or {})
    coverage = metadata.get("coverage") if isinstance(metadata, dict) else None
    coverage_text = "Unknown"
    if isinstance(coverage, list):
        if not coverage:
            coverage_text = "Unknown"
        else:
            from .coverage import summarize
            counts = summarize(coverage)
            complete = counts["complete"]
            total = len(coverage)
            coverage_text = (f"{complete} / {total} ({complete / total:.1%}) complete; "
                             f"{counts['partial']} partial, {counts['unknown']} unknown, "
                             f"{counts['excluded']} excluded")
    elif isinstance(coverage, dict):
        completed, total = coverage.get("completed"), coverage.get("total")
        if isinstance(completed, int) and isinstance(total, int) and 0 <= completed <= total and total > 0:
            coverage_text = f"{completed} / {total} ({completed / total:.1%})"
    elif isinstance(metadata, dict):
        supplied = metadata.get("coverage_summary")
        total = metadata.get("target_count")
        if isinstance(supplied, dict) and isinstance(total, int) and total > 0:
            complete = supplied.get("complete")
            if isinstance(complete, int) and 0 <= complete <= total:
                partial = supplied.get("partial", 0)
                unknown = supplied.get("unknown", 0)
                excluded = supplied.get("excluded", 0)
                if all(isinstance(v, int) and v >= 0 for v in (partial, unknown, excluded)):
                    coverage_text = (f"{complete} / {total} ({complete / total:.1%}) complete; "
                                     f"{partial} partial, {unknown} unknown, {excluded} excluded")
    demo = bool(summary["demo"])
    name = escape(xlsx_name, quote=True)
    href = quote(xlsx_name, safe="")
    labels = {
        "en": ("SourceLedger final report", "Open final workbook", "Observed product records", "Collection tasks", "Rows with recorded prices", "Comparable price rows", "Source sites", "Category", "Records", "Coverage", "Accuracy", "Unknown", "Not measured", "No source records", "Category unavailable: no recorded category field.", "Counts describe recorded rows, not market completeness or price accuracy.", "Included data", "Collection status", "Started", "Finished", "Source", "Workbook", "Review note", "Recorded prices and commercial conditions may still need human review.", "COLLECTION REPORT", "Demo data — not market prices"),
        "ko": ("SourceLedger 최종 보고서", "최종 워크북 열기", "관측 제품 기록", "수집 작업", "기록된 가격이 있는 행", "비교 가능 가격 행", "출처 사이트", "분류", "기록", "수집 범위", "정확도", "알 수 없음", "측정되지 않음", "출처 기록 없음", "분류 정보 없음: 분류 필드가 기록되지 않았습니다.", "집계는 기록된 행만 나타내며 시장 전체의 수집 완전성이나 가격 정확도를 뜻하지 않습니다.", "포함된 데이터", "수집 상태", "시작", "종료", "출처", "워크북", "검토 안내", "기록된 가격과 거래 조건은 사람의 검토가 더 필요할 수 있습니다.", "수집 보고서", "데모 데이터 — 시장 가격이 아닙니다"),
        "ja": ("SourceLedger 最終レポート", "最終ブックを開く", "観測製品の記録", "収集タスク", "記録価格のある行", "比較可能な価格行", "情報源サイト", "分類", "記録", "カバレッジ", "精度", "不明", "未測定", "情報源の記録なし", "分類情報なし: 分類項目は記録されていません。", "集計は記録された行のみを示し、市場全体の網羅性や価格の正確さを示すものではありません。", "含まれるデータ", "収集状態", "開始", "終了", "情報源", "ブック", "確認事項", "記録された価格と取引条件は人による確認が必要な場合があります。", "収集レポート", "デモデータ — 市場価格ではありません"),
    }
    topic = metadata.get("topic") if isinstance(metadata, dict) else None
    topic_fields = {key: str(topic[key]) for key in ("industry", "product", "market")
                    if isinstance(topic, dict) and topic.get(key) not in (None, "")}
    if isinstance(topic, str) and topic.strip():
        topic_fields = {"topic": topic}
    status = metadata.get("status") if isinstance(metadata, dict) else None
    started = metadata.get("started_at") if isinstance(metadata, dict) else None
    finished = metadata.get("finished_at") if isinstance(metadata, dict) else None
    included_names = [name for name in ("Results", "Collection Status", "Price Comparison", "Observation History",
                                         "Source Evidence", "Review Required", "Catalog coverage", "Collection Issues",
                                         "Collection Changes", "Item Changes", "Rental Quotes", "Supporting Evidence")
                      if name in summary["sheet_names"]]
    sections = []
    for lang, lab in labels.items():
        localized_site_rows = ("".join(f"<tr><th scope='row'>{escape(site or lab[11])}</th><td>{values[0]}</td><td>{values[1]}</td></tr>"
                                       for site, values in sorted(sites.items(), key=lambda item: item[0] or ""))
                               if sites else f"<tr><td colspan='3'>{lab[13]}</td></tr>")
        localized_category_rows = ("".join(f"<tr><th scope='row'>{escape(category or lab[11])}</th><td>{count}</td></tr>"
                                           for category, count in sorted(categories.items(), key=lambda item: item[0] or ""))
                                   if category_available else f"<tr><td colspan='2'>{lab[14]}</td></tr>")
        localized_coverage = coverage_text.replace("Unknown", lab[11]).replace(" complete; ", {"en": " complete; ", "ko": " 완료; ", "ja": " 完了; "}[lang]).replace(" partial, ", {"en": " partial, ", "ko": " 부분, ", "ja": " 部分, "}[lang]).replace(" unknown, ", {"en": " unknown, ", "ko": " 미확인, ", "ja": " 不明, "}[lang]).replace(" excluded", {"en": " excluded", "ko": " 제외", "ja": " 除外"}[lang])
        topic_labels = {"en": ("Industry", "Product", "Market", "Topic"),
                        "ko": ("산업", "제품", "시장", "주제"),
                        "ja": ("業界", "製品", "市場", "テーマ")}
        topic_keys = ("industry", "product", "market", "topic")
        topic_line = " · ".join(f"{topic_labels[lang][topic_keys.index(key)]}: {escape(value)}"
                                for key, value in topic_fields.items())
        timing = " · ".join(f"{label}: {escape(str(value))}" for label, value in ((lab[18], started), (lab[19], finished))
                            if value not in (None, ""))
        included = "".join(f"<li>{escape(sheet)}</li>" for sheet in included_names)
        category_panel = (f"<h2>{lab[7]}</h2><p class='subheading'>{lab[8]}</p>"
                          f"<table><thead><tr><th>{lab[7]}</th><th class='num'>{lab[8]}</th></tr></thead><tbody>{localized_category_rows}</tbody></table>"
                          if category_available else
                          f"<h2>{lab[16]}</h2><p class='subheading'>{lab[14]}</p><ul class='included'>{included}</ul>")
        sections.append(
            f"<section class='language {lang}' lang='{lang}'>"
            f"{f'<div class=demo-banner role=note>{lab[25]}</div>' if demo else ''}"
            f"<div class='heading'><div><h1>{lab[0]}</h1>{f'<p class=muted>{topic_line}</p>' if topic_line else ''}</div>"
            f"{f'<span class=badge>{lab[17]}: {escape(str(status))}</span>' if status not in (None, '') else ''}</div>"
            f"{f'<p class=meta>{timing}</p>' if timing else ''}"
            f"<div class='file'><div><strong>{lab[21]}</strong><p class='small muted'>{name}</p></div>"
            f"<a href='{href}'>{lab[1]} ↗</a></div>"
            f"<div class='stats'><article class='stat'><span>{lab[2]}</span><b>{len(results):,}</b></article>"
            f"<article class='stat'><span>{lab[3]}</span><b>{summary['task_count']:,}</b></article>"
            f"<article class='stat'><span>{lab[4]}</span><b>{price_count:,}</b></article>"
            f"<article class='stat'><span>{lab[5]}</span><b>{summary['comparison_count']:,}</b></article></div>"
            f"<div class='columns'><article class='panel'><h2>{lab[6]}</h2><p class='subheading'>{lab[15]}</p>"
            f"<table><thead><tr><th>{lab[20]}</th><th class='num'>{lab[8]}</th><th class='num'>{lab[4]}</th>"
            f"</tr></thead><tbody>{localized_site_rows}</tbody></table></article>"
            f"<article class='panel'>{category_panel}</article></div>"
            f"<div class='quality'><article class='panel'><h2>{lab[9]}</h2><div class='qvalue'>{escape(localized_coverage)}</div>"
            f"<p class='small muted'>{lab[15]}</p></article>"
            f"<article class='panel'><h2>{lab[10]}</h2><div class='qvalue'>{lab[12]}</div>"
            f"<p class='small muted'>{lab[15]}</p></article></div>"
            f"<aside class='notice'><strong>{lab[22]}</strong><p>{lab[23]}</p></aside>"
            f"<footer>SourceLedger <span>{lab[15]}</span></footer></section>")
    return ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width,initial-scale=1'>"
            "<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'\">"
            "<title>SourceLedger final report</title><style>"
            ":root{color-scheme:light;--bg:#f3f5f1;--paper:#fff;--ink:#193e38;--sub:#61756d;--line:#dbe5dd;--green:#1d745b;--soft:#eaf3ec;--amber:#865918;--amberbg:#fff4df}"
            "*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 'Segoe UI','Malgun Gothic',system-ui,sans-serif}"
            "main{max-width:1180px;margin:auto;padding:30px}.top,.heading,.file,footer{display:flex;justify-content:space-between;align-items:center;gap:16px;flex-wrap:wrap}"
            ".brand{font-size:19px;font-weight:650}.brand small{font-size:11px;letter-spacing:2px;margin-left:12px;color:var(--sub)}"
            "h1{font-size:30px;letter-spacing:-1px;margin:20px 0 2px}h2{font-size:16px;margin:0}p{margin:4px 0}.muted,.subheading{color:var(--sub)}"
            ".small,.subheading,.meta{font-size:12px}.meta{margin:13px 0 18px}.badge{padding:5px 11px;border-radius:30px;background:var(--amberbg);color:var(--amber);font-size:12px}"
            "a{color:var(--green);text-underline-offset:4px}a:focus-visible{outline:3px solid #2b8c70;outline-offset:3px}"
            ".file{background:var(--soft);border-radius:12px;padding:16px 20px;margin:18px 0}.file a{font-weight:650;font-size:15px}"
            ".stats{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}.stat,.panel{background:var(--paper);border:1px solid var(--line);border-radius:12px;padding:19px}"
            ".stat span{display:block;font-size:12px;color:var(--sub)}.stat b{display:block;font-size:34px;line-height:1.3;letter-spacing:-1px;font-variant-numeric:tabular-nums;margin:5px 0}"
            ".columns{display:grid;grid-template-columns:1.4fr 1fr;gap:16px;margin-top:16px}.subheading{margin:4px 0 14px}"
            "table{width:100%;border-collapse:collapse;font-size:13px}th{font-weight:400;font-size:12px;color:var(--sub);text-align:left;padding:0 8px 10px}"
            "td{padding:12px 8px;border-top:1px solid var(--line)}th:first-child,td:first-child{padding-left:0}.num{text-align:right;font-variant-numeric:tabular-nums}"
            ".included{columns:2;gap:22px;padding-left:18px;margin:10px 0}.included li{break-inside:avoid;margin:0 0 8px}"
            ".quality{display:grid;grid-template-columns:1fr 1fr;gap:16px;margin-top:16px}.qvalue{font-size:25px;font-weight:650;margin:6px 0}"
            ".notice{background:var(--amberbg);color:var(--amber);padding:14px 18px;border-radius:10px;margin-top:16px}"
            ".demo-banner{background:#8e2525;color:#fff;border-radius:8px;padding:10px 14px;font-weight:700;margin-top:16px}"
            "footer{font-size:11px;color:var(--sub);margin-top:17px}.language{display:none}"
            "input[name=lang]{position:absolute;width:1px;height:1px;opacity:0;clip-path:inset(50%)}"
            ".languages{display:flex;gap:6px}.languages label{cursor:pointer;border:1px solid var(--line);border-radius:6px;background:var(--paper);padding:8px 12px;font-size:12px}"
            "#en:checked~main .en,#ko:checked~main .ko,#ja:checked~main .ja{display:block}"
            "#en:checked~main label[for=en],#ko:checked~main label[for=ko],#ja:checked~main label[for=ja]{background:var(--ink);color:#fff}"
            "#en:focus-visible~main label[for=en],#ko:focus-visible~main label[for=ko],#ja:focus-visible~main label[for=ja]{outline:3px solid #2b8c70;outline-offset:3px}"
            "@media(max-width:760px){main{padding:18px}.stats{grid-template-columns:1fr 1fr}.columns,.quality{grid-template-columns:1fr}h1{font-size:25px}.panel{padding:16px}}"
            "@media(max-width:380px){main{padding:12px}.stat{padding:12px}.stat b{font-size:28px}.brand small{display:none}.included{columns:1}}"
            "@media print{@page{size:A4 landscape;margin:8mm}main{padding:0;max-width:none}body{font-size:11px;background:#fff}.languages{display:none}"
            ".stat,.panel{padding:10px}.columns,.quality{margin-top:10px}.stat b{font-size:25px}h1{font-size:23px;margin-top:8px}.notice{padding:8px}footer{margin-top:8px}"
            "article,.file,.notice{break-inside:avoid}}"
            "</style></head><body>"
            "<input type='radio' name='lang' id='en' checked aria-label='English'>"
            "<input type='radio' name='lang' id='ko' aria-label='한국어'>"
            "<input type='radio' name='lang' id='ja' aria-label='日本語'>"
            "<main><div class='top'><div class='brand'>SourceLedger <small>COLLECTION REPORT</small></div>"
            "<div class='languages' role='group' aria-label='Language'><label for='en'>English</label>"
            "<label for='ko'>한국어</label><label for='ja'>日本語</label></div></div>"
            + "".join(sections) + "</main></body></html>"), {
                "observed_product_records": len(results), "collection_tasks": summary["task_count"],
                "recorded_price_rows": price_count,
                "comparable_price_rows": summary["comparison_count"],
                "source_sites": sum(site is not None for site in sites),
            }


def summary_matches(summary, path: Path) -> bool:
    """A summary is reused only for the exact file state it describes."""
    if not isinstance(summary, dict) or not isinstance(summary.get("file"), (list, tuple)):
        return False
    try:
        stat = Path(path).stat()
    except OSError:
        return False
    return list(summary["file"]) == [stat.st_size, stat.st_mtime_ns] and {
        "result_rows", "price_column", "evidence_urls", "categories", "task_count",
        "comparison_count", "demo", "sheet_names"}.issubset(summary)


def _temp_near(path: Path, suffix: str) -> Path:
    fd, name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=suffix, dir=path.parent)
    os.close(fd)
    return Path(name)


def finalize_reports(report_paths: list[Path], output_path: Path, *, metadata: dict | None = None) -> dict:
    """Finalize exactly the supplied reports into a workbook and offline HTML dashboard.

    A single report at ``output_path`` is reused without rewriting its workbook.
    Batch inputs must be distinct from the output so source workbooks stay intact.
    ``metadata["workbook_summary"]`` (from the exporter) replaces re-reading the
    workbook when it still describes the file byte-for-byte (size and mtime).
    """
    paths = [Path(os.path.abspath(p)) for p in report_paths]
    output = Path(os.path.abspath(output_path))
    dashboard = output.with_suffix(".html")
    if not paths or len(set(paths)) != len(paths):
        raise ValueError("provide one or more distinct report paths")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("final report path must end in .xlsx")
    if len(paths) > 1 and output in paths:
        raise ValueError("batch output cannot overwrite a source report")
    for path in [*paths, output, dashboard]:
        if any(part.is_symlink() for part in (path, *path.parents)):
            raise ValueError(f"symbolic links are not supported: {path}")
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
        _validate_xlsx(path)  # The one full CRC check of each input.
    output.parent.mkdir(parents=True, exist_ok=True)
    same = len(paths) == 1 and paths[0] == output
    supplied = (metadata or {}).get("workbook_summary")
    summary = supplied if same and summary_matches(supplied, output) else None
    staged_book = None if same else _temp_near(output, ".xlsx")
    staged_html = _temp_near(dashboard, ".html")
    old_book = old_html = None
    try:
        if same:
            workbook_path = output
        elif len(paths) == 1:
            shutil.copyfile(paths[0], staged_book)
            workbook_path = staged_book
            _validate_xlsx(workbook_path, deep=False)
        else:
            _merge(paths, staged_book)
            workbook_path = staged_book
            _validate_xlsx(workbook_path)
        if summary is None:
            summary = workbook_summary(workbook_path)
        html, counts = _dashboard(summary, output.name, metadata or {})
        staged_html.write_text(html, encoding="utf-8")
        if not staged_html.read_text(encoding="utf-8").startswith("<!doctype html>"):
            raise ValueError("invalid HTML dashboard")
        if not same and output.exists():
            old_book = _temp_near(output, ".bak")
            shutil.copyfile(output, old_book)
        if dashboard.exists():
            old_html = _temp_near(dashboard, ".bak")
            shutil.copyfile(dashboard, old_html)
        try:
            if not same:
                os.replace(staged_book, output)
            os.replace(staged_html, dashboard)
        except Exception:
            if not same:
                if old_book is not None:
                    os.replace(old_book, output)
                else:
                    output.unlink(missing_ok=True)
            if old_html is not None:
                os.replace(old_html, dashboard)
            else:
                dashboard.unlink(missing_ok=True)
            raise
        return {"report_path": str(output), "dashboard_path": str(dashboard), **counts}
    finally:
        for path in (staged_book, staged_html, old_book, old_html):
            if path is not None:
                path.unlink(missing_ok=True)
