"""Small audit worksheets that can be written without reloading a whole report.

A ``SheetSpec`` describes one literal table. The same spec is written by
xlsxwriter during a report export, or later replaced inside an existing XLSX by
editing only the package parts that name worksheets. Replacing a sheet never
re-reads or re-serializes the other worksheets, so a large evidence report is
not round-tripped through openpyxl just to refresh a coverage or change table.
Unsupported package shapes fall back to a single openpyxl load/save.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from html import unescape
from io import BytesIO
import json
import math
import os
from pathlib import Path
from posixpath import basename as _pbasename, dirname as _pdirname, join as _pjoin, normpath as _pnormpath
import re
import tempfile
from typing import Any
from xml.etree import ElementTree
import zipfile

# Excel stores at most 32,767 UTF-16 code units in a cell. Longer text would
# be truncated silently by xlsxwriter or rejected by Excel; the report says so.
EXCEL_CELL_LIMIT = 32767
TRUNCATION_MARKER = "\n[Truncated for Excel: {shown} of {total} characters shown; the full value is kept in the SQLite ledger.]"

_ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WORKSHEET_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
_WORKSHEET_CONTENT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"

# Colors keep openpyxl's 00 alpha so reports refreshed in place read the same
# as before; Excel ignores the alpha byte for solid fills.
STYLES: dict[str, dict[str, Any]] = {
    "header_blue": {"bold": True, "color": "FFFFFF", "fill": "1F4E78"},
    "header_navy_wrap": {"bold": True, "color": "FFFFFF", "fill": "17365D", "wrap": True},
    "bold": {"bold": True},
    "wrap_top": {"wrap": True},
}


@dataclass
class SheetSpec:
    name: str
    rows: list[list[Any]]
    widths: dict[int, float] = field(default_factory=dict)  # 1-based column -> width
    freeze: tuple[int, int] | None = None  # 1-based (row, column) of the first unfrozen cell
    auto_filter: str | None = None
    row_styles: dict[int, str] = field(default_factory=dict)  # 1-based row -> style across all columns
    cell_style: str | None = None  # style for every other cell in the table bounds
    row_heights: dict[int, float] = field(default_factory=dict)

    @property
    def max_col(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def cells(self):
        """Yield (row, column, value, style) for every materialized cell."""
        width = self.max_col
        for r, values in enumerate(self.rows, 1):
            style = self.row_styles.get(r, self.cell_style)
            for c in range(1, width + 1):
                value = cell_value(values[c - 1] if c <= len(values) else None)
                if value is None and style is None:
                    continue
                yield r, c, value, style


def _utf16_len(text: str) -> int:
    return len(text.encode("utf-16-le", "surrogatepass")) // 2


def cell_text(text: str) -> str:
    """Return text Excel can store, with an explicit marker when it was cut."""
    if _ILLEGAL_XML.search(text):
        text = _ILLEGAL_XML.sub(lambda m: "_x%04X_" % ord(m.group()), text)
    if len(text) <= EXCEL_CELL_LIMIT // 2:
        return text  # Even all-astral text fits; skip counting UTF-16 units.
    total = _utf16_len(text)
    if total <= EXCEL_CELL_LIMIT:
        return text
    marker = TRUNCATION_MARKER.format(shown="{shown}", total=f"{total:,}")
    budget = EXCEL_CELL_LIMIT - _utf16_len(marker.format(shown=f"{EXCEL_CELL_LIMIT:,}"))
    cut = text[:budget]
    while (excess := _utf16_len(cut) - budget) > 0:  # astral characters use two units
        cut = cut[:len(cut) - excess]
    return cut + marker.format(shown=f"{len(cut):,}")


def cell_value(value: Any) -> Any:
    """Map a Python value to the literal scalar an audit cell stores."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, Decimal)) or (isinstance(value, float) and math.isfinite(value)):
        if isinstance(value, Decimal) and not value.is_finite():
            return cell_text(str(value))
        return value
    if isinstance(value, str):
        # openpyxl never materialized empty strings; keep them blank.
        return cell_text(value) if value else None
    if isinstance(value, (dict, list, tuple)):
        return cell_text(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str))
    return cell_text(str(value))


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ---------------------------------------------------------------- xlsxwriter

def write_xlsxwriter(workbook: Any, spec: SheetSpec, cache: dict[str, Any]) -> None:
    """Write one spec into an open xlsxwriter workbook (single export pass)."""
    ws = workbook.add_worksheet(spec.name)
    for key in {spec.cell_style, *spec.row_styles.values()} - {None} - set(cache):
        style = STYLES[key]
        options: dict[str, Any] = {}
        if style.get("bold"):
            options["bold"] = True
        if style.get("color"):
            options["font_color"] = "#" + style["color"]
        if style.get("fill"):
            options.update(pattern=1, bg_color="#" + style["fill"])
        if style.get("wrap"):
            options.update(text_wrap=True, valign="top")
        cache[key] = workbook.add_format(options)
    for r, c, value, style in spec.cells():
        fmt = cache.get(style) if style else None
        if value is None:
            ws.write_blank(r - 1, c - 1, None, fmt)
        elif isinstance(value, bool):
            ws.write_boolean(r - 1, c - 1, value, fmt)
        elif isinstance(value, (int, float, Decimal)):
            ws.write_number(r - 1, c - 1, value, fmt)
        else:
            ws.write_string(r - 1, c - 1, value, fmt)
    for col, width in spec.widths.items():
        ws.set_column(col - 1, col - 1, width)
    for row, height in spec.row_heights.items():
        ws.set_row(row - 1, height)
    if spec.freeze:
        ws.freeze_panes(spec.freeze[0] - 1, spec.freeze[1] - 1)
    if spec.auto_filter:
        ws.autofilter(spec.auto_filter)


# ------------------------------------------------------------------ openpyxl

def _write_openpyxl(wb: Any, spec: SheetSpec) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    if spec.name in wb:
        del wb[spec.name]
    ws = wb.create_sheet(spec.name)
    for r, c, value, style in spec.cells():
        cell = ws.cell(r, c)
        if value is not None:
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # Source text must never become an Excel formula.
        if style:
            options = STYLES[style]
            if options.get("bold") or options.get("color"):
                cell.font = Font(bold=bool(options.get("bold")), color=options.get("color"))
            if options.get("fill"):
                cell.fill = PatternFill("solid", fgColor=options["fill"])
            if options.get("wrap"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in spec.widths.items():
        ws.column_dimensions[_column_letter(col)].width = width
    for row, height in spec.row_heights.items():
        ws.row_dimensions[row].height = height
    if spec.freeze:
        ws.freeze_panes = f"{_column_letter(spec.freeze[1])}{spec.freeze[0]}"
    if spec.auto_filter:
        ws.auto_filter.ref = spec.auto_filter


def _replace_openpyxl(path: Path, specs: list[SheetSpec]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(path.read_bytes()))
    temp = None
    try:
        for spec in specs:
            _write_openpyxl(wb, spec)
        temp = _temp_near(path)
        wb.save(temp)
        validate_package(temp)
        os.replace(temp, path)
        temp = None
    finally:
        wb.close()
        if temp is not None:
            temp.unlink(missing_ok=True)


# ------------------------------------------------------------- package splice

class SpliceUnsupported(Exception):
    """The workbook package has a shape this editor does not rewrite."""


def _temp_near(path: Path) -> Path:
    fd, name = tempfile.mkstemp(prefix=path.stem + ".", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    return Path(name)


def validate_package(path: Path, *, deep: bool = False) -> None:
    """Check an XLSX zip. ``deep`` also verifies every member CRC."""
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names:
            raise ValueError("invalid XLSX archive")
        if deep and archive.testzip() is not None:
            raise ValueError("invalid XLSX archive")


def _attrs(element: str) -> dict[str, str]:
    return dict(re.findall(r'([\w:.-]+)\s*=\s*"([^"]*)"', element))


def _xml_attr(value: str) -> str:
    return (value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


def _xml_text(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _elements(body: str, tag: str) -> list[str]:
    """Self-closing child elements only; anything else in ``body`` is unsupported."""
    items = re.findall(rf"<{tag}\b[^>]*?/>", body)
    rest = re.sub(rf"<{tag}\b[^>]*?/>", "", body)
    if rest.strip():
        raise SpliceUnsupported(f"unexpected {tag} content")
    return items


def _part_for_target(base_dir: str, target: str) -> str:
    if target.startswith("/"):
        return target[1:]
    return _pnormpath(_pjoin(base_dir, target))


def _rels_name(part: str) -> str:
    return _pjoin(_pdirname(part), "_rels", _pbasename(part) + ".rels")


def _abs_ref(ref: str) -> str:
    return ":".join(re.sub(r"([A-Z]+)(\d+)", r"$\1$\2", piece) for piece in ref.split(":"))


def _quoted_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _style_parts(styles: str, keys: list[str]) -> tuple[str, dict[str, int]]:
    """Append fonts/fills/xfs for ``keys``; return new styles.xml and xf ids."""
    if not keys:
        return styles, {}

    def grow(xml: str, tag: str, child: str, additions: list[str]) -> tuple[str, int]:
        match = re.search(rf"<{tag}\b([^>]*?)>(.*?)</{tag}>", xml, re.S)
        if match is None:
            raise SpliceUnsupported(f"styles without {tag}")
        attrs, body = match.group(1), match.group(2)
        start = len(re.findall(rf"<{child}\b", body))
        body = body + "".join(additions)
        count = len(re.findall(rf"<{child}\b", body))
        attrs = re.sub(r'\s*count="\d*"', "", attrs) + f' count="{count}"'
        return xml[:match.start()] + f"<{tag}{attrs}>{body}</{tag}>" + xml[match.end():], start

    fonts = re.search(r"<fonts\b[^>]*?>(.*?)</fonts>", styles, re.S)
    if fonts is None:
        raise SpliceUnsupported("styles without fonts")
    base = re.match(r"\s*<font\b[^>]*?(?:/>|>(.*?)</font>)", fonts.group(1), re.S)
    if base is None:
        raise SpliceUnsupported("styles without a default font")
    base_children = re.sub(r"<(?:b|color)\b[^>]*?/>", "", base.group(1) or "")
    font_xml, fill_xml, font_ids, fill_ids = [], [], {}, {}
    for key in keys:
        style = STYLES[key]
        if style.get("bold") or style.get("color"):
            font_ids[key] = len(font_xml)
            font_xml.append("<font>" + ("<b/>" if style.get("bold") else "")
                            + (f'<color rgb="00{style["color"]}"/>' if style.get("color") else "")
                            + base_children + "</font>")
        if style.get("fill"):
            fill_ids[key] = len(fill_xml)
            fill_xml.append(f'<fill><patternFill patternType="solid"><fgColor rgb="00{style["fill"]}"/>'
                            '<bgColor rgb="00000000"/></patternFill></fill>')
    styles, first_font = grow(styles, "fonts", "font", font_xml)
    styles, first_fill = grow(styles, "fills", "fill", fill_xml)
    xfs, xf_ids = [], {}
    cell_xfs = re.search(r"<cellXfs\b[^>]*?>(.*?)</cellXfs>", styles, re.S)
    if cell_xfs is None:
        raise SpliceUnsupported("styles without cellXfs")
    first_xf = len(re.findall(r"<xf\b", cell_xfs.group(1)))
    for key in keys:
        style = STYLES[key]
        font = first_font + font_ids[key] if key in font_ids else 0
        fill = first_fill + fill_ids[key] if key in fill_ids else 0
        xf = (f'<xf numFmtId="0" fontId="{font}" fillId="{fill}" borderId="0" xfId="0"'
              + (' applyFont="1"' if key in font_ids else "") + (' applyFill="1"' if key in fill_ids else ""))
        if style.get("wrap"):
            xf += ' applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>'
        else:
            xf += "/>"
        xf_ids[key] = first_xf + len(xfs)
        xfs.append(xf)
    styles, _ = grow(styles, "cellXfs", "xf", xfs)
    return styles, xf_ids


def sheet_xml(spec: SheetSpec, xf_ids: dict[str, int]) -> bytes:
    """Serialize one spec as a worksheet part with inline strings."""
    width, height = spec.max_col, len(spec.rows)
    parts = [f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
             f'<worksheet xmlns="{_MAIN_NS}" xmlns:r="{_REL_NS}">']
    parts.append(f'<dimension ref="A1{":" + _column_letter(width) + str(height) if width and height else ""}"/>')
    parts.append('<sheetViews><sheetView workbookViewId="0">')
    if spec.freeze and spec.freeze != (1, 1):
        row, col = spec.freeze
        cell = f"{_column_letter(col)}{row}"
        pane = "bottomRight" if row > 1 and col > 1 else "bottomLeft" if row > 1 else "topRight"
        parts.append("<pane" + (f' xSplit="{col - 1}"' if col > 1 else "") + (f' ySplit="{row - 1}"' if row > 1 else "")
                     + f' topLeftCell="{cell}" activePane="{pane}" state="frozen"/>'
                     f'<selection pane="{pane}" activeCell="{cell}" sqref="{cell}"/>')
    parts.append('</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/>')
    if spec.widths:
        parts.append("<cols>" + "".join(f'<col min="{col}" max="{col}" width="{width_value}" customWidth="1"/>'
                                        for col, width_value in sorted(spec.widths.items())) + "</cols>")
    parts.append("<sheetData>")
    current = None
    for r, c, value, style in spec.cells():
        if r != current:
            if current is not None:
                parts.append("</row>")
            height_attr = (f' ht="{spec.row_heights[r]}" customHeight="1"' if r in spec.row_heights else "")
            parts.append(f'<row r="{r}"{height_attr}>')
            current = r
        ref = f"{_column_letter(c)}{r}"
        style_attr = f' s="{xf_ids[style]}"' if style else ""
        if value is None:
            parts.append(f'<c r="{ref}"{style_attr}/>')
        elif isinstance(value, bool):
            parts.append(f'<c r="{ref}"{style_attr} t="b"><v>{int(value)}</v></c>')
        elif isinstance(value, (int, float, Decimal)):
            parts.append(f'<c r="{ref}"{style_attr}><v>{value}</v></c>')
        else:
            space = ' xml:space="preserve"' if value != value.strip() or "\n" in value else ""
            parts.append(f'<c r="{ref}"{style_attr} t="inlineStr"><is><t{space}>{_xml_text(value)}</t></is></c>')
    if current is not None:
        parts.append("</row>")
    parts.append("</sheetData>")
    if spec.auto_filter:
        parts.append(f'<autoFilter ref="{spec.auto_filter}"/>')
    parts.append('<pageMargins left="0.75" right="0.75" top="1" bottom="1" header="0.5" footer="0.5"/></worksheet>')
    return "".join(parts).encode("utf-8")


def _splice(path: Path, specs: list[SheetSpec]) -> None:
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as source:
        infos = source.infolist()
        names = [info.filename for info in infos]
        if len(set(names)) != len(names) or any(info.flag_bits & 0x1 for info in infos):
            raise SpliceUnsupported("duplicate or encrypted members")
        try:
            content_types = source.read("[Content_Types].xml").decode("utf-8")
            workbook = source.read("xl/workbook.xml").decode("utf-8")
            rels = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        except KeyError as exc:
            raise SpliceUnsupported("non-standard workbook part names") from exc

        sheets_block = re.search(r"<sheets\b[^>]*?>(.*?)</sheets>", workbook, re.S)
        if sheets_block is None:
            raise SpliceUnsupported("workbook without sheets")
        sheets = []
        for element in _elements(sheets_block.group(1), "sheet"):
            attrs = _attrs(element)
            rel_keys = [key for key in attrs if key.endswith(":id")]
            if "name" not in attrs or "sheetId" not in attrs or len(rel_keys) != 1:
                raise SpliceUnsupported("unexpected sheet entry")
            sheets.append({"element": element, "name": unescape(attrs["name"]),
                           "sheet_id": int(attrs["sheetId"]), "rid": attrs[rel_keys[0]]})
        rel_body = re.search(r"<Relationships\b[^>]*?>(.*?)</Relationships>", rels, re.S)
        if rel_body is None:
            raise SpliceUnsupported("unexpected workbook relationships")
        relationships = [(element, _attrs(element)) for element in _elements(rel_body.group(1), "Relationship")]
        by_id = {attrs.get("Id"): (element, attrs) for element, attrs in relationships}
        styles_targets = [attrs["Target"] for _, attrs in relationships
                          if attrs.get("Type", "").endswith("/styles") and "Target" in attrs]
        if len(styles_targets) != 1:
            raise SpliceUnsupported("workbook without one styles part")
        styles_name = _part_for_target("xl", styles_targets[0])
        try:
            styles = source.read(styles_name).decode("utf-8")
        except KeyError as exc:
            raise SpliceUnsupported("missing styles part") from exc

        wanted = {spec.name for spec in specs}
        if len(wanted) != len(specs):
            raise SpliceUnsupported("duplicate replacement sheet")
        removed_parts: set[str] = set()
        removed_indexes: set[int] = set()
        for index, sheet in enumerate(sheets):
            if sheet["name"] not in wanted:
                if sheet["name"].casefold() in {name.casefold() for name in wanted}:
                    raise SpliceUnsupported("sheet name differs only by case")
                continue
            element, attrs = by_id.get(sheet["rid"], (None, {}))
            if element is None or attrs.get("Type") != _WORKSHEET_TYPE or attrs.get("TargetMode") == "External":
                raise SpliceUnsupported("replaced sheet is not a plain worksheet")
            part = _part_for_target("xl", attrs["Target"])
            sheet_rels = _rels_name(part)
            if sheet_rels in names:
                body = source.read(sheet_rels).decode("utf-8")
                block = re.search(r"<Relationships\b[^>]*?>(.*?)</Relationships>", body, re.S)
                if block is None or any(_attrs(item).get("TargetMode") != "External"
                                        for item in _elements(block.group(1), "Relationship")):
                    raise SpliceUnsupported("replaced sheet owns other package parts")
                removed_parts.add(sheet_rels)
            removed_parts.add(part)
            removed_indexes.add(index)
        if removed_indexes and any(name == "xl/calcChain.xml" or name.startswith(("xl/pivotCache", "xl/externalLinks"))
                                   for name in names):
            raise SpliceUnsupported("workbook has parts that index sheets")
        removed_names = {sheets[index]["name"] for index in removed_indexes}
        removed_rids = {sheets[index]["rid"] for index in removed_indexes}

        kept = [sheet for index, sheet in enumerate(sheets) if index not in removed_indexes]
        new_index = {old: new for new, old in enumerate(i for i in range(len(sheets)) if i not in removed_indexes)}

        # Defined names: drop the removed sheets' local names, renumber others.
        # openpyxl writes an empty ``<definedNames />``; a second block is invalid.
        defined_pattern = r"<definedNames\b[^>]*?(?:/>|>(.*?)</definedNames>)"
        if len(re.findall(defined_pattern, workbook, re.S)) > 1:
            raise SpliceUnsupported("several defined name blocks")
        defined = re.search(defined_pattern, workbook, re.S)
        defined_items: list[str] = []
        if defined is not None:
            body = defined.group(1) or ""
            items = re.findall(r"<definedName\b[^>]*?(?:/>|>.*?</definedName>)", body, re.S)
            if re.sub(r"<definedName\b[^>]*?(?:/>|>.*?</definedName>)", "", body, flags=re.S).strip():
                raise SpliceUnsupported("unexpected defined names")
            for item in items:
                head = re.match(r"<definedName\b[^>]*?>", item).group()
                attrs = _attrs(head)
                if "localSheetId" in attrs:
                    old = int(attrs["localSheetId"])
                    if old in removed_indexes:
                        continue
                    if old not in new_index:
                        raise SpliceUnsupported("defined name refers to a missing sheet")
                    item = item.replace(f'localSheetId="{attrs["localSheetId"]}"', f'localSheetId="{new_index[old]}"', 1)
                else:
                    text = unescape(re.sub(r"<[^>]+>", "", item))
                    if any(name in text for name in removed_names):
                        raise SpliceUnsupported("global defined name refers to a replaced sheet")
                defined_items.append(item)

        all_numbers = [int(match) for name in names for match in re.findall(r"^xl/worksheets/sheet(\d+)\.xml$", name)]
        next_part = max(all_numbers, default=0) + 1
        next_sheet_id = max((sheet["sheet_id"] for sheet in sheets), default=0) + 1
        existing_rids = {attrs.get("Id") for _, attrs in relationships}

        style_keys = []
        for spec in specs:
            for key in [spec.cell_style, *spec.row_styles.values()]:
                if key and key not in style_keys:
                    style_keys.append(key)
        styles, xf_ids = _style_parts(styles, style_keys)

        sheet_elements = [sheet["element"] for sheet in kept]
        new_rels, new_overrides = [], []
        for offset, spec in enumerate(specs):
            part = f"xl/worksheets/sheet{next_part + offset}.xml"
            rid = f"rIdSL{next_sheet_id + offset}"
            while rid in existing_rids:
                rid += "x"
            existing_rids.add(rid)
            replacements[part] = sheet_xml(spec, xf_ids)
            sheet_elements.append(f'<sheet xmlns:r="{_REL_NS}" name="{_xml_attr(spec.name)}" '
                                  f'sheetId="{next_sheet_id + offset}" r:id="{rid}"/>')
            new_rels.append(f'<Relationship Id="{rid}" Type="{_WORKSHEET_TYPE}" '
                            f'Target="worksheets/sheet{next_part + offset}.xml"/>')
            new_overrides.append(f'<Override PartName="/{part}" ContentType="{_WORKSHEET_CONTENT}"/>')
            if spec.auto_filter:
                defined_items.append(f'<definedName name="_xlnm._FilterDatabase" localSheetId="{len(kept) + offset}" '
                                     f'hidden="1">{_xml_text(_quoted_sheet(spec.name))}!{_abs_ref(spec.auto_filter)}</definedName>')

        workbook = workbook[:sheets_block.start(1)] + "".join(sheet_elements) + workbook[sheets_block.end(1):]
        defined_xml = "<definedNames>" + "".join(defined_items) + "</definedNames>" if defined_items else ""
        if defined is not None:
            workbook = re.sub(defined_pattern, lambda _: defined_xml, workbook, count=1, flags=re.S)
        elif defined_xml:
            workbook = workbook.replace("</sheets>", "</sheets>" + defined_xml, 1)
        view = re.search(r"<workbookView\b[^>]*?/?>", workbook)
        if view is not None:
            element = view.group()
            for attribute in ("activeTab", "firstSheet"):
                value = _attrs(element).get(attribute)
                if value is None:
                    continue
                old = int(value)
                replacement = "" if old in removed_indexes or old not in new_index else f' {attribute}="{new_index[old]}"'
                element = re.sub(rf'\s{attribute}="{value}"', replacement, element, count=1)
            workbook = workbook[:view.start()] + element + workbook[view.end():]

        rel_items = [element for element, attrs in relationships if attrs.get("Id") not in removed_rids] + new_rels
        rels = rels[:rel_body.start(1)] + "".join(rel_items) + rels[rel_body.end(1):]
        for part in removed_parts:
            content_types = re.sub(rf'<Override\b[^>]*?PartName="/{re.escape(part)}"[^>]*?/>', "", content_types)
        if "</Types>" not in content_types:
            raise SpliceUnsupported("unexpected content types")
        content_types = content_types.replace("</Types>", "".join(new_overrides) + "</Types>", 1)

        replacements.update({"[Content_Types].xml": content_types.encode("utf-8"),
                             "xl/workbook.xml": workbook.encode("utf-8"),
                             "xl/_rels/workbook.xml.rels": rels.encode("utf-8"),
                             styles_name: styles.encode("utf-8")})
        for name, data in replacements.items():
            try:
                ElementTree.fromstring(data)
            except ElementTree.ParseError as exc:  # Never publish a malformed package.
                raise SpliceUnsupported(f"generated part is not well-formed: {name}") from exc

        temp = _temp_near(path)
        try:
            with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
                for info in infos:
                    if info.filename in removed_parts:
                        continue
                    data = replacements.pop(info.filename, None)
                    clone = zipfile.ZipInfo(info.filename, date_time=info.date_time)
                    clone.compress_type = zipfile.ZIP_DEFLATED
                    clone.external_attr = info.external_attr
                    target.writestr(clone, source.read(info) if data is None else data)
                for name, data in replacements.items():
                    target.writestr(name, data)
            validate_package(temp)
        except BaseException:
            temp.unlink(missing_ok=True)
            raise
    try:
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def replace_sheets(path: str | Path, specs: list[SheetSpec]) -> Path:
    """Atomically replace (or add) sheets at the end of an existing workbook.

    Replaced sheets move to the end, matching the former delete-and-append
    behaviour. Other worksheets are carried over byte-for-byte in content.
    """
    path = Path(path)
    if not specs:
        return path
    try:
        _splice(path, specs)
    except (SpliceUnsupported, zipfile.BadZipFile, UnicodeDecodeError, ValueError):
        _replace_openpyxl(path, specs)
    return path
