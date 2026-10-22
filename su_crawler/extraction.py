"""Evidence-preserving extraction for files and fetched web responses."""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import replace
from datetime import date, datetime
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from .models import Candidate, FetchResult, Source
from .adapters.common import display_state


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _evidence(location: str, raw: Any, *, display: str | None = None,
              proof_kind: str | None = None, hidden_by: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    proof: dict[str, Any] = {"location": location, "raw": raw}
    if display is not None:
        proof["display_state"] = display
    if proof_kind is not None:
        proof["proof_kind"] = proof_kind
    if hidden_by:
        proof["hidden_by"] = hidden_by
    return proof


def _json_value(value: Any) -> Any:
    """Keep file-derived values JSON-serializable without inventing a value."""
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def _candidate_from_mapping(
    values: dict[str, Any], *, locator: str, method: str, columns: dict[str, str] | None = None,
    evidence_mode: str = "structured_record",
) -> Candidate:
    fields: dict[str, Any] = {}
    evidence: dict[str, dict[str, Any]] = {}
    specs: dict[str, str] = {}
    mapping = columns or {key: key for key in values}
    for field, key in mapping.items():
        raw = _json_value(values.get(key))
        if raw is None:
            continue
        if field.startswith("spec:"):
            spec = field.removeprefix("spec:")
            specs[spec] = str(raw)
            evidence[field] = _evidence(locator, raw, display="not_applicable", proof_kind="record_value")
        else:
            fields[field] = raw
            evidence[field] = _evidence(locator, raw, display="not_applicable", proof_kind="record_value")
    return Candidate(fields=fields, evidence=evidence, locator=locator, extraction_method=method, specs=specs,
                     source_visibility="not_applicable", evidence_mode=evidence_mode)


def _displayed_text(node: Tag, mode: str) -> str:
    parts = []
    for value in node.descendants:
        if not isinstance(value, NavigableString) or not str(value).strip():
            continue
        parent = value.parent if isinstance(value.parent, Tag) else None
        if display_state(parent, mode)[0] == "hidden":
            continue
        parts.append(str(value).strip())
    return re.sub(r"\s+", " ", " ".join(parts)).strip()


def _selected_value(node: Any, selector: str, mode: str) -> tuple[Any, Any, bool]:
    attr = re.search(r"::attr\(([^)]+)\)$", selector)
    css = selector[: attr.start()] if attr else selector
    selected = node.select_one(css)
    if selected is None:
        return None, None, bool(attr)
    if attr:
        value = selected.get(attr.group(1))
    elif display_state(selected, mode)[0] == "hidden":
        # Preserve the observed raw value even when it must be excluded from
        # verification/comparison. Visible wrappers still omit hidden children.
        value = selected.get_text(" ", strip=True)
    else:
        value = _displayed_text(selected, mode)
    return value, selected, bool(attr)


def _extract_html(result: FetchResult, source: Source) -> list[Candidate]:
    soup = BeautifulSoup(result.content, "html.parser")
    mode = "rendered_dom" if result.backend == "playwright" else "static_html"
    candidates: list[Candidate] = []
    if source.selectors:
        rows = soup.select(source.row_selector) if source.row_selector else [soup]
        for index, row in enumerate(rows, start=1):
            fields: dict[str, Any] = {}
            evidence: dict[str, dict[str, Any]] = {}
            specs: dict[str, str] = {}
            for field, selector in source.selectors.items():
                raw, selected, is_attribute = _selected_value(row, selector, mode)
                if raw is None:
                    continue
                location = f"css:{source.row_selector or ':document'}[{index}] {selector}"
                if field.startswith("spec:"):
                    specs[field.removeprefix("spec:")] = str(raw)
                else:
                    fields[field] = raw
                state, hidden_by = display_state(selected, mode)
                # Attribute values are DOM metadata rather than displayed text.
                # They remain useful source proof, but do not establish display.
                if is_attribute and state == "visible":
                    state = "unconfirmed"
                evidence[field] = _evidence(
                    location, raw, display=state,
                    proof_kind="dom_attribute" if is_attribute else "rendered_text",
                    hidden_by=hidden_by,
                )
            _price_marker_currency(fields, evidence)
            if fields or specs:
                price_state = (evidence.get("price") or {}).get("display_state")
                visibility = price_state if price_state in {"visible", "hidden", "unconfirmed"} else (
                    "visible" if mode == "rendered_dom" else "unconfirmed"
                )
                candidates.append(Candidate(fields, evidence, f"css-row:{index}", "html_css", specs,
                                            source_visibility=visibility, evidence_mode=mode))
        from .network_evidence import corroborate
        return corroborate(candidates, result)

    # Share the same bounded structured-data parser, visible-price checks and
    # review semantics as catalog collection.
    from .generic_extraction import apply_market_currency, extract_generic

    return apply_market_currency(extract_generic(result, source), source.market_currency)


def _price_marker_currency(fields: dict[str, Any], evidence: dict[str, dict[str, Any]]) -> None:
    """A selected price such as ``12,900원`` also proves its currency.

    Only a marker that names exactly one currency qualifies; ``$`` or ``¥``
    alone leaves currency missing for review.
    """
    from .amounts import currency_codes, displayed_amount

    if "currency" in fields or not isinstance(fields.get("price"), str) or "price" not in evidence:
        return
    amount, _ = displayed_amount(fields["price"])
    codes = currency_codes(amount["marker"]) if amount else frozenset()
    if len(codes) != 1:
        return
    proof = evidence["price"]
    fields["currency"] = next(iter(codes))
    evidence["currency"] = {**{key: value for key, value in proof.items() if key != "raw"},
                            "raw": amount["marker"], "source_text": fields["price"],
                            "normalization": {"rule": "currency_marker"}}


def _extract_csv(result: FetchResult, source: Source) -> list[Candidate]:
    text = result.content.decode("utf-8-sig", errors="replace")
    return [
        _candidate_from_mapping(dict(row), locator=f"csv:row:{index}", method="csv", columns=source.columns)
        for index, row in enumerate(csv.DictReader(io.StringIO(text)), start=2)
    ]


def _extract_xlsx(result: FetchResult, source: Source) -> list[Candidate]:
    values_book = load_workbook(io.BytesIO(result.content), read_only=True, data_only=True)
    formulas_book = load_workbook(io.BytesIO(result.content), read_only=True, data_only=False)
    try:
        sheet_name = source.sheet or values_book.sheetnames[0]
        values_ws, formulas_ws = values_book[sheet_name], formulas_book[sheet_name]
        value_rows, formula_rows = values_ws.iter_rows(values_only=True), formulas_ws.iter_rows(values_only=True)
        headers = next(value_rows, ())
        next(formula_rows, ())
        header_index = {str(value): index for index, value in enumerate(headers) if value is not None}
        candidates: list[Candidate] = []
        for row_number, (row, formula_row) in enumerate(zip(value_rows, formula_rows), start=2):
            fields: dict[str, Any] = {}
            evidence: dict[str, dict[str, Any]] = {}
            specs: dict[str, str] = {}
            for field, header in source.columns.items():
                index = header_index.get(header)
                if index is None or index >= len(row):
                    continue
                raw, formula = row[index], formula_row[index]
                location = f"xlsx:{sheet_name}!{get_column_letter(index + 1)}{row_number}"
                if isinstance(formula, str) and formula.startswith("=") and raw is None:
                    evidence[field] = {**_evidence(location, formula, display="not_applicable", proof_kind="record_value"), "formula_cache_missing": True}
                    continue
                if raw is None:
                    continue
                raw = _json_value(raw)
                if field.startswith("spec:"):
                    specs[field.removeprefix("spec:")] = str(raw)
                else:
                    fields[field] = raw
                evidence[field] = _evidence(location, raw, display="not_applicable", proof_kind="record_value")
            if fields or specs or evidence:
                candidates.append(Candidate(fields, evidence, f"xlsx:{sheet_name}:row:{row_number}", "xlsx", specs,
                                            source_visibility="not_applicable", evidence_mode="structured_record"))
        return candidates
    finally:
        values_book.close()
        formulas_book.close()


def _extract_json(result: FetchResult, source: Source) -> list[Candidate]:
    payload = json.loads(result.content.decode("utf-8-sig"), parse_float=str)
    rows = payload if isinstance(payload, list) else payload.get("items", payload.get("products", [payload])) if isinstance(payload, dict) else []
    return [
        _candidate_from_mapping(row, locator=f"json:item:{index}", method="json", columns=source.columns)
        for index, row in enumerate(rows, start=1) if isinstance(row, dict)
    ]


def _extract_pdf(result: FetchResult, source: Source) -> list[Candidate]:
    if not source.pdf_pattern:
        return []
    reader = PdfReader(io.BytesIO(result.content))
    pattern = re.compile(source.pdf_pattern or "", re.MULTILINE)
    candidates: list[Candidate] = []
    any_text = False
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        any_text = any_text or bool(text.strip())
        for match_index, match in enumerate(pattern.finditer(text), start=1):
            values = {key: value for key, value in match.groupdict().items() if value is not None}
            candidates.append(_candidate_from_mapping(values, locator=f"pdf:page:{page_number}:match:{match_index}", method="pdf_text_pattern", evidence_mode="document_text"))
    if not candidates and not any_text:
        return [Candidate({}, {"_extraction": {"location": "pdf", "raw": "no embedded text; OCR not implemented"}},
                          "pdf", "pdf_ocr_not_implemented", source_visibility="not_applicable",
                          evidence_mode="document_text")]
    return candidates


def extract(result: FetchResult, source: Source) -> list[Candidate]:
    """Extract only source-supplied values; invalid/unavailable fetches produce no values."""
    if result.status != "fetched":
        return []
    if result.variant_snapshots:
        candidates: list[Candidate] = []
        seen_scroll_rows: set[str] = set()
        for index, snapshot in enumerate(result.variant_snapshots):
            content = snapshot.get("content")
            if not isinstance(content, bytes):
                continue
            item = replace(result, content=content, final_url=snapshot.get("final_url", result.final_url),
                           fetched_at=snapshot.get("fetched_at", result.fetched_at), variant_snapshots=[])
            selections = snapshot.get("selections") or {}
            selection_values = snapshot.get("selection_values") or {}
            price_settlement = snapshot.get("price_settlement")
            settlement_status = (price_settlement.get("status") if isinstance(price_settlement, dict)
                                 and isinstance(price_settlement.get("status"), str) else None)
            for candidate in extract(item, source):
                if isinstance(selections, dict) and "scroll_items" in selections:
                    key = json.dumps([candidate.locator, candidate.fields, candidate.specs],
                                     ensure_ascii=False, sort_keys=True, default=str)
                    if key in seen_scroll_rows:
                        continue
                    seen_scroll_rows.add(key)
                candidate.locator = f"capture:{index + 1}/{candidate.locator}"
                candidate.derived_values["capture_snapshot_index"] = index
                # A visible ready marker may already be visible from the prior
                # choice. It does not prove that a quoted price has settled;
                # only a reviewed price element that changed after the
                # selection, once the network was quiet, does.
                if selection_values:
                    if settlement_status is not None:
                        candidate.derived_values["price_settlement"] = settlement_status
                    if settlement_status == "identical_after_quiet":
                        candidate.review_flags.append("identical_amounts_after_network_quiet_review")
                    elif settlement_status != "changed":
                        candidate.review_flags.append("selection_price_settlement_unconfirmed")
                if isinstance(selections, dict):
                    for selector, label in selections.items():
                        if selector in {"page", "scroll_items"}:
                            continue
                        candidate.specs[f"selected:{selector}"] = str(label)
                        candidate.evidence[f"spec:selected:{selector}"] = _evidence(
                            f"capture:{index + 1} {selector}", str(label),
                            display="visible", proof_kind="browser_selection"
                        )
                        if isinstance(selection_values, dict) and selector in selection_values:
                            value = str(selection_values[selector])
                            candidate.specs[f"selected_value:{selector}"] = value
                            candidate.evidence[f"spec:selected_value:{selector}"] = _evidence(
                                f"capture:{index + 1} {selector}::value", value,
                                display="unconfirmed", proof_kind="source_metadata"
                            )
                    candidate.evidence["_selection"] = _evidence(
                        f"capture:{index + 1}", dict(selections), display="visible", proof_kind="browser_selection"
                    )
                candidates.append(candidate)
        return candidates
    if source.adapter:
        from .adapters import extract_adapter

        return extract_adapter(result, source)
    media_type = result.media_type.lower()
    if "html" in media_type:
        return _extract_html(result, source)
    if "csv" in media_type:
        return _extract_csv(result, source)
    if "spreadsheet" in media_type or "excel" in media_type or result.content[:2] == b"PK":
        return _extract_xlsx(result, source)
    if "json" in media_type:
        return _extract_json(result, source)
    if "pdf" in media_type or result.content[:4] == b"%PDF":
        return _extract_pdf(result, source)
    return []
