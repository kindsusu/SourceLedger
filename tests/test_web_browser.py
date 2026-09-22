"""Real local browser flow; no market websites or external services are contacted."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import shutil
import threading
import time

import pytest

from su_crawler.assistant_runtime import runtime_status, stop_worker
from su_crawler.browser_runtime import select_browser_runtime
from su_crawler.web_server import build_web_server


def launch_browser(driver):
    runtime = select_browser_runtime(driver.chromium)
    if runtime is None:
        pytest.skip("No local Chromium runtime installed")
    return driver.chromium.launch(headless=True, **runtime.launch_options())


def test_browser_onboarding_job_and_report(tmp_path):
    playwright = pytest.importorskip("playwright.sync_api")
    root = Path(__file__).resolve().parents[1]
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    fixtures = workspace / "fixtures"
    fixtures.mkdir()
    config = json.loads((root / "examples/verification.json").read_text(encoding="utf-8"))
    config["output_dir"] = "outputs"
    for source in config["sources"]:
        original = root / "examples" / source["location"]
        shutil.copyfile(original, fixtures / original.name)
    (workspace / "verification.json").write_text(json.dumps(config), encoding="utf-8")
    server = build_web_server(workspace, port=0)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with playwright.sync_playwright() as driver:
            browser = launch_browser(driver)
            page = browser.new_page(viewport={"width": 1440, "height": 1000})
            page_errors = []
            console_errors = []
            page.on("pageerror", lambda error: page_errors.append(str(error)))
            page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
            page.goto(server.url)
            page.get_by_label("Industry", exact=True).fill("Synthetic components")
            page.get_by_label("Product", exact=True).fill("Test part")
            page.get_by_label("Market", exact=True).fill("Offline fixture")
            page.get_by_role("button", name="Create workspace").click()
            page.locator('nav a[data-route="sources"]').click()
            page.get_by_label("Source URL").fill("https://example.invalid/catalog")
            page.locator('#source-form input[name="name"]').fill('<img src=x onerror=alert(1)>')
            page.get_by_role("button", name="Register source").click()
            playwright.expect(page.locator("#source-list")).to_contain_text("example.invalid")
            playwright.expect(page.locator("#source-list")).to_contain_text('<img src=x onerror=alert(1)>')
            assert page.locator("#source-list img").count() == 0
            playwright.expect(page.locator("#identifier-rows .key-value-row")).to_have_count(1)
            fields = page.locator("#identifier-rows .key-value-row").first.locator("input")
            fields.nth(0).fill("model")
            fields.nth(1).fill("TEST-A")
            page.get_by_role("button", name="Save product identity").click()
            page.get_by_label("Source URL").fill("https://example.invalid/unsaved")
            page.get_by_label("Source URL").focus()
            # Wait for an actual polling response and verify the editing context survived.
            with page.expect_response(lambda response: response.url.endswith("/api/bootstrap"), timeout=10000):
                pass
            playwright.expect(page.get_by_label("Source URL")).to_have_value("https://example.invalid/unsaved")
            playwright.expect(page.get_by_label("Source URL")).to_be_focused()
            for width in (375, 768, 1024, 1440):
                page.set_viewport_size({"width": width, "height": 1000})
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), width
            page.locator('nav a[data-route="runs"]').click()
            page.locator("details.advanced summary").click()
            page.locator("#advanced-operation").select_option("run")
            page.locator('#advanced-fields input[name="config_path"]').fill("verification.json")
            page.get_by_role("button", name="Queue advanced job").click()
            playwright.expect(page.locator("#job-list")).to_contain_text("Queued", ignore_case=True)
            page.get_by_role("button", name="Start worker", exact=True).click()
            playwright.expect(page.locator("#job-detail")).to_contain_text("Succeeded", ignore_case=True, timeout=30000)
            with page.expect_download() as download_event:
                page.get_by_role("link", name="Download XLSX").click()
            download = download_event.value
            assert download.failure() is None
            from openpyxl import load_workbook
            saved_report = tmp_path / "downloaded-report.xlsx"
            download.save_as(saved_report)
            workbook = load_workbook(saved_report, read_only=True)
            try:
                assert "Source Evidence" in workbook.sheetnames
            finally:
                workbook.close()
            assert page.locator("#job-detail table tbody tr").count() >= 2
            if importlib.util.find_spec("mcp") is not None:
                page.locator('nav a[data-route="connections"]').click()
                for client in ("codex", "claude-code", "claude-desktop"):
                    page.locator('#connection-form select[name="client"]').select_option(client)
                    with page.expect_response(lambda response: response.url.endswith("/api/connections")):
                        page.get_by_role("button", name="Generate snippet").click()
                    playwright.expect(page.locator("#connection-snippet")).to_contain_text("serve-assistant")
            assert page_errors == []
            assert console_errors == []
            browser.close()
    finally:
        stop_worker(workspace)
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline and runtime_status(workspace)["status"] != "stopped":
            time.sleep(0.1)
        assert runtime_status(workspace)["status"] == "stopped"
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_observed_amount_stays_separate_from_calculator_estimate(tmp_path):
    playwright = pytest.importorskip("playwright.sync_api")
    workspace = tmp_path / "workspace"
    server = build_web_server(workspace, port=0)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    job = {
        "id": "job-observation", "operation": "run", "status": "succeeded", "attempt": 1,
        "created_at": "2026-01-01T00:00:00Z", "updated_at": "2026-01-01T00:01:00Z",
        "result": {
            "execution_status": "succeeded", "evidence_status": "completed",
            "run_id": "run-observation", "output_dir": str(workspace / "outputs"),
        },
    }
    bootstrap = {
        "csrf_token": "test-token", "version": "test", "workspace_root": str(workspace),
        "status": "configured", "workspace": {
            "industry": "Fixtures", "market": "Offline", "product": {
                "name": "Test part", "identifiers": {"model": "TEST-A"}, "required_specs": {},
            }, "sources": [],
        },
        "research": None, "worker": {"status": "stopped"}, "jobs": [job], "mcp_available": False,
    }
    observations = {
        "total": 1, "offset": 0, "limit": 50, "rows": [{
            "source_id": "source-fixture", "source_name": "Fixture source",
            "source_url": "https://example.invalid/item", "amount": None, "currency": None,
            "raw_fields": {"price": "Call for quote"}, "value_origin": "calculator_estimate",
            "derived_amount": "123", "verification_level": "review",
            "collected_at": "2026-01-01T00:00:30Z",
        }],
    }
    try:
        with playwright.sync_playwright() as driver:
            browser = launch_browser(driver)
            page = browser.new_page(viewport={"width": 1024, "height": 800})
            page.route("**/api/bootstrap", lambda route: route.fulfill(json=bootstrap))
            page.route("**/api/jobs/job-observation/observations?*", lambda route: route.fulfill(json=observations))
            page.route("**/api/jobs/job-observation", lambda route: route.fulfill(json=job))
            page.goto(f"{server.url}/#runs")
            page.get_by_role("button", name="Run job, Succeeded").click()
            row = page.locator("#job-detail table tbody tr")
            playwright.expect(row).to_have_count(1)
            headers = page.locator("#job-detail table th").all_text_contents()
            cells = row.locator("td").all_text_contents()
            observed = cells[headers.index("Observed amount")]
            estimate = cells[headers.index("Estimated amount")]
            assert observed == "—"
            assert estimate == "123"
            assert observed not in {"Call for quote", "123"}
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
