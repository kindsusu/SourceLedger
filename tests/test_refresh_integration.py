"""Reviewed repeat settings are shared by UI, MCP, and collection snapshots."""
import asyncio

import pytest

from su_crawler import research_plans as plans
from su_crawler.execution_policy import build_execution_brief


def test_refresh_policy_requires_revision_review_and_survives_preview(tmp_path):
    plan = plans.create_plan(tmp_path, 'Inspect components')['plan']
    assert plan['refresh_policy'] == {'mode': 'full', 'full_interval_days': 7}
    plan = plans.revise_plan(tmp_path, plan['id'], plan['revision'], {
        'topic': {'industry': 'Parts', 'product': 'Bolt', 'market': 'US'},
        'added_candidates': [{'url': 'https://example.test/catalog', 'kind': 'category'}],
        'refresh_policy': {'mode': 'incremental', 'full_interval_days': 14},
    })['plan']
    plan = plans.confirm_plan(tmp_path, plan['id'], plan['revision'], True)['plan']
    original = plan
    plan = plans.revise_plan(tmp_path, plan['id'], plan['revision'], {
        'refresh_policy': {'mode': 'full', 'full_interval_days': 14}})['plan']
    assert plan['state'] == 'preview' and plan['confirmed_at'] is None
    assert plans.confirmed_snapshot(tmp_path, original['id'], original['revision'])['refresh_policy']['mode'] == 'incremental'
    brief = build_execution_brief(plan)
    assert brief['refresh_policy'] == {'mode': 'full', 'full_interval_days': 14}
    plan = plans.submit_plan_preview(tmp_path, plan['id'], plan['revision'], {
        'topic': plan['topic'], 'candidates': [], 'summary': 'Refined request'})['plan']
    assert plan['refresh_policy']['mode'] == 'full'


@pytest.mark.parametrize('value', [{'mode': 'skip'}, {'full_interval_days': 0},
    {'full_interval_days': True}, {'full_interval_days': 1.2}, {'full_interval_days': 366},
    {'secret': 'no'}, []])
def test_invalid_repeat_settings_rejected_without_revision_change(tmp_path, value):
    plan = plans.create_plan(tmp_path, 'Inspect components')['plan']
    with pytest.raises(ValueError):
        plans.revise_plan(tmp_path, plan['id'], plan['revision'], {'refresh_policy': value})
    assert plans.get_plan(tmp_path, plan['id'])['plan']['revision'] == 1


def test_mcp_repeat_policy_roundtrip(tmp_path):
    pytest.importorskip('mcp')
    from su_crawler.assistant_server import build_assistant_server
    server = build_assistant_server(tmp_path)
    def call(name, args):
        return asyncio.run(server.call_tool(name, args))[1]
    plan = call('create_research_plan', {'request_text': 'Inspect parts'})['plan']
    changed = call('update_research_plan', {'plan_id': plan['id'], 'expected_revision': plan['revision'],
        'changes': {'refresh_policy': {'mode': 'incremental', 'full_interval_days': 2}}})['plan']
    brief = call('get_research_execution_plan', {'plan_id': plan['id'], 'expected_revision': changed['revision']})
    assert brief['execution_plan']['refresh_policy']['full_interval_days'] == 2


def test_batch_refresh_summary_is_bounded_and_preserves_totals(tmp_path):
    from su_crawler.research_batch import _summary
    batches = []
    for mode in ('full', 'incremental'):
        batches.append({'status': 'completed', 'coverage_counts': {'visited': 40},
            'coverage_summary': {'complete': 40}, 'target_count': 40, 'budget_seconds_used': 1,
            'refresh': {'effective_mode': mode, 'reason': 'first_run' if mode == 'full' else 'scheduled_incremental',
                'counts': {'added': 2, 'unchanged': 38},
                'changes': [{'url': f'https://example.test/{mode}/{i}'} for i in range(40)],
                'item_counts': {'baseline': 40, 'price_changed': 30},
                'item_changes': [{'status': 'baseline'} for _ in range(70)],
                'item_changes_total': 70, 'item_changes_truncated': 0,
                'changes_truncated': 0}})
    summary = _summary({'batches': batches, 'batch_count': 2, 'target_count': 80,
        'job_id': 'test', 'started_at': None, 'finished_at': None}, output=tmp_path)
    refresh = summary['refresh']
    assert refresh['effective_mode'] == 'mixed'
    assert refresh['counts'] == {'added': 4, 'unchanged': 76}
    assert len(refresh['changes']) == 50 and refresh['changes_truncated'] == 30
    assert refresh['item_counts'] == {'baseline': 80, 'price_changed': 60}
    assert len(refresh['item_changes']) == 100
    assert refresh['item_changes_total'] == 140 and refresh['item_changes_truncated'] == 40


def test_web_refresh_policy_roundtrip(tmp_path):
    from test_web_server import running_server, authorized, json_response
    with running_server(tmp_path) as server:
        status, created = json_response(authorized(server, 'POST', '/api/plans', {'request_text': 'Inspect parts'}))
        assert status == 200
        plan = created['plan']
        status, edited = json_response(authorized(server, 'POST', '/api/plans/edit', {
            'plan_id': plan['id'], 'expected_revision': plan['revision'],
            'changes': {'refresh_policy': {'mode': 'incremental', 'full_interval_days': 3}}}))
        assert status == 200
        assert edited['plan']['refresh_policy']['full_interval_days'] == 3


"""Product-level refresh evidence, durable history and report regeneration."""
import json

from openpyxl import load_workbook

from su_crawler.catalog_collection import collect_plan
from su_crawler.catalog_refresh import item_changes, refresh_for_run
from su_crawler.models import FetchResult


def candidate(sku='A', price='100', **fields):
    return {'fields': {'sku': sku, 'name': 'Product', 'price': price, 'currency': 'USD', **fields}, 'specs': {}}


def test_product_changes_distinguish_prices_conditions_and_missing_evidence():
    old = candidate()
    assert item_changes([], [old], baseline=True)[0]['status'] == 'baseline'
    assert item_changes([], [old])[0]['status'] == 'added'
    changed = item_changes([old], [candidate(price='120')])[0]
    assert changed['status'] == 'price_changed'
    assert changed['condition_comparison'] == 'not_observed'
    assert changed['before']['fields']['price'] == '100'
    assert changed['after']['fields']['price'] == '120'
    assert changed['changed_fields'] == ['price']
    condition = item_changes([old], [candidate(tax='excluded')])[0]
    assert condition['status'] == 'conditions_changed'
    assert condition['changed_fields'] == ['tax']
    mixed = item_changes([candidate(deposit_amount='500000')],
                         [candidate(price='120', deposit_amount='600000')])[0]
    assert mixed['status'] == 'price_and_conditions_changed'
    assert mixed['changed_fields'] == ['deposit_amount', 'price']
    assert mixed['condition_comparison'] == 'observed_fields_only'
    renamed = item_changes([candidate(name='Original')],
                           [candidate(name='Renamed', price='120')])[0]
    assert renamed['status'] == 'price_changed'
    assert renamed['changed_fields'] == ['name', 'price']
    inquiry = candidate(price=None, price_type='inquiry')
    transition = item_changes([old], [inquiry])[0]
    assert transition['status'] == 'price_and_conditions_changed'
    assert transition['after']['fields']['price'] is None
    assert item_changes([candidate(price_type='sale')],
                         [candidate(price_type='list')])[0]['status'] == 'conditions_changed'
    spec_incomplete = {**candidate(), 'specs': {'grade': None}}
    assert item_changes([spec_incomplete], [{**candidate(price='120'),
                                            'specs': {'grade': None}}])[0]['condition_comparison'] == 'incomplete'
    assert item_changes([old], [], successful=False)[0]['after'] is None
    assert item_changes([], [], successful=False)[0]['product_name'] == ''
    assert item_changes([old], [], inventory_complete=True)[0]['status'] == 'removed_candidate'


def test_ambiguous_identifiers_never_create_one_to_many_price_matches():
    old = candidate()
    rows = item_changes([old], [candidate(price='110'), candidate(price='120')])
    assert [row['status'] for row in rows] == ['added', 'added', 'unavailable']
    rows = item_changes([old, old], [candidate(price='120')])
    assert [row['status'] for row in rows] == ['added', 'unavailable', 'unavailable']
    variant = {**candidate(price='120'), 'specs': {'size': 'XL'}}
    assert [row['status'] for row in item_changes([old], [variant])] == ['added', 'unavailable']


def test_contract_terms_identify_distinct_offers_without_matching_mutable_conditions():
    old = [candidate(term_months=term, deposit_amount='500000', annual_mileage_km='30000')
           for term in ('12', '24', '36', '48', '60')]
    new = [candidate(term_months=term, price='120', deposit_amount='600000', annual_mileage_km='25000')
           for term in ('60', '48', '36', '24', '12')]
    rows = item_changes(old, new)
    assert len(rows) == 5 and all(row['status'] == 'price_and_conditions_changed' for row in rows)
    assert all(row['before']['fields']['term_months'] == row['after']['fields']['term_months'] for row in rows)
    assert all(row['changed_fields'] == ['annual_mileage_km', 'deposit_amount', 'price'] for row in rows)
    ambiguous = item_changes([old[0]], [new[-1], new[-1]])
    assert [row['status'] for row in ambiguous] == ['added', 'added', 'unavailable']


def test_contract_mode_and_price_basis_do_not_pair_different_offers():
    old = candidate(term_months='36', contract_mode='rent', price_basis='monthly')
    daily = candidate(term_months='36', contract_mode='rent', price_basis='daily', price='120')
    lease = candidate(term_months='36', contract_mode='lease', price_basis='monthly', price='120')
    assert [row['status'] for row in item_changes([old], [daily])] == ['added', 'unavailable']
    assert [row['status'] for row in item_changes([old], [lease])] == ['added', 'unavailable']


def test_full_default_retains_all_item_changes_and_regenerates_report(tmp_path):
    url = 'https://shop.example/catalog'
    plan = {'id': 'product-history', 'revision': 1, 'state': 'confirmed',
            'topic': {'industry': 'retail', 'product': 'items', 'market': 'public'},
            'candidates': [{'id': 'catalog', 'name': 'Catalog', 'url': url, 'selected': True}]}
    state = {'price': '100', 'failed': False, 'count': 105, 'availability': 'InStock'}

    def fetch(source, _base, backend, **kwargs):
        if backend != 'http' or state['failed']:
            return FetchResult(source.id, 'failed', backend, final_url=url)
        products = [{'@type': 'Product', 'name': '=Product', 'sku': f'SKU-{number}',
                     'offers': {'@type': 'Offer', 'price': state['price'], 'priceCurrency': 'USD',
                                'availability': state['availability']}}
                    for number in range(state['count'])]
        html = '<script type="application/ld+json">' + json.dumps(products) + '</script>'
        return FetchResult(source.id, 'fetched', backend, content=html.encode(), final_url=url)

    output = tmp_path / 'out'
    first = collect_plan(tmp_path, plan, output_dir=output, collector=fetch, max_pages=1, follow_links=False)
    assert first['refresh']['policy']['mode'] == 'full'
    assert first['refresh']['item_counts']['baseline'] == 105
    assert first['refresh']['counts']['added'] == 0
    state['price'] = '120'
    second = collect_plan(tmp_path, plan, output_dir=output, collector=fetch, max_pages=1, follow_links=False)
    refresh = second['refresh']
    assert refresh['effective_mode'] == 'full'
    assert refresh['item_counts']['price_changed'] == 105
    assert len(refresh['item_changes']) == 100
    assert refresh['item_changes_total'] == 105 and refresh['item_changes_truncated'] == 5
    durable = refresh_for_run(output, second['run_id'])
    assert len(durable['item_changes']) == 105
    assert durable['item_changes_truncated'] == 0

    from su_crawler.config import load_config
    from su_crawler.pipeline import report_for_run
    from su_crawler.storage import Store
    store = Store(output)
    try:
        report_for_run(load_config(second['config_path']), store, second['run_id'])
    finally:
        store.close()
    book = load_workbook(second['report_path'])
    try:
        assert book['Collection Changes']['D3'].value is None
        sheet = book['Item Changes']
        assert sheet.max_row == 106
        assert sheet['A2'].value == 'Price changed'
        assert sheet['K2'].value.startswith('The displayed price changed.')
        assert 'unrecorded terms remain unknown' in sheet['K2'].value
        assert sheet['J2'].value == 'Price'
        assert sheet['H2'].value == 'Availability: InStock'
        assert sheet['H2'].alignment.wrap_text
        assert sheet['A1'].font.bold and sheet['A1'].fill.fgColor.rgb == '0017365D'
        assert sheet['B2'].value == '=Product' and sheet['B2'].data_type == 's'
        assert sheet['D2'].value == '100' and sheet['E2'].value == '120'
        assert sheet['F2'].value == 'USD' and sheet['G2'].value == 'USD'
    finally:
        book.close()
    state['availability'] = 'PreOrder'
    state['count'] = 106
    conditions = collect_plan(tmp_path, plan, output_dir=output, collector=fetch, max_pages=1, follow_links=False)
    assert conditions['refresh']['item_counts']['conditions_changed'] == 105
    assert conditions['refresh']['item_counts']['added'] == 1
    assert conditions['refresh']['item_counts']['price_changed'] == 0
    state['price'] = 'Contact for quote'
    inquiry = collect_plan(tmp_path, plan, output_dir=output, collector=fetch, max_pages=1, follow_links=False)
    assert inquiry['refresh']['item_counts']['price_changed'] == 106
    assert inquiry['refresh']['item_changes'][0]['after']['fields']['price'] == 'Contact for quote'
    state['failed'] = True
    failed = collect_plan(tmp_path, plan, output_dir=output, collector=fetch, max_pages=1, follow_links=False)
    assert failed['refresh']['item_counts']['unavailable'] == 106
    assert failed['refresh']['item_changes'][0]['after'] is None
    assert not failed['refresh']['full_recheck_succeeded']


def test_failed_first_page_is_reported_without_inventing_a_product(tmp_path):
    plan = {'id': 'empty-failure', 'revision': 1, 'state': 'confirmed',
            'topic': {'industry': 'retail', 'product': 'items', 'market': 'public'},
            'candidates': [{'id': 'one', 'name': 'Catalog', 'url': 'https://shop.example/item', 'selected': True}]}

    def failed(source, _base, backend, **kwargs):
        return FetchResult(source.id, 'failed', backend)

    result = collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=failed,
                          max_pages=1, follow_links=False)
    row = result['refresh']['item_changes'][0]
    assert row['status'] == 'unavailable' and row['product_name'] == ''
    assert row['before'] is None and row['after'] is None


def test_item_change_report_explains_price_availability_and_preserves_values(tmp_path):
    from openpyxl import Workbook
    from su_crawler.report import append_refresh_summary
    path = tmp_path / 'readable.xlsx'
    book = Workbook()
    book.save(path)
    book.close()
    previous = candidate(price='100', tax=0)
    previous['specs'] = {'size': 'M'}
    current = candidate(price='120', tax='included')
    current['specs'] = {'size': 'M'}
    rows = item_changes([previous], [current])
    rows += item_changes([candidate()], [candidate(price=None)])
    rows += item_changes([candidate()], [candidate(price='Contact for quote')])
    rows += item_changes([candidate()], [candidate(currency='EUR')])
    rows += item_changes([], [candidate(price=0)], baseline=True)
    append_refresh_summary(path, {'item_changes': rows})
    book = load_workbook(path)
    try:
        sheet = book['Item Changes']
        assert sheet['A2'].value == 'Price and observed conditions changed'
        assert 'not a like-for-like price comparison' in sheet['K2'].value
        assert sheet['H2'].value == 'Tax: 0\nSize: M'
        assert sheet['I2'].value == 'Tax: included\nSize: M'
        assert sheet['A3'].value == 'Price no longer displayed'
        assert sheet['E3'].value is None
        assert 'current price is unknown' in sheet['K3'].value
        assert sheet['A4'].value == 'Price inquiry required'
        assert sheet['E4'].value == 'Contact for quote'
        assert 'No numeric current price is inferred' in sheet['K4'].value
        assert sheet['A5'].value == 'Currency changed'
        assert sheet['K5'].value.startswith('The displayed currency changed; the displayed price text is unchanged.')
        assert sheet['A6'].value == 'First capture'
        assert sheet['D6'].value is None and sheet['E6'].value == 0
        assert sheet['E6'].data_type == 'n'
    finally:
        book.close()
