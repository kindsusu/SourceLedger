from __future__ import annotations

from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import json
import os
import subprocess
import sys
import threading

from openpyxl import load_workbook

from su_crawler.storage import Store


ROOT = Path(__file__).resolve().parents[1]


class ResearchSite(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/robots.txt":
            self._send(b"User-agent: *\nAllow: /\n", "text/plain")
        elif self.path == "/catalog":
            self._send(
                b'<html><body><a href="/product/test-a">Product A</a>'
                b'<a href="/about">About</a><a href="/product/test-a#details">Duplicate</a></body></html>'
            )
        elif self.path == "/product/test-a":
            self._send(b"""<!doctype html><html><body><article class="product">
                <span class="model">TEST-A</span>
                <span class="price">12000</span>
                <span class="currency">KRW</span>
                <span class="unit">each</span>
                <span class="pack">1</span>
                <span class="tax">included</span>
                <span class="price-type">retail</span>
                <span class="basis">each</span>
                <span class="grade">industrial</span>
            </article></body></html>""")
        elif self.path == "/about":
            self._send(b"<html><body>Company information only</body></html>")
        else:
            self._send(b"missing", status=404, media_type="text/plain")

    def _send(self, body: bytes, media_type: str = "text/html", status: int = 200):
        self.send_response(status)
        self.send_header("Content-Type", f"{media_type}; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


@contextmanager
def research_site():
    server = ThreadingHTTPServer(("127.0.0.1", 0), ResearchSite)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def cli(*arguments: str, expected: int = 0) -> dict:
    environment = dict(os.environ, PYTHONIOENCODING="utf-8")
    result = subprocess.run(
        [sys.executable, "-m", "su_crawler", *arguments],
        cwd=ROOT, env=environment, capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == expected, result.stderr or result.stdout
    return json.loads(result.stdout)


def test_cli_research_discovery_verification_activation_and_collection(tmp_path):
    workspace = tmp_path / "research.json"
    draft_path = tmp_path / "draft.json"
    samples_path = tmp_path / "known-samples.json"
    receipt_path = tmp_path / "verification.json"
    active_path = tmp_path / "active.json"

    with research_site() as origin:
        initialized = cli(
            "init", "--workspace", str(workspace), "--industry", "Industrial pumps",
            "--product", "Process pump", "--market", "South Korea",
        )
        assert initialized["locale"] == "en"
        assert initialized["product"]["identifiers"] == {}

        added = cli(
            "source-add", "--workspace", str(workspace), "--url", origin + "/catalog",
            "--scope", "internal", "--name", "Controlled catalog",
        )
        seed_id = added["sources"][0]["id"]
        discovered = cli(
            "source-discover", "--workspace", str(workspace), "--source-id", seed_id, "--limit", "10",
        )
        assert discovered["status"] == "discovered"
        assert discovered["added_count"] == 2

        product = cli(
            "product-set", "--workspace", str(workspace), "--identifier", "model=TEST-A",
            "--spec", "grade=industrial",
        )
        assert product["product"]["identifiers"] == {"model": "TEST-A"}
        draft_result = cli("draft", "--workspace", str(workspace), "--output", str(draft_path))
        assert draft_result["ready"] is False

        draft = json.loads(draft_path.read_text(encoding="utf-8"))
        product_source = next(source for source in draft["sources"] if source["location"] == origin + "/product/test-a")
        # Candidate discovery cannot infer extraction rules. A human-reviewed rule
        # and explicit source selection are required before verification.
        product_source.update({
            "row_selector": ".product",
            "selectors": {
                "model": ".model", "price": ".price", "currency": ".currency",
                "unit": ".unit", "pack_quantity": ".pack", "tax": ".tax",
                "price_type": ".price-type", "price_basis": ".basis", "spec:grade": ".grade",
            },
        })
        draft["sources"] = [product_source]
        draft_path.write_text(json.dumps(draft, ensure_ascii=False, indent=2), encoding="utf-8")

        samples = {
            "schema": "source-ledger/known-samples/v1",
            "samples": [{
                "source_id": product_source["id"], "product_id": "product",
                "price": "12000", "currency": "KRW",
            }],
        }
        samples_path.write_text(json.dumps(samples, indent=2), encoding="utf-8")

        receipt = cli(
            "verify", "--config", str(draft_path), "--receipt", str(receipt_path),
            "--samples", str(samples_path),
        )
        assert receipt["eligible"] is True
        assert receipt["reasons"] == []
        assert len(receipt["evidence_proofs"]) == 1

        activated = cli(
            "activate", "--config", str(draft_path), "--receipt", str(receipt_path),
            "--output", str(active_path),
        )
        assert activated["status"] == "activated"
        run = cli("run", "--config", str(active_path))

    assert run["status"] == "completed"
    assert Path(run["report_path"]).is_file()
    store = Store(tmp_path / "outputs")
    try:
        rows = store.observations(run["id"])
    finally:
        store.close()
    assert len(rows) == 1
    row = rows[0]
    assert row["source_id"] == product_source["id"]
    assert row["source_url"].endswith("/product/test-a")
    assert row["status"] == "verified"
    assert row["amount"] == "12000"
    assert row["normalized_amount"] == "12000"
    assert row["currency"] == "KRW"
    assert Path(row["evidence_path"]).is_file()
    workbook = load_workbook(run["report_path"], read_only=True)
    try:
        assert "Price Comparison" in workbook.sheetnames
        comparison_rows = list(workbook["Price Comparison"].values)
        assert len(comparison_rows) == 2
        assert any("12000" in str(cell) for cell in comparison_rows[1])
    finally:
        workbook.close()


def test_public_candidate_cannot_discover_loopback(tmp_path):
    workspace = tmp_path / "public-research.json"
    with research_site() as origin:
        cli(
            "init", "--workspace", str(workspace), "--industry", "Industrial pumps",
            "--product", "Process pump", "--market", "South Korea",
        )
        added = cli("source-add", "--workspace", str(workspace), "--url", origin + "/catalog")
        source_id = added["sources"][0]["id"]
        result = cli(
            "source-discover", "--workspace", str(workspace), "--source-id", source_id, "--limit", "10",
        )
    assert result["status"] == "policy_denied"
    assert result["added"] == []
    workspace_data = json.loads(workspace.read_text(encoding="utf-8"))
    assert [source["id"] for source in workspace_data["sources"]] == [source_id]
