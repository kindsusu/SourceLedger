"""A single approved inventory runs as resumable 50-page-or-smaller batches."""
from __future__ import annotations

from pathlib import Path
import json

import pytest

from su_crawler import research_batch as batch, research_plans as plans
from su_crawler.catalog_checkpoint import read_checkpoint
from su_crawler.models import FetchResult
from su_crawler.storage import Store


def _confirmed(root: Path, count: int = 45) -> dict:
    created = plans.create_plan(root, "Approved example product inventory")['plan']
    candidates = [{"url": f"https://shop.example/item/{i}", "name": f"Item {i}"} for i in range(count)]
    revised = plans.revise_plan(root, created['id'], created['revision'], {
        "topic": {"industry": "Retail", "product": "Example", "market": "Japan"},
        "added_candidates": candidates,
    })['plan']
    return plans.confirm_plan(root, revised['id'], revised['revision'], True)['plan']


def test_approved_inventory_exceeds_single_job_cap_and_is_pinned(tmp_path):
    plan = _confirmed(tmp_path, 525)
    args = {"plan_id": plan['id'], "expected_revision": plan['revision'], "batch_size": 20,
            "max_seconds": 600, "concurrency": 2}
    job = batch.queue_batch_run(tmp_path, args)
    assert job['arguments']['url_fingerprint'] == batch._url_fingerprint(plan['candidates'])
    assert batch.queue_batch_run(tmp_path, args)['id'] == job['id']
    assert len(plan['candidates']) == 525
    assert batch._initial_journal(job['id'], job['arguments'], plan['candidates'])['batch_count'] == 27
    changed = plans.revise_plan(tmp_path, plan['id'], plan['revision'], {'summary': 'New scope'})['plan']
    assert changed['revision'] > plan['revision']
    with pytest.raises(ValueError, match='changed'):
        batch.queue_batch_run(tmp_path, args)


def test_server_batches_resume_without_recollecting_finished_batches(tmp_path, monkeypatch):
    plan = _confirmed(tmp_path, 5)
    job = batch.queue_batch_run(tmp_path, {"plan_id": plan['id'], "expected_revision": plan['revision'],
                                          "batch_size": 2, "max_seconds": 600, "concurrency": 2})
    calls = []
    fail_once = True

    def collect(root, snapshot, **kwargs):
        nonlocal fail_once
        index = len(calls)
        calls.append((tuple(item['url'] for item in snapshot['candidates']), kwargs))
        if index == 1 and fail_once:
            fail_once = False
            raise RuntimeError('simulated interruption')
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = 'Results'
        workbook.active.append(['Product', 'Source', 'Recorded Price (not a final quote)', 'Currency', 'Result Status'])
        workbook.active.append([f'Item {index}', 'shop.example', 12, 'USD', 'review'])
        workbook.save(tmp_path / f'report-{index}.xlsx')
        return {"run_id": f"run-{index}", "status": "completed",
                "report_path": str(tmp_path / f"report-{index}.xlsx"),
                "budget": {"seconds_used": 1.5},
                "coverage": [{"url": item['url'], "status": "visited"} for item in snapshot['candidates']]}

    # The collector is the only external boundary. The test never opens the web.
    monkeypatch.setattr('su_crawler.catalog_collection.collect_plan', collect)
    monkeypatch.setattr(batch, '_observation_counts', lambda output, run_id: {
        'total': 2, 'priced': 2, 'review': 2, 'verified': 0})
    with pytest.raises(RuntimeError, match='simulated interruption'):
        batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    journal = batch.read_batch_journal(tmp_path, plan['id'], job['id'])
    assert len(journal['batches']) == 1
    result = batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert len(calls) == 4
    assert [len(call[0]) for call in calls] == [2, 2, 2, 1]
    assert calls[0][0] != calls[2][0]
    assert all(call[1]['follow_links'] is False and call[1]['concurrency'] == 2 for call in calls)
    assert result['result']['completed_batches'] == 3
    assert result['result']['target_count'] == 5
    assert result['result']['coverage_counts'] == {'visited': 5}
    assert result['result']['observation_counts']['total'] == 6
    assert result['result']['total_collection_seconds'] == 4.5
    assert len(batch.list_batch_reports(tmp_path, job['id'])['batches']) == 3
    assert batch.batch_progress(tmp_path, job['id'])['pending_targets'] == 0


def test_real_collection_shared_ledger_distinct_checkpoints_and_reports(tmp_path, monkeypatch):
    plan = _confirmed(tmp_path, 3)
    job = batch.queue_batch_run(tmp_path, {'plan_id': plan['id'], 'expected_revision': plan['revision'],
                                          'batch_size': 2, 'max_seconds': 120, 'concurrency': 2})
    calls = []

    def fetch(source, base_dir, backend):
        calls.append(source.location)
        payload = json.dumps({'@type': 'Product', 'name': source.location,
                              'offers': {'@type': 'Offer', 'price': '12', 'priceCurrency': 'USD'}}).encode()
        return FetchResult(source.id, 'fetched', backend,
                           content=b'<script type="application/ld+json">' + payload + b'</script>',
                           final_url=source.location)

    monkeypatch.setattr('su_crawler.catalog_collection.collect', fetch)
    result = batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert result['result']['completed_batches'] == 2
    assert result['result']['coverage_counts'] == {'visited': 3}
    reports = batch.list_batch_reports(tmp_path, job['id'])['batches']
    assert len({row['run_id'] for row in reports}) == 2
    assert len({row['checkpoint_id'] for row in reports}) == 2
    output = tmp_path / 'catalog-prices' / plan['id']
    store = Store(output)
    try:
        counts = [len(store.observations(row['run_id'])) for row in reports]
    finally:
        store.close()
    assert counts == [2, 1]
    assert result['result']['observation_counts']['total'] == 3
    from openpyxl import load_workbook
    final_path = Path(result['result']['final_report_path'])
    assert result['report_path'] == str(final_path)
    assert result['dashboard_path'] == result['result']['dashboard_path']
    assert final_path.is_file()
    assert Path(result['result']['dashboard_path']).is_file()
    with final_path.open('rb') as stream:
        workbook = load_workbook(stream, read_only=True)
        assert workbook['Results'].max_row == 4
        workbook.close()
    assert result['result']['coverage_summary']['unknown'] == 3
    assert result['result']['status'] == 'needs_review'
    assert all(Path(row['report_path']).is_file() for row in reports)
    assert all(read_checkpoint(output, row['checkpoint_id'])['result']['run_id'] == row['run_id'] for row in reports)
    calls_before_replay = len(calls)
    again = batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert again['result']['completed_batches'] == 2
    assert len(calls) == calls_before_replay
    assert again['result']['final_report_path'] == str(final_path)
    assert Path(again['result']['dashboard_path']).is_file()
    from su_crawler import final_report
    finalize = final_report.finalize_reports
    def fail_export(*args, **kwargs):
        raise OSError('synthetic final export failure')
    monkeypatch.setattr(final_report, 'finalize_reports', fail_export)
    with pytest.raises(OSError, match='synthetic final export failure'):
        batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert batch.read_batch_journal(tmp_path, plan['id'], job['id']).get('final_artifacts') is None
    assert len(calls) == calls_before_replay
    monkeypatch.setattr(final_report, 'finalize_reports', finalize)
    recovered = batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert Path(recovered['report_path']).is_file()
    assert len(calls) == calls_before_replay
    monkeypatch.setattr('su_crawler.catalog_collection.collection_rules_version', lambda: 'updated-rules')
    updated = batch.execute_batch_run(tmp_path, job['arguments'], job_id=job['id'])
    assert len(calls) == calls_before_replay
    assert updated['result']['observation_counts']['total'] == 3
    journal = batch.read_batch_journal(tmp_path, plan['id'], job['id'])
    assert len(journal['prior_results']) == 1
    old_runs = {row['run_id'] for row in journal['prior_results'][0]['batches']}
    new_runs = {row['run_id'] for row in journal['batches']}
    assert old_runs.isdisjoint(new_runs)
    store = Store(output)
    try:
        assert sum(len(store.observations(run_id)) for run_id in old_runs) == 3
        assert sum(len(store.observations(run_id)) for run_id in new_runs) == 3
    finally:
        store.close()


def test_partial_batch_status_is_preserved(tmp_path):
    journal = {'job_id': 'a' * 32, 'target_count': 2, 'batch_count': 2,
               'started_at': 'now', 'finished_at': 'later',
               'batches': [
                   {'target_count': 1, 'status': 'partial', 'coverage_counts': {'unprocessed': 1},
                    'observation_counts': {'total': 0}, 'budget_seconds_used': 5},
                   {'target_count': 1, 'status': 'completed', 'coverage_counts': {'visited': 1},
                    'observation_counts': {'total': 1}, 'budget_seconds_used': 2},
               ]}
    summary = batch._summary(journal, output=tmp_path)
    assert summary['status'] == 'partial'
    assert summary['coverage_counts'] == {'unprocessed': 1, 'visited': 1}
    assert summary['observation_counts']['total'] == 1


@pytest.mark.parametrize('override', [
    {'batch_size': 51}, {'max_seconds': 601}, {'concurrency': 5},
    {'batch_size': True}, {'follow_links': True},
])
def test_batch_budgets_and_expansion_are_bounded(tmp_path, override):
    plan = _confirmed(tmp_path)
    args = {'plan_id': plan['id'], 'expected_revision': plan['revision'], **override}
    with pytest.raises(ValueError):
        batch.queue_batch_run(tmp_path, args)
