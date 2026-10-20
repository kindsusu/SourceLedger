"""Offline integration checks for section coverage and its durable evidence."""
from pathlib import Path

from openpyxl import load_workbook
import pytest

from su_crawler.catalog_collection import collect_plan
from su_crawler.models import FetchResult
from su_crawler.storage import Store


FIXTURE = Path(__file__).parent / "fixtures" / "adapters" / "gongcar_rentpass.html"
URL = "https://gongcar.kr/rentpass"


def _plan():
    return {"id": "catalog", "state": "confirmed", "topic": {
        "industry": "Vehicle rental", "product": "Passenger vehicle", "market": "South Korea"},
        "include_terms": ["가상차량25 가상차량14"], "exclude_terms": [], "conditions": [], "excluded_urls": [],
        "candidates": [{"id": "rentpass", "url": URL, "selected": True}],
    }


def test_scope_exclusions_do_not_turn_discovered_sections_into_missing_cards(tmp_path, monkeypatch):
    payload = FIXTURE.read_bytes()
    seen = []

    def collector(source, _base, backend):
        seen.append((backend, source.recipe))
        if backend == "http":
            return FetchResult(source.id, "failed", backend, final_url=URL, message="offline fixture uses browser")
        return FetchResult(source.id, "fetched", backend, content=payload, final_url=URL,
                           media_type="text/html", trace=[
                               {"event": "list_expansion", "section": "popular", "status": "complete",
                                "reason": "button_disappeared", "observed_count": 3, "click_count": 2},
                               {"event": "list_expansion", "section": "promotion", "status": "complete",
                                "reason": "overflow_inspected", "observed_count": 1, "click_count": 1},
                           ])

    result = collect_plan(tmp_path, _plan(), output_dir=tmp_path / "out", max_pages=1,
                          follow_links=False, collector=collector, checkpoint_id="scope-filter")
    page = result["coverage"][0]
    assert page["products"] == 1
    assert page["scope_counts"]["excluded"] == 3
    sections = {section["section"]: section for section in page["sections"]}
    assert sections["rentpass_popular"]["observed_count"] == 3
    assert sections["rentpass_popular"]["extracted_count"] == 3
    assert sections["rentpass_popular"]["price_count"] == 3
    assert sections["rentpass_popular"]["unparsed_count"] == 0
    assert sections["rentpass_popular"]["status"] == "complete"
    assert sections["rentpass_popular"]["click_count"] == 2
    assert sections["rentpass_promotion"]["observed_count"] == 1
    assert sections["rentpass_promotion"]["status"] == "complete"
    # Gongcar sections are complete only from the rendered page: the browser
    # runs first and its usable capture makes the HTTP request unnecessary.
    assert [backend for backend, _ in seen] == ["playwright"]
    assert any(step["action"] == "expand_list" for step in seen[0][1])
    assert seen[0][1][0] == {"action": "click", "selector": '#promo.is-on button[data-act="promo-close"]',
                             "if_visible": True, "appearance_wait_seconds": 2, "timeout_seconds": 2}
    assert any(event.get("event") == "http_fallback_skipped" for event in page["trace"])

    store = Store(tmp_path / "out")
    try:
        assert store.run_coverage(result["run_id"]) == result["coverage"]
        assert store.observations(result["run_id"])
        from su_crawler.config import load_config
        from su_crawler.pipeline import report_for_run
        report_for_run(load_config(result['config_path']), store, result['run_id'])
    finally:
        store.close()
    assert sections["rentpass_popular"]["evidence_ref"]
    wb = load_workbook(result["report_path"], read_only=True)
    assert "Catalog coverage" in wb.sheetnames
    assert {row[1] for row in list(wb["Catalog coverage"].values)[1:]} == set(sections)

    # A changed browser recipe cannot accept a checkpoint made with the old
    # expansion contract and silently report its prior list as complete.
    from su_crawler import catalog_collection
    original_recipe = catalog_collection._ready_recipe
    monkeypatch.setattr(catalog_collection, "_ready_recipe",
                        lambda adapter, url: [*original_recipe(adapter, url), {"action": "wait_for", "selector": "body"}])
    with pytest.raises(ValueError, match="checkpoint scope or budget changed"):
        collect_plan(tmp_path, _plan(), output_dir=tmp_path / "out", max_pages=1,
                     follow_links=False, collector=collector, checkpoint_id="scope-filter")


@pytest.mark.parametrize("browser_status", ["timeout", "tool_unavailable", "failed"])
def test_browser_failure_falls_back_to_http(tmp_path, browser_status):
    payload = FIXTURE.read_bytes()
    seen = []

    def collector(source, _base, backend):
        seen.append(backend)
        if backend == "playwright":
            return FetchResult(source.id, browser_status, backend, message="browser did not finish")
        return FetchResult(source.id, "fetched", backend, content=payload, final_url=URL, media_type="text/html")

    plan = {**_plan(), "include_terms": []}
    result = collect_plan(tmp_path, plan, output_dir=tmp_path / "out", max_pages=1,
                          follow_links=False, collector=collector, checkpoint_id="fallback")
    assert seen == ["playwright", "http"]
    page = result["coverage"][0]
    assert page["status"] == "visited" and page["products"] > 0
    assert {event.get("reason") for event in page["trace"] if event.get("event") == "http_fallback"} == {browser_status}
    # Static HTML never completes a Gongcar section.
    assert all(section["status"] != "complete" for section in page["sections"])
    store = Store(tmp_path / "out")
    try:
        assert store.observations(result["run_id"])
    finally:
        store.close()


def test_browser_capture_without_rows_still_tries_http(tmp_path):
    payload = FIXTURE.read_bytes()
    seen = []

    def collector(source, _base, backend):
        seen.append(backend)
        content = b"<html><body></body></html>" if backend == "playwright" else payload
        return FetchResult(source.id, "fetched", backend, content=content, final_url=URL, media_type="text/html")

    plan = {**_plan(), "include_terms": []}
    result = collect_plan(tmp_path, plan, output_dir=tmp_path / "out", max_pages=1,
                          follow_links=False, collector=collector, checkpoint_id="empty-browser")
    assert seen == ["playwright", "http"]
    assert result["coverage"][0]["products"] > 0
