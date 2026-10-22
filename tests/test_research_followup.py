"""Offline recovery and evidence tests; external models/services are never used."""
import asyncio
import json
from pathlib import Path

import pytest

from su_crawler import assistant_runtime as runtime, catalog_collection, research_plans as plans
from su_crawler.assistant_server import build_assistant_server
from su_crawler.assistant_workspace import execute_job
from su_crawler.catalog_checkpoint import read_checkpoint
from su_crawler.models import FetchResult, utc_now
from su_crawler.plan_jobs import queue_plan_job
from su_crawler.research_followup import research_gaps, submit_browser_evidence, validate_retry
from su_crawler.storage import Store
from test_catalog_collection import plan, html_product


def completed_job(root, monkeypatch, urls):
    p = plans.create_plan(root, "Inspect selected fixture products in Korea")["plan"]
    p = plans.revise_plan(root, p["id"], p["revision"], {
        "topic": {"industry": "retail", "product": "fixtures", "market": "Korea"},
        "added_candidates": [{"url": url, "name": "Fixture"} for url in urls]})["plan"]
    p = plans.confirm_plan(root, p["id"], p["revision"], True)["plan"]
    job = queue_plan_job(root, "research_plan", {"plan_id": p["id"], "expected_revision": p["revision"], "max_pages": len(urls)})
    result = execute_job(root, "research_plan", job["arguments"], job_id=job["id"])
    job["status"] = "succeeded"
    runtime._write_json(runtime._job_path(root, job["id"]), job)
    runtime._write_json(runtime._job_path(root, job["id"]).with_name("result.json"), result)
    return job, result


def test_job_result_edit_cannot_hide_persisted_unknown_enumeration(tmp_path, monkeypatch):
    url = "https://shop.example/item"
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(
        s.id, "fetched", backend, content=html_product("Widget", price="10", currency="USD"), final_url=s.location))
    job, result = completed_job(tmp_path, monkeypatch, [url])
    page = result["result"]["coverage"][0]
    assert page["status"] == "visited"
    page["sections"] = [{"section": "popular", "status": "partial", "observed_count": 4,
                         "extracted_count": 4, "price_count": 4, "unparsed_count": 0,
                         "reason": "button_stalled", "click_count": 1}]
    runtime._write_json(runtime._job_path(tmp_path, job["id"]).with_name("result.json"), result)
    gap = research_gaps(tmp_path, job["id"])["items"][0]
    assert gap["status"] == "partial" and gap["retryable"] is False
    assert any(issue['category'] == 'unknown_enumeration' for issue in gap['collection_issues'])
    assert gap["sections"] == page["sections"]
    args = {"retry_of": job["id"], "retry_urls": [url], **{
        key: job["arguments"][key] for key in ("plan_id", "expected_revision", "plan_fingerprint")}}
    with pytest.raises(ValueError, match='review-only prices need evidence'):
        validate_retry(tmp_path, args)
    page["sections"][0]["status"] = "complete"
    runtime._write_json(runtime._job_path(tmp_path, job["id"]).with_name("result.json"), result)
    gap = research_gaps(tmp_path, job["id"])["items"][0]
    assert gap["status"] == "needs_review" and gap["retryable"] is False
    with pytest.raises(ValueError, match="review-only prices"):
        validate_retry(tmp_path, args)


def test_browser_submission_mcp_is_review_idempotent_and_in_xlsx(tmp_path, monkeypatch):
    url = "https://shop.example/item"
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "fetched", backend, content=b"<html>No metadata</html>", final_url=s.location))
    job, result = completed_job(tmp_path, monkeypatch, [url])
    assert research_gaps(tmp_path, job["id"])["items"][0]["status"] == "no_data"
    server = build_assistant_server(tmp_path)
    payload = dict(job_id=job["id"], url=url, captured_at=utc_now(), page_text="Widget option A 12.30 USD", fields={"name": "Widget", "price": "12.30", "currency": "USD", "option": "option A"}, locator="price card A")
    first = asyncio.run(server.call_tool("submit_browser_evidence", payload))[1]
    second = asyncio.run(server.call_tool("submit_browser_evidence", payload))[1]
    assert first["submission_id"] == second["submission_id"]
    assert first["observation"]["status"] == "review" and not first["observation"]["comparable"]
    assert first["observation"]["amount"] == "12.30"
    presentation = asyncio.run(server.call_tool("get_job_observations", {"job_id": job["id"], "detail": True}))[1]
    assert presentation["evidence_count"] == 1 and presentation["unlinked_evidence_count"] == 1
    assert presentation["result_count"] == 1
    assert presentation["rows"][0]["id"] == first["submission_id"]
    assert presentation["rows"][0]["primary_candidate"] is True
    assert presentation["rows"][0]["status"] == "review"
    from openpyxl import load_workbook
    wb = load_workbook(result["report_path"], read_only=True, data_only=True)
    try:
        values = [str(v) for ws in wb for row in ws.values for v in row if v is not None]
        assert "12.30" in values and first["submission_id"] in values
    finally:
        wb.close()
    with pytest.raises(ValueError, match="exact support"):
        submit_browser_evidence(tmp_path, **{**payload, "fields": {"name": "Widget", "price": "999"}})
    with pytest.raises(ValueError, match="approved page"):
        submit_browser_evidence(tmp_path, **{**payload, "url": "https://other.example/item"})
    with pytest.raises(ValueError, match="timezone"):
        submit_browser_evidence(tmp_path, **{**payload, "captured_at": "2026-01-01T00:00:00"})
    missing = submit_browser_evidence(tmp_path, **{**payload, "page_text": "Widget price on request", "fields": {"name": "Widget"}})
    assert missing["observation"]["amount"] is None and missing["observation"]["currency"] is None


def test_browser_rental_starting_price_is_review_candidate_with_literal_proof(tmp_path, monkeypatch):
    url = "https://gongcar.kr/rentpass"
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "fetched", backend, content=b"<html>No metadata</html>", final_url=s.location))
    job, result = completed_job(tmp_path, monkeypatch, [url])
    page_text = "더 뉴 모닝 월 290,000원~ 48개월 / 보증금 30%"
    response = submit_browser_evidence(
        tmp_path, job["id"], url=url, captured_at=utc_now(), page_text=page_text,
        fields={"name": "더 뉴 모닝", "price": "월 290,000원~", "contract_term": "48개월", "deposit": "보증금 30%"},
        locator="더 뉴 모닝 card",
    )
    raw = response["observation"]
    assert raw["raw_fields"]["price"] == "월 290,000원~"
    assert raw["evidence"]["price"]["raw"] == "월 290,000원~"
    assert raw["amount"] == "290000" and raw["derived_values"]["price_kind"] == "from"
    assert raw["derived_values"]["literal_interpretations"]["term_months"] == "48"
    assert raw["derived_values"]["literal_interpretations"]["deposit_percent"] == "30"
    assert raw["derived_values"]["currency_from_price_marker"] == "KRW"
    assert raw["currency"] is None
    assert "deposit_amount" not in raw["derived_values"]["literal_interpretations"]
    assert raw["price_profile"] == "rental" and raw["status"] == "review" and not raw["comparable"]
    from su_crawler.assistant_workspace import result_observations
    view = result_observations(tmp_path, runtime.get_job(tmp_path, job["id"]))
    assert view["result_count"] == 1 and view["rows"][0]["primary_candidate"]
    assert view["rows"][0]["field_status"]["term_months"]["value"] == "48"
    assert view["rows"][0]["field_status"]["term_months"]["status"] == "derived"
    assert view["rows"][0]["field_status"]["deposit_percent"]["value"] == "30"
    assert "deposit_amount" in view["rows"][0]["missing_conditions"]
    from openpyxl import load_workbook
    wb = load_workbook(result["report_path"], read_only=True, data_only=True)
    try:
        result_rows = list(wb["Results"].values)
        assert len(result_rows) == 2
        assert result_rows[1][2] == 290000
        assert "starting at" in result_rows[1][5]
        assert "월 290,000원~" in result_rows[1][5]
    finally:
        wb.close()


def test_combined_contract_text_and_repeated_capture_make_one_review_row(tmp_path, monkeypatch):
    url = "https://gongcar.kr/rentpass"
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "fetched", backend, content=b"<html>No metadata</html>", final_url=s.location))
    job, _ = completed_job(tmp_path, monkeypatch, [url])
    fields = {"name": "더 뉴 레이", "price": "월 310,000원~", "terms": "48개월 / 보증금 30% 기준"}
    args = dict(url=url, page_text="더 뉴 레이 월 310,000원~ 48개월 / 보증금 30% 기준", fields=fields, locator="레이 card")
    first = submit_browser_evidence(tmp_path, job["id"], captured_at=utc_now(), **args)
    second = submit_browser_evidence(tmp_path, job["id"], captured_at=utc_now(), **args)
    assert first["submission_id"] != second["submission_id"]
    parsed = first["observation"]["derived_values"]["literal_interpretations"]
    assert parsed["term_months"] == "48"
    assert parsed["deposit_percent"] == "30"
    assert parsed["price_basis"] == "monthly"
    from su_crawler.assistant_workspace import result_observations
    view = result_observations(tmp_path, runtime.get_job(tmp_path, job["id"]))
    assert view["result_count"] == 1 and view["evidence_count"] == 2
    assert view["linked_evidence_count"] == 1 and view["unlinked_evidence_count"] == 1
    assert view["rows"][0]["price_profile"] == "rental"
    assert view["rows"][0]["field_status"]["price_basis"]["status"] == "derived"


def test_generic_terms_and_insurance_do_not_turn_food_into_rental(tmp_path, monkeypatch):
    url = "https://shop.example/food"
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "fetched", backend, content=b"<html>No metadata</html>", final_url=s.location))
    job, _ = completed_job(tmp_path, monkeypatch, [url])
    result = submit_browser_evidence(
        tmp_path, job["id"], url=url, captured_at=utc_now(),
        page_text="Apple 20,000원 delivery terms shipping insurance",
        fields={"name": "Apple", "price": "20,000원", "terms": "delivery terms", "insurance": "shipping insurance"},
        locator="Apple card",
    )
    assert result["observation"]["price_profile"] == "unit"
    assert "term_months" not in result["observation"]["derived_values"]["literal_interpretations"]


def test_resume_keeps_completed_pages_and_charges_interrupted_page(tmp_path):
    urls = ["https://shop.example/a", "https://shop.example/b", "https://shop.example/c"]
    calls = []
    def fetch(source, base, backend):
        calls.append(source.location)
        if source.location == urls[1]:
            raise KeyboardInterrupt("simulate hard process death")
        return FetchResult(source.id, "fetched", backend, content=html_product(source.location, price="10", currency="USD"), final_url=source.location)
    output = tmp_path / "out"
    with pytest.raises(KeyboardInterrupt):
        catalog_collection.collect_plan(tmp_path, plan(urls), output_dir=output, max_pages=3, max_seconds=180, collector=fetch, checkpoint_id="run")
    assert read_checkpoint(output, "run")["dispatched"] == 2
    calls.clear()
    result = catalog_collection.collect_plan(tmp_path, plan(urls), output_dir=output, max_pages=3, max_seconds=180, collector=fetch, checkpoint_id="run")
    assert set(calls) == {urls[2]}
    assert result["budget"]["pages_used"] == 3 and result["budget"]["seconds_used"] >= 60
    assert result["status"] == "partial"
    assert any(p["status"] == "interrupted" for p in result["coverage"])
    calls.clear()
    again = catalog_collection.collect_plan(tmp_path, plan(urls), output_dir=output, max_pages=3, max_seconds=180, collector=fetch, checkpoint_id="run")
    assert again == result and not calls
    with pytest.raises(ValueError, match="budget changed"):
        catalog_collection.collect_plan(tmp_path, plan(urls), output_dir=output, max_pages=4, max_seconds=180, collector=fetch, checkpoint_id="run")


def test_retry_only_failed_page_reuses_success_and_rejects_new_url(tmp_path, monkeypatch):
    urls = ["https://shop.example/a", "https://shop.example/b"]
    calls = []
    def fetch(s, b, backend):
        calls.append(s.location)
        if s.location == urls[1] and calls.count(urls[1]) <= 3:
            return FetchResult(s.id, "timeout", backend, message="temporary failure", final_url=s.location)
        return FetchResult(s.id, "fetched", backend, content=html_product("Widget", price="10", currency="USD"), final_url=s.location)
    monkeypatch.setattr(catalog_collection, "collect", fetch)
    # A timed-out page has no robots allow on record; the guard asks again (offline here).
    from su_crawler import host_capture_policy
    monkeypatch.setattr(host_capture_policy, "_live_robots_check", lambda host: lambda url: ("allowed", "test"))
    job, result = completed_job(tmp_path, monkeypatch, urls)
    evidence = submit_browser_evidence(tmp_path, job["id"], url=urls[1], captured_at=utc_now(),
                                      page_text="Widget unavailable", fields={"name": "Widget"}, locator="card")
    args = {"plan_id": job["arguments"]["plan_id"], "expected_revision": job["arguments"]["expected_revision"], "retry_of": job["id"], "retry_urls": [urls[1]], "max_pages": 1}
    with pytest.raises(ValueError, match="unresolved pages"):
        queue_plan_job(tmp_path, "research_plan", {**args, "retry_urls": ["https://shop.example/new"]})
    retry = queue_plan_job(tmp_path, "research_plan", args)
    calls.clear()
    combined = execute_job(tmp_path, "research_plan", retry["arguments"], job_id=retry["id"])
    assert set(calls) == {urls[1]}
    assert {p["url"] for p in combined["result"]["coverage"]} == set(urls)
    store = Store(Path(combined["output_dir"]))
    try:
        rows = store.observations(combined["run_id"])
        assert sum(r["extraction_method"] == "json_ld_product_offer" for r in rows) == 1
        supplements = [r for r in rows if r["extraction_method"] == "host_browser_submission"]
        assert len(supplements) == 1 and supplements[0]["status"] == "review"
        assert supplements[0]["derived_values"]["parent_submission_id"] == evidence["submission_id"]
    finally:
        store.close()


def test_plan_resume_mcp_starts_worker_but_does_not_reset_budget(tmp_path, monkeypatch):
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "failed", backend))
    job, result = completed_job(tmp_path, monkeypatch, ["https://shop.example/item"])
    job["status"] = "interrupted"
    runtime._write_json(runtime._job_path(tmp_path, job["id"]), job)
    launches = []
    monkeypatch.setattr(runtime, "start_worker", lambda root: launches.append(root))
    server = build_assistant_server(tmp_path)
    queued = asyncio.run(server.call_tool("resume_job_execution", {"job_id": job["id"]}))[1]
    assert queued["status"] == "queued" and launches
    assert queued["arguments"] == job["arguments"]
    resumed = execute_job(tmp_path, "research_plan", queued["arguments"], job_id=job["id"])
    assert resumed["run_id"] == result["run_id"]


def test_jetcar_discovery_does_not_report_navigation_as_missing_products():
    from su_crawler.catalog_collection import _links
    url = "https://www.jetcar.kr/sub0201/2159"
    result = FetchResult("s", "fetched", "http", content=b'<a href="/content/provision">Terms</a><a href="/sub0201/2160">Car</a><a href="/bbs/login.php">Login</a>')
    assert _links(result, url, "www.jetcar.kr", set()) == ["https://www.jetcar.kr/sub0201/2160"]


def test_restart_after_ledger_commit_reuses_sealed_run(tmp_path, monkeypatch):
    url = "https://shop.example/a"
    calls = []
    def fetch(s, b, backend):
        calls.append(s.location)
        return FetchResult(s.id, "fetched", backend, content=html_product("Widget", price="10", currency="USD", links='<a href="/b">B</a>'), final_url=s.location)
    execute = catalog_collection.execute
    def interrupted(*args, **kwargs):
        execute(*args, **kwargs)
        raise KeyboardInterrupt("report receipt not yet committed")
    monkeypatch.setattr(catalog_collection, "execute", interrupted)
    kwargs = dict(output_dir=tmp_path / "out", max_pages=1, collector=fetch, checkpoint_id="receipt")
    with pytest.raises(KeyboardInterrupt):
        catalog_collection.collect_plan(tmp_path, plan([url]), **kwargs)
    assert read_checkpoint(tmp_path / "out", "receipt")["collection_finished"]
    monkeypatch.setattr(catalog_collection, "execute", execute)
    calls.clear()
    result = catalog_collection.collect_plan(tmp_path, plan([url]), **kwargs)
    assert not calls and len(result["coverage"]) == 2
    store = Store(tmp_path / "out")
    try:
        assert store.db.execute("SELECT count(*) FROM runs").fetchone()[0] == 1
        assert len(store.observations(result["run_id"])) == 1
    finally:
        store.close()


def _blocked_job(tmp_path, monkeypatch, url="https://shop.example.test/item"):
    monkeypatch.setattr(catalog_collection, "collect", lambda s, b, backend: FetchResult(s.id, "fetched", backend, content=b"<html>No metadata</html>", final_url=s.location))
    return completed_job(tmp_path, monkeypatch, [url])


def _evidence_args(url):
    return dict(url=url, captured_at="2026-10-01T00:00:00+00:00", page_text="Widget 12.30 USD",
                fields={"name": "Widget", "price": "12.30"}, locator="card")


def _submission_state(result):
    output = Path(result["output_dir"])
    store = Store(output)
    try:
        count = store.db.execute("SELECT count(*) FROM browser_submissions").fetchone()[0]
    finally:
        store.close()
    folder = output / "browser-evidence"
    return count, sorted(p.name for p in folder.iterdir()) if folder.exists() else []


def test_final_url_is_optional_and_policy_record_is_stored_on_the_row(tmp_path, monkeypatch):
    url = "https://shop.example.test/item"
    job, result = _blocked_job(tmp_path, monkeypatch, url)
    args = _evidence_args(url)
    plain = submit_browser_evidence(tmp_path, job["id"], **args)
    record = plain["observation"]["host_capture_policy"]
    assert record["redirected"] is False and record["robots"] == "not_needed" and record["final_url"] == url
    assert record["final_url_reported"] is False
    assert any("did not report the final URL" in flag for flag in plain["observation"]["review_flags"])
    assert plain["observation"]["status"] == "review" and not plain["observation"]["comparable"]
    store = Store(Path(result["output_dir"]))
    try:
        stored = [row for row in store.observations(result["run_id"]) if row["id"] == plain["submission_id"]][0]
    finally:
        store.close()
    assert stored["host_capture_policy"] == record
    saved = json.loads(Path(stored["evidence_path"]).read_text(encoding="utf-8"))
    assert saved["final_url"] == url and saved["capture_policy"] == record
    reported = submit_browser_evidence(tmp_path, job["id"], final_url=url, **args)
    assert reported["submission_id"] != plain["submission_id"]
    assert reported["observation"]["host_capture_policy"]["final_url_reported"] is True
    assert not any("did not report the final URL" in flag for flag in reported["observation"]["review_flags"])
    same = submit_browser_evidence(tmp_path, job["id"], final_url=url + "#frag", **args)
    assert same["submission_id"] == reported["submission_id"]
    view = asyncio.run(build_assistant_server(tmp_path).call_tool("get_job_observations", {"job_id": job["id"]}))[1]
    assert view["evidence_count"] == 2


def test_redirected_final_url_needs_same_host_and_allowed_robots(tmp_path, monkeypatch):
    from su_crawler import host_capture_policy as policy
    url = "https://shop.example.test/item"
    job, result = _blocked_job(tmp_path, monkeypatch, url)
    decisions = iter([("disallowed", "robots.txt disallows collection"), ("unavailable", "Unable to check robots.txt"), ("allowed", "ok")])
    monkeypatch.setattr(policy, "_live_robots_check", lambda host: (lambda final: next(decisions)))
    final = "https://www.shop.example.test/item/2"
    with pytest.raises(ValueError, match="approved page's host"):
        submit_browser_evidence(tmp_path, job["id"], final_url="https://other.example.test/item", **_evidence_args(url))
    with pytest.raises(ValueError, match="disallowed by robots"):
        submit_browser_evidence(tmp_path, job["id"], final_url=final, **_evidence_args(url))
    with pytest.raises(ValueError, match="could not be verified"):
        submit_browser_evidence(tmp_path, job["id"], final_url=final, **_evidence_args(url))
    assert _submission_state(result) == (0, [])
    ok = submit_browser_evidence(tmp_path, job["id"], final_url=final, **_evidence_args(url))
    assert ok["observation"]["host_capture_policy"]["redirected"] is True
    assert ok["observation"]["host_capture_policy"]["final_url"] == final
    assert ok["observation"]["status"] == "review" and not ok["observation"]["comparable"]


def test_policy_denied_page_rejects_host_evidence_and_writes_nothing(tmp_path, monkeypatch):
    url = "https://shop.example.test/item"
    job, result = _blocked_job(tmp_path, monkeypatch, url)
    result["result"]["coverage"][0]["status"] = "policy_denied"
    runtime._write_json(runtime._job_path(tmp_path, job["id"]).with_name("result.json"), result)
    with pytest.raises(ValueError, match="denied by collection policy"):
        submit_browser_evidence(tmp_path, job["id"], **_evidence_args(url))
    with pytest.raises(Exception, match="denied by collection policy"):
        asyncio.run(build_assistant_server(tmp_path).call_tool("submit_browser_evidence", {"job_id": job["id"], **_evidence_args(url)}))
    assert _submission_state(result) == (0, [])
    (tmp_path / "shot.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(ValueError, match="denied by collection policy"):
        submit_browser_evidence(tmp_path, job["id"], screenshot_path="shot.png", **_evidence_args(url))
    assert _submission_state(result) == (0, [])


def test_final_url_with_credentials_or_private_address_is_rejected(tmp_path, monkeypatch):
    url = "https://shop.example.test/item"
    job, result = _blocked_job(tmp_path, monkeypatch, url)
    for final in ("https://user:pw@shop.example.test/item", "ftp://shop.example.test/item", "https://127.0.0.1/item"):
        with pytest.raises(ValueError):
            submit_browser_evidence(tmp_path, job["id"], final_url=final, **_evidence_args(url))
    assert _submission_state(result) == (0, [])
