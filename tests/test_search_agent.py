"""A mocked public-service chain through search, proposal, verification, and reporting."""
from __future__ import annotations

import json
import socket

import httpx
from openpyxl import load_workbook

from su_crawler import agent, research, search
from su_crawler.models import FetchResult
from su_crawler.storage import Store


_HTML = b"""<!doctype html><article class="product-row">
<span data-field="model">PX-100</span><span data-field="price">12000</span>
<span data-field="currency">KRW</span><span data-field="unit">each</span>
<span data-field="pack_quantity">1</span><span data-field="tax">included</span>
<span data-field="price_type">retail</span><span data-field="price_basis">each</span>
</article>"""


def test_agent_search_to_verified_shared_ledger_and_xlsx(tmp_path, monkeypatch):
    workspace = tmp_path / "research.json"
    research.init_workspace(workspace, industry="Industrial pumps", product="Process pump", market="South Korea")
    research.set_product(workspace, identifiers={"model": "PX-100"})
    provider = tmp_path / "searxng.json"
    provider.write_text(json.dumps({
        "provider": "searxng", "endpoint": "https://search.example/search",
        "allow_private_endpoint": False, "timeout_seconds": 15, "language": "en",
        "allowed_result_domains": ["vendor.example"],
    }), encoding="utf-8")

    queries: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "search.example":
            queries.append(request.url.params["q"])
            return httpx.Response(200, json={"results": [
                {"url": "https://unrelated.example/item", "title": "Ignore"},
                {"url": "https://vendor.example/product#listing", "title": "Vendor PX-100", "content": "not price evidence"},
            ]})
        if request.url.host == "vendor.example" and request.url.path == "/robots.txt":
            return httpx.Response(200, content=b"User-agent: *\nAllow: /\n", headers={"content-type": "text/plain"})
        if request.url.host == "vendor.example" and request.url.path == "/product":
            return httpx.Response(200, content=_HTML, headers={"content-type": "text/html; charset=utf-8"})
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    real_client = httpx.Client
    monkeypatch.setattr(httpx, "Client", lambda **kwargs: real_client(transport=transport, **kwargs))
    monkeypatch.setattr(socket, "getaddrinfo", lambda host, port, **kwargs: [
        (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", port or 443)),
    ])

    def fixture_collect(source, base_dir, backend, **kwargs):
        assert source.location == "https://vendor.example/product"
        return FetchResult(
            source.id, "fetched", backend, content=_HTML, media_type="text/html; charset=utf-8",
            final_url=source.location,
        )

    # Static HTML deliberately cannot establish display evidence. Supply a
    # synthetic rendered capture for proposal and final collection so this
    # test never starts a real browser or contacts vendor.example.
    monkeypatch.setattr("su_crawler.proposals.collect", fixture_collect)
    monkeypatch.setattr("su_crawler.collectors.collect", fixture_collect)

    state = agent.run_agent(workspace, run_dir=tmp_path / "run", search_config_path=provider,
                            max_sources=3, max_model_calls=0, max_seconds=30)

    assert state["status"] == "completed"
    assert state["usage"]["model_calls"] == 0
    assert state["search"]["status"] == "searched"
    assert queries == ["Industrial pumps Process pump South Korea PX-100"]
    saved = research.load_workspace(workspace)["sources"]
    assert len(saved) == 1
    assert saved[0]["location"] == "https://vendor.example/product"
    assert saved[0]["provenance"]["method"] == "search"
    assert saved[0]["provenance"]["provider"] == "searxng"
    assert state["tasks"][0]["status"] == "proposed"
    assert state["collection"]["eligible"] is True
    assert "activation" not in state["collection"]  # Activation is intentionally not requested.

    store = Store(tmp_path / "outputs")
    try:
        observations = store.observations(state["collection"]["run_id"])
    finally:
        store.close()
    assert len(observations) == 1
    assert observations[0]["status"] == "verified"
    assert observations[0]["amount"] == "12000"
    report = state["collection"]["report_path"]
    book = load_workbook(report, read_only=True)
    try:
        assert "Price Comparison" in book.sheetnames
    finally:
        book.close()
