from dataclasses import replace
from pathlib import Path
import json
import pytest

from su_crawler.config import load_config
from su_crawler.models import Candidate, FetchResult, Product
from su_crawler.pipeline import _candidate_quality, _conflicts, _has_sufficient_evidence, execute, report_for_run
from su_crawler.storage import Store, workspace_lock, RunBusyError

ROOT = Path(__file__).resolve().parents[1]


def demo(tmp_path):
    return replace(load_config(ROOT / "examples/demo.json"), output_dir=str(tmp_path))


def test_deferred_export_keeps_evidence_and_resumes_without_recollection(tmp_path):
    config = demo(tmp_path)
    result = execute(config, export=False)
    assert not result.get("report_path")
    assert not list(tmp_path.glob("*.xlsx"))
    store = Store(tmp_path)
    before = store.observations(result["id"])
    assert before and all(Path(row["evidence_path"]).is_file() for row in before)
    assert store.source_tasks(result["id"], "catalog") == [
        row for row in store.tasks(result["id"]) if row["source_id"] == "catalog"]
    store.close()
    result = execute(config, resume_id=result["id"], collector=lambda *a: pytest.fail("recollected"))
    assert Path(result["report_path"]).is_file()
    store = Store(tmp_path)
    assert store.observations(result["id"]) == before
    store.close()


def test_end_to_end_and_resume_does_not_duplicate(tmp_path):
    config = demo(tmp_path)
    result = execute(config, max_tasks=1)
    assert result["status"] == "paused"
    store = Store(tmp_path)
    assert len(store.tasks(result["id"])) == 5
    before = store.observations(result["id"])
    assert len(before) == 1
    store.close()
    final = execute(config, resume_id=result["id"])
    assert final["status"] in {"completed", "partial"}
    assert Path(final["report_path"]).is_file()
    dashboard = Path(final["dashboard_path"])
    assert dashboard.is_file()
    dashboard.unlink()
    restored = execute(config, resume_id=result["id"], collector=lambda *a: pytest.fail("completed run recollected"))
    assert Path(restored["dashboard_path"]).is_file()
    execute(config, resume_id=result["id"], collector=lambda *a: pytest.fail("completed run recollected"))
    store = Store(tmp_path)
    rows = store.observations(result["id"])
    assert len(rows) == 4
    assert len({r["id"] for r in rows}) == 4
    assert sum(r["status"] == "verified" for r in rows) == 1
    catalog = next(r for r in rows if r["source_id"] == "catalog" and r["product_id"] == "demo_a")
    assert catalog["evidence_mode"] == "static_html"
    assert catalog["source_visibility"] == "unconfirmed"
    assert catalog["status"] == "review" and catalog["comparable"] is False
    assert all(r["amount"] is None for r in rows if r["product_id"] == "demo_b")
    assert all(Path(r["evidence_path"]).is_file() for r in rows)
    store.close()


def test_partial_source_failure_does_not_abort(tmp_path):
    config = demo(tmp_path)
    from su_crawler.collectors import collect
    def collector(source, base, backend):
        return FetchResult(source.id, "needs_auth", backend, message="인증 필요") if source.id == "catalog" else collect(source, base, backend)
    result = execute(config, collector=collector)
    store = Store(tmp_path)
    assert result["status"] == "partial"
    assert len(store.tasks(result["id"])) == 5
    assert len(store.observations(result["id"])) == 2
    assert Path(result["report_path"]).is_file()
    store.close()


def test_failed_browser_attempt_persists_prior_capture_without_observations(tmp_path):
    config = demo(tmp_path)
    source = next(s for s in config.sources if s.id == "catalog")
    config = replace(config, sources=[replace(source, kind="web", location="https://example.org/item",
                                              allowed_domains=["example.org"], backends=["playwright"],
                                              recipe=[{"action": "enumerate_options", "selectors": ["#size"]}],
                                      max_attempts=1)])
    content = b"<html><body><strong>Observed source text</strong></body></html>"

    def collector(item, base, backend):
        return FetchResult(item.id, "timeout", backend, message="Later option timed out",
                           variant_snapshots=[{"content": content, "final_url": item.location,
                                               "fetched_at": "2026-01-01T00:00:00+00:00",
                                               "selections": {"#size": "Large"}}])

    run = execute(config, collector=collector)
    store = Store(tmp_path)
    try:
        attempts = [dict(row) for row in store.db.execute("SELECT * FROM attempts WHERE run_id=?", (run["id"],))]
        assert len(attempts) == 1 and attempts[0]["status"] == "timeout"
        refs = json.loads(attempts[0]["data"])["prior_capture_evidence"]
        assert len(refs) == 1
        assert Path(refs[0]["content_path"]).read_bytes() == content
        assert json.loads(Path(refs[0]["receipt_path"]).read_text(encoding="utf-8"))["hash"] == refs[0]["content_sha256"]
        assert store.observations(run["id"]) == []
        assert all(task["status"] == "timeout" for task in store.tasks(run["id"]))
    finally:
        store.close()


def test_report_reexport_uses_live_issue_status(tmp_path):
    from openpyxl import load_workbook
    from su_crawler.collection_issues import sync_coverage_issues

    config = demo(tmp_path)
    run = execute(config)
    store = Store(tmp_path)
    try:
        coverage = [{"url": "https://example.org/item", "source_id": "catalog", "status": "visited",
                     "sections": [{"section": "offers", "status": "unknown"}]}]
        store.save_run_coverage(run["id"], coverage)
        issue = sync_coverage_issues(store, run["id"], coverage)[0]
        store.update_collection_issue(issue["id"], status="needs_review")
        report_for_run(config, store, run["id"])
        assert store.run_coverage(run["id"]) == coverage
        workbook = load_workbook(run["report_path"], read_only=True, data_only=True)
        try:
            assert workbook["Collection Issues"]["D2"].value == "needs_review"
        finally:
            workbook.close()
    finally:
        store.close()


def test_changed_config_resume_and_lock(tmp_path):
    config = demo(tmp_path)
    run = execute(config, max_tasks=1)
    with pytest.raises(ValueError):
        execute(replace(config, name="changed"), resume_id=run["id"])
    with workspace_lock(tmp_path):
        with pytest.raises(RunBusyError):
            execute(config)


def test_stale_reexport_preserves_observation(tmp_path):
    config = demo(tmp_path)
    run = execute(config)
    store = Store(tmp_path)
    row = next(r for r in store.observations(run["id"]) if r["status"] == "verified")
    row["collected_at"] = "2000-01-01T00:00:00+00:00"
    with store.db:
        store.db.execute("UPDATE observations SET data=? WHERE id=?", (json.dumps(row), row["id"]))
    report_for_run(config, store, run["id"])
    assert next(r for r in store.observations(run["id"]) if r["id"] == row["id"])["amount"] == row["amount"]
    from openpyxl import load_workbook
    wb = load_workbook(run["report_path"], read_only=True)
    assert any("stale" in str(r) for r in wb["Observation History"].values)
    wb.close()
    store.close()


def test_http_failure_browser_fallback(tmp_path):
    config = demo(tmp_path)
    source = replace(config.sources[0], kind="web", location="https://example.com/item", backends=["http", "playwright"], min_interval_seconds=0)
    config = replace(config, sources=[source])
    calls = []
    def collector(source, base, backend):
        calls.append(backend)
        if backend == "http":
            return FetchResult(source.id, "blocked", backend, message="403")
        return FetchResult(source.id, "fetched", backend, content=(ROOT / "examples/fixtures/catalog.html").read_bytes(), final_url=source.location)
    result = execute(config, collector=collector)
    assert calls[:2] == ["http", "playwright"]
    store = Store(tmp_path)
    assert any(r["amount"] == "12000" for r in store.observations(result["id"]))
    store.close()


def test_browser_shell_does_not_discard_partial_http_evidence(tmp_path):
    config = demo(tmp_path)
    source = replace(config.sources[0], kind="web", location="https://example.com/item", backends=["http", "playwright"],
                     row_selector=".item", selectors={"model": ".model", "price": ".price", "currency": ".currency"},
                     min_interval_seconds=0, max_attempts=1)
    from su_crawler.models import Product
    config = replace(config, sources=[replace(source, product_ids=["a", "b"])],
                     products=[Product("a", "A", identifiers={"model": "A"}), Product("b", "B", identifiers={"model": "B"})])
    page = b'<div class="item"><b class="model">A</b><b class="price">12000</b><b class="currency">KRW</b></div>'
    def collector(source, base, backend):
        return FetchResult(source.id, "fetched", backend, content=page if backend == "http" else b'<div>Loading</div>',
                           media_type="text/html", final_url=source.location)
    result = execute(config, collector=collector)
    store = Store(tmp_path)
    rows = store.observations(result["id"])
    assert len(rows) == 1 and rows[0]["amount"] == "12000"
    assert next(t for t in store.tasks(result["id"]) if t["product_id"] == "a")["backend"] == "http"
    store.close()


def test_complementary_backends_keep_each_products_original_evidence(tmp_path):
    from su_crawler.models import Product

    config = demo(tmp_path)
    source = replace(
        config.sources[0], kind="web", location="https://example.com/item",
        backends=["http", "playwright"], product_ids=["a", "b"], row_selector=".item",
        selectors={"model": ".model", "price": ".price", "currency": ".currency"},
        min_interval_seconds=0, max_attempts=1,
    )
    config = replace(config, sources=[source], products=[
        Product("a", "A", identifiers={"model": "A"}),
        Product("b", "B", identifiers={"model": "B"}),
    ])
    pages = {
        "http": b'<div class="item"><b class="model">A</b><b class="price">12000</b><b class="currency">KRW</b></div>',
        "playwright": b'<div class="item"><b class="model">B</b><b class="price">23000</b><b class="currency">KRW</b></div>',
    }

    def collector(source, base, backend):
        return FetchResult(source.id, "fetched", backend, content=pages[backend],
                           media_type="text/html", final_url=source.location)

    result = execute(config, collector=collector)
    store = Store(tmp_path)
    rows = {row["product_id"]: row for row in store.observations(result["id"])}
    tasks = {task["product_id"]: task for task in store.tasks(result["id"])}
    store.close()

    assert rows["a"]["amount"] == "12000" and tasks["a"]["backend"] == "http"
    assert rows["b"]["amount"] == "23000" and tasks["b"]["backend"] == "playwright"
    assert rows["a"]["evidence_sha256"] != rows["b"]["evidence_sha256"]
    assert Path(rows["a"]["evidence_path"]).read_bytes() == pages["http"]
    assert Path(rows["b"]["evidence_path"]).read_bytes() == pages["playwright"]


def test_conflict_review_preserves_each_observed_amount():
    rows = [
        {"status": "verified", "amount": "10", "normalized_amount": "10", "comparable": True,
         "comparison_key": "same", "verification_level": "evidence_validated",
         "raw_fields": {"price": "10", "currency": "USD", "condition": "new"}},
        {"status": "verified", "amount": "12", "normalized_amount": "12", "comparable": True,
         "comparison_key": "same", "verification_level": "evidence_validated",
         "raw_fields": {"price": "12", "currency": "USD", "condition": "new"}},
    ]
    _conflicts(rows)
    assert [row["amount"] for row in rows] == ["10", "12"]
    assert all(row["status"] == "review" and row["normalized_amount"] is None for row in rows)
    assert all(not row["comparable"] and row["comparison_key"] is None for row in rows)
    assert all(row["verification_level"] == "review" for row in rows)


def test_expired_at_reexport_is_excluded_without_rewriting_history(tmp_path):
    from openpyxl import load_workbook
    config = demo(tmp_path)
    run = execute(config)
    store = Store(tmp_path)
    row = next(r for r in store.observations(run["id"]) if r["status"] == "verified")
    row["raw_fields"]["valid_to"] = "2000-01-01"
    with store.db:
        store.db.execute("UPDATE observations SET data=? WHERE id=?", (json.dumps(row), row["id"]))
    report_for_run(config, store, run["id"])
    wb = load_workbook(run["report_path"], read_only=True)
    assert wb["Price Comparison"].max_row == 2  # banner + header; the only comparable record is expired
    assert any("validity date at report time" in str(r) for r in wb["Review Required"].values)
    wb.close()
    assert next(r for r in store.observations(run["id"]) if r["id"] == row["id"])["comparable"] is True
    store.close()


def test_prefetched_responses_choose_quality_per_product_and_keep_original_evidence(tmp_path):
    config = demo(tmp_path)
    source = replace(
        config.sources[0], kind="web", location="https://example.com/items",
        backends=["http", "playwright"], product_ids=["a", "b"], max_attempts=1,
    )
    config = replace(config, sources=[source], products=[
        Product("a", "A", identifiers={"model": "A"}),
        Product("b", "B", identifiers={"model": "B"}),
    ])

    def candidate(model, price, locator, visibility, mode):
        fields = {"model": model, "price": price, "currency": "KRW"}
        return Candidate(fields, {key: {"location": locator, "raw": value,
                                        "display_state": visibility} for key, value in fields.items()},
                         locator, mode, source_visibility=visibility, evidence_mode=mode)

    http_candidates = [
        candidate("A", "100", "http:a:monthly", "hidden", "static_html"),
        candidate("A", "90", "http:a:prepay", "hidden", "static_html"),
        candidate("B", "200", "http:b", "hidden", "static_html"),
    ]
    browser_candidates = [candidate("A", "110", "browser:a", "visible", "rendered_dom")]
    http = FetchResult(source.id, "fetched", "http", content=b"static", final_url=source.location)
    browser = FetchResult(source.id, "fetched", "playwright", content=b"rendered", final_url=source.location,
                          screenshot=b"png bytes")

    result = execute(config, prefetched={source.id: [(http, http_candidates), (browser, browser_candidates)]})
    store = Store(tmp_path)
    tasks = {task["product_id"]: task for task in store.tasks(result["id"])}
    rows = store.observations(result["id"])
    store.close()

    assert tasks["a"]["backend"] == "playwright"
    assert tasks["b"]["backend"] == "http"
    assert [row["locator"] for row in rows if row["product_id"] == "a"] == ["browser:a"]
    assert [row["locator"] for row in rows if row["product_id"] == "b"] == ["http:b"]
    rendered = next(row for row in rows if row["product_id"] == "a")
    artifacts = rendered["evidence_artifacts"]
    assert set(artifacts) == {"content", "receipt", "screenshot"}
    receipt = json.loads(Path(artifacts["receipt"]["path"]).read_text(encoding="utf-8"))
    assert receipt["artifacts"] == {key: artifacts[key] for key in ("content", "screenshot")}


def test_captured_responses_are_recorded_after_the_run_budget(tmp_path):
    """The run budget bounds new fetches, not recording what was already collected."""
    config = demo(tmp_path)
    source = replace(config.sources[0], kind="web", location="https://example.com/items",
                     backends=["http"], product_ids=["a"], max_attempts=1)
    config = replace(config, sources=[source], products=[Product("a", "A", identifiers={"model": "A"})],
                     max_run_seconds=0)
    fields = {"model": "A", "price": "100", "currency": "KRW"}
    captured = Candidate(fields, {key: {"location": "row", "raw": value} for key, value in fields.items()}, "row", "html")
    response = FetchResult(source.id, "fetched", "http", content=b"static", final_url=source.location)
    result = execute(config, prefetched={source.id: [(response, [captured])]})
    store = Store(tmp_path)
    tasks, rows = store.tasks(result["id"]), store.observations(result["id"])
    store.close()
    assert [task["status"] for task in tasks] != ["budget_exhausted"]
    assert [row["locator"] for row in rows] == ["row"]


@pytest.mark.parametrize(("price", "separator"), [
    ("12,000", "."),
    ("KRW 12000", "."),
    ("12.000,50", ","),
])
def test_candidate_quality_uses_validation_price_parser(price, separator):
    candidate = Candidate(
        {"model": "A", "price": price},
        {"model": {"location": "row", "raw": "A", "display_state": "visible"},
         "price": {"location": "row", "raw": price, "display_state": "visible"}},
        "row", "html_css", source_visibility="visible", evidence_mode="rendered_dom",
    )
    assert _candidate_quality(candidate, separator)[0] == 1
    assert _has_sufficient_evidence([candidate], separator)


@pytest.mark.parametrize("proof", [
    {"location": "", "raw": "12000", "display_state": "visible"},
    {"location": "row", "raw": "999", "display_state": "visible"},
    {"location": "row", "raw": "12000", "display_state": "hidden"},
])
def test_candidate_quality_rejects_unproven_price(proof):
    candidate = Candidate(
        {"model": "A", "price": "12000"}, {"price": proof}, "row", "html_css",
        source_visibility="visible", evidence_mode="rendered_dom",
    )
    assert _candidate_quality(candidate)[0] == 0
    assert not _has_sufficient_evidence([candidate])


def test_report_export_never_reloads_the_finished_workbook(tmp_path, monkeypatch):
    import openpyxl
    from su_crawler import final_report, xlsx_sheets
    from su_crawler.collection_issues import sync_coverage_issues

    config = demo(tmp_path)
    run = execute(config)
    store = Store(tmp_path)
    try:
        coverage = [{"url": "https://example.org/item", "source_id": "catalog", "status": "visited",
                     "sections": [{"section": "offers", "status": "unknown"}]}]
        store.save_run_coverage(run["id"], coverage)
        sync_coverage_issues(store, run["id"], coverage)
        fail = lambda *a, **k: pytest.fail("finished workbook was reloaded")
        monkeypatch.setattr(openpyxl, "load_workbook", fail)
        monkeypatch.setattr(final_report, "load_workbook", fail)
        monkeypatch.setattr(xlsx_sheets, "_replace_openpyxl", fail)
        result = report_for_run(config, store, run["id"])
        monkeypatch.undo()
        assert Path(result["dashboard_path"]).is_file()
        book = openpyxl.load_workbook(result["report_path"], read_only=True)
        try:
            assert book.sheetnames[-2:] == ["Catalog coverage", "Collection Issues"]
            assert list(book["Collection Issues"].values)[1][3] == "open"
        finally:
            book.close()
    finally:
        store.close()


def test_ledger_indexes_and_atomic_attempt_records(tmp_path):
    store = Store(tmp_path)
    try:
        indexes = {row[0] for row in store.db.execute("SELECT name FROM sqlite_master WHERE type='index'")}
        assert {"observations_run", "observations_task", "attempts_run_source", "browser_submissions_run"} <= indexes
        plan = store.db.execute("EXPLAIN QUERY PLAN DELETE FROM observations WHERE task_id=?", ("t",)).fetchall()
        assert any("observations_task" in str(tuple(row)) for row in plan)
        with store.db:
            store.db.execute("INSERT INTO runs(id,config_hash,status,started_at,demo) VALUES('r','h','running','now',0)")
            for task in ("a", "b"):
                store.db.execute("INSERT INTO tasks(id,run_id,source_id,product_id,status,updated_at) VALUES(?,?,?,?,?,?)",
                                 (task, "r", "s", task, "queued", "now"))
        store.record_attempt("r", "s", ["a", "b"], "http", "fetched", "now", {"message": "", "trace": []})
        assert [row["attempts"] for row in store.tasks("r")] == [1, 1]
        with pytest.raises(Exception):
            # A failed attempt record must not leave counted attempts behind.
            store.record_attempt("missing-run", "s", ["a", "b"], "http", "fetched", "now", {"x": object()})
        assert [row["attempts"] for row in store.tasks("r")] == [1, 1]
        store.finish_tasks([("a", "verified", "", "http", []), ("b", "failed", "x", "http", [])])
        assert [row["status"] for row in store.tasks("r")] == ["verified", "failed"]
    finally:
        store.close()


def test_indexed_matching_returns_exactly_what_the_linear_scan_returns():
    from su_crawler.pipeline import _matches, _matching
    fields = [{"item_id": "A"}, {"item_id": " a "}, {"item_id": "B", "sku": "s1"}, {}, {"item_id": ""},
              {"sku": "S1"}, {"item_id": "B", "sku": "S1"}]
    candidates = [Candidate(f, {}, f"row:{i}", "test") for i, f in enumerate(fields)]
    products = [Product("p1", "A", identifiers={"item_id": "a"}), Product("p2", "B", identifiers={"item_id": "B", "sku": "s1"}),
                Product("p3", "none", identifiers={}), Product("p4", "empty", identifiers={"item_id": ""}),
                Product("p5", "sku", identifiers={"sku": "s1"})]
    cache = {}
    for product in products:
        expected = [c for c in candidates if _matches(c, product)]
        assert _matching(candidates, product, cache) == expected
        assert _matching(candidates, product, cache) == expected  # cached index
