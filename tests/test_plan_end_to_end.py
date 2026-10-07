"""Real local plan flow; only the market fetch and worker process launch are replaced."""
from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import threading
import uuid

import pytest

from su_crawler import assistant_runtime, catalog_collection, plan_jobs
from su_crawler.assistant_server import build_assistant_server
from su_crawler.assistant_workspace import result_observations, result_report
from su_crawler.models import FetchResult
from su_crawler.research_plans import confirmed_snapshot, get_plan
from test_web_browser import launch_browser
from test_web_server import request, json_response, running_server


REQUEST = (
    "Research Acme 100% cotton shirts sold in Japan.\n"
    "Use only the Acme product page below. Exclude used items.\n"
    "Keep original price evidence and show any unmet conditions for review."
)
PRODUCT_URL = "https://catalog.example.test/acme-shirt"
EVIDENCE_URL = "https://catalog.example.test/search/acme"
PREVIEW = {
    "summary": "Acme cotton shirt sources in Japan",
    "topic": {"industry": "Clothing", "product": "Cotton shirts", "market": "Japan"},
    "categories": ["Shirts"], "include_terms": [], "exclude_terms": [],
    "conditions": [
        {"field": "brand", "operator": "equals", "value": "Acme"},
        {"field": "material", "operator": "equals", "value": "cotton 100%"},
        {"field": "condition", "operator": "equals", "value": "new"},
    ],
    "candidates": [{"name": "Acme cotton shirt", "url": PRODUCT_URL,
                    "evidence_url": EVIDENCE_URL, "reason": "Public product page", "kind": "product"}],
    "note": "Suggested reference; price not verified",
}
HTML = (b'<html><body><script type="application/ld+json">'
        b'{"@context":"https://schema.org","@type":"Product","name":"Acme cotton shirt",'
        b'"brand":{"@type":"Brand","name":"Acme"},"material":"100% cotton",'
        b'"itemCondition":"https://schema.org/NewCondition",'
        b'"offers":{"@type":"Offer","price":"19.00","priceCurrency":"JPY"}}'
        b'</script></body></html>')


def _call_mcp(server, name: str, arguments: dict):
    # FastMCP's async call runs in a different thread from Playwright's sync loop.
    with ThreadPoolExecutor(max_workers=1) as pool:
        return pool.submit(lambda: asyncio.run(server.call_tool(name, arguments))).result(timeout=15)[1]


def test_connected_preview_to_confirmed_collection_and_report(tmp_path, monkeypatch):
    pytest.importorskip("mcp")
    playwright = pytest.importorskip("playwright.sync_api")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    collected = []
    worker_threads = []

    def fake_collect(source, base_dir, backend):
        collected.append((source.location, backend))
        return FetchResult(source.id, "fetched", backend, content=HTML, final_url=source.location)

    def start_thread_worker(root):
        base = Path(root)
        token = uuid.uuid4().hex
        runtime_dir = base / "runtime"
        runtime_dir.mkdir(exist_ok=True)
        assistant_runtime._write_json(runtime_dir / "state.json", {
            "status": "starting", "pid": None, "token": token, "current_job_id": None,
            "updated_at": assistant_runtime.utc_now(),
        })
        thread = threading.Thread(target=assistant_runtime._run_worker,
                                  args=(base, token, 4.0), daemon=True)
        thread.start()
        worker_threads.append(thread)
        return {"status": "starting"}

    monkeypatch.setattr(catalog_collection, "collect", fake_collect)
    monkeypatch.setattr(plan_jobs.runtime, "start_worker", start_thread_worker)
    mcp_server = build_assistant_server(workspace)
    artifacts = Path(__file__).resolve().parents[1] / ".tmp"
    artifacts.mkdir(exist_ok=True)
    page_errors = []
    job_id = None

    try:
        with running_server(workspace) as http_server, playwright.sync_playwright() as driver:
            browser = launch_browser(driver)
            try:
                page = browser.new_page(viewport={"width": 1440, "height": 960}, accept_downloads=True)
                page.set_default_timeout(10000)
                page.on("pageerror", lambda error: page_errors.append(str(error)))
                page.goto(f"{http_server.url}/#plan")
                playwright.expect(page.locator("#plan-ai-path")).to_have_value("local")
                page.locator("#plan-ai-path").select_option("host")
                page.locator("#plan-request-text").fill(REQUEST)
                page.locator("#plan-create").click()
                playwright.expect(page.locator("#plan-workspace")).to_be_visible()
                _, first = json_response(request(http_server, "GET", "/api/bootstrap"))
                assert first["status"] == "needs_setup" and first["jobs"] == []
                assert not (workspace / "research.json").exists()
                draft = first["research_plans"]["plans"][0]
                assert draft["request_text"] == REQUEST and draft["revision"] == 1
                assert _call_mcp(mcp_server, "submit_research_preview", {
                    "plan_id": draft["id"], "expected_revision": 1, "preview": PREVIEW,
                })["plan"]["revision"] == 2
                playwright.expect(page.locator("#plan-candidates .plan-candidate")).to_have_count(1, timeout=10000)
                playwright.expect(page.locator("#plan-candidates a.plan-url")).to_have_count(2)
                assert page.locator("#plan-candidates a.plan-url").first.get_attribute("href") == PRODUCT_URL
                assert page.locator("#plan-product").input_value() == "Cotton shirts"
                assert page.locator("#plan-exclude").input_value() == ""
                assert get_plan(workspace, draft["id"])["plan"]["conditions"] == PREVIEW["conditions"]
                assert get_plan(workspace, draft["id"])["plan"]["request_text"] == REQUEST
                page.screenshot(path=str(artifacts / "research-plan-desktop.png"), full_page=True)

                choice = page.locator("#plan-candidates [data-plan-select]")
                choice.uncheck()
                choice.check()
                page.locator("#plan-start").click()
                playwright.expect(page.locator("#job-detail")).to_contain_text("Succeeded", ignore_case=True, timeout=30000)
                _, final = json_response(request(http_server, "GET", "/api/bootstrap"))
                approved = final["research_plans"]["plans"][0]
                assert approved["state"] == "confirmed"
                assert approved["request_text"] == REQUEST
                assert approved["conditions"] == PREVIEW["conditions"]
                assert confirmed_snapshot(workspace, approved["id"], approved["revision"]) == approved
                assert (workspace / "research.json").is_file()
                jobs = [job for job in final["jobs"] if job["operation"] == "research_plan"]
                assert len(jobs) == 1 and jobs[0]["status"] == "succeeded"
                job_id = jobs[0]["id"]
                assert collected and all(url == PRODUCT_URL for url, _ in collected)
                job = json_response(request(http_server, "GET", f"/api/jobs/{job_id}"))[1]
                result = job["result"]
                assert result["plan_revision"] == approved["revision"]
                assert result["result"]["coverage"]
                assert result["result"]["coverage"][0]["url"] == PRODUCT_URL
                assert result["result"]["coverage"][0]["scope_counts"]["matched"] == 1
                assert result["result"]["coverage"][0]["scope_counts"]["unknown"] == 0
                rows = result_observations(workspace, job)
                assert rows["total"] == 1
                assert rows["rows"][0]["raw_fields"]["name"] == "Acme cotton shirt"
                assert rows["rows"][0]["status"] == "review"
                report = Path(result_report(workspace, job)["path"])
                assert report.is_file() and report.stat().st_size > 0
                from openpyxl import load_workbook
                workbook = load_workbook(report, read_only=True)
                try:
                    assert "Source Evidence" in workbook.sheetnames
                finally:
                    workbook.close()
                playwright.expect(page.locator("#job-detail .plan-coverage-row")).to_have_count(1)
                playwright.expect(page.locator("#job-detail table tbody tr")).to_have_count(1)
                detail_text = page.locator("#job-detail").inner_text()
                assert "Product condition checks" in detail_text
                assert "matched" in detail_text.lower()
                assert all(key not in detail_text for key in
                           ("plan.scopeNote", "plan.scopeBound", "plan.scopeDetails"))
                with page.expect_download() as download_event:
                    page.get_by_role("link", name="Download XLSX").click()
                assert download_event.value.failure() is None
                page.set_viewport_size({"width": 375, "height": 850})
                assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
                page.screenshot(path=str(artifacts / "research-plan-mobile.png"), full_page=True)
                assert page_errors == []
            finally:
                browser.close()
    finally:
        for thread in worker_threads:
            thread.join(timeout=10)
    assert job_id is not None
