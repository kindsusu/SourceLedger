"""Cross-component snapshot provenance through catalog, SQLite and XLSX."""
import hashlib
from pathlib import Path

from su_crawler.catalog_collection import collect_plan
from su_crawler.models import FetchResult
from su_crawler.storage import Store


def test_discovery_retains_links_from_all_pagination_snapshots():
    from su_crawler.catalog_collection import _links
    result = FetchResult('listing', 'fetched', 'playwright', b'<a href="/third">C</a>',
                         final_url='https://store.example/catalog', variant_snapshots=[
                             {'content': b'<a href="/first">A</a><a href="https://outside.example/x">X</a>'},
                             {'content': b'<a href="/second">B</a><a href="/first">A</a>'},
                         ])
    assert set(_links(result, result.final_url, 'store.example', {'https://store.example/second'})) == {
        'https://store.example/first', 'https://store.example/third'}


def test_failed_page_keeps_prior_capture_reference_in_issue_report(tmp_path):
    url = 'https://store.example/item'
    def fetch(source, base, backend):
        return FetchResult(source.id, 'timeout', backend, message='Later state timed out',
                           variant_snapshots=[{'content': b'<p>Earlier state</p>', 'final_url': url}])
    plan = {'id': 'failed-options', 'revision': 1, 'state': 'confirmed',
            'topic': {'industry': 'Parts', 'product': 'Bolt', 'market': 'US'},
            'candidates': [{'url': url, 'selected': True}],
            'collection_rules': {url: {'selectors': {'name': 'p'}}}}
    result = collect_plan(tmp_path, plan, output_dir=tmp_path/'out', max_pages=1,
                          collector=fetch, follow_links=False)
    store = Store(tmp_path/'out')
    try:
        assert not store.observations(result['run_id'])
        assert result['coverage'][0]['status'] == 'timeout'
        issue = store.collection_issues(result['run_id'])[0]
        assert Path(issue['evidence_ref']).is_file()
    finally:
        store.close()
    from openpyxl import load_workbook
    book = load_workbook(result['report_path'], read_only=True)
    assert book['Collection Issues']['I2'].value == issue['evidence_ref']
    book.close()


def test_each_option_keeps_its_own_source_bytes_and_ledger_identity(tmp_path):
    url = 'https://store.example/shirt'
    snapshots = [{'content': f'<h1>Shirt</h1><b>{price}</b><i>USD</i>'.encode(),
                  'final_url': url, 'fetched_at': '2026-10-14T00:00:00+00:00',
                  'selections': {'#size': size}}
                 for size, price in [('M', 10), ('L', 12)]]
    def fetch(source, base, backend):
        return FetchResult(source.id, 'fetched', backend, snapshots[-1]['content'], final_url=url,
                           variant_snapshots=snapshots,
                           trace=[{'event': 'option_enumeration', 'status': 'complete',
                                   'reason': 'options_exhausted', 'captured_count': 2}])
    plan = {'id': 'options', 'revision': 1, 'state': 'confirmed',
            'topic': {'industry': 'Clothing', 'product': 'Shirt', 'market': 'US'},
            'candidates': [{'url': url, 'selected': True}],
            'collection_rules': {url: {'selectors': {'name': 'h1', 'price': 'b', 'currency': 'i'}}}}
    result = collect_plan(tmp_path, plan, output_dir=tmp_path/'out', max_pages=1,
                          collector=fetch, follow_links=False, checkpoint_id='snapshots')
    ledger = Store(tmp_path/'out')
    try:
        rows = ledger.observations(result['run_id'])
        assert len(rows) == 2
        assert {row['raw_fields']['spec:selected:#size'] for row in rows} == {'M', 'L'}
        assert {row['raw_fields']['price'] for row in rows} == {'10', '12'}
        for row in rows:
            content = Path(row['evidence_path']).read_bytes()
            assert hashlib.sha256(content).hexdigest() == row['evidence_sha256']
            assert ('<b>' + row['raw_fields']['price'] + '</b>').encode() in content
        sections = result['coverage'][0]['sections']
        assert next(s for s in sections if s['section'] == 'option_enumeration')['status'] == 'complete'
        assert result['coverage_summary']['unknown'] == 1
        assert ledger.collection_issues(result['run_id'])
    finally:
        ledger.close()
    from openpyxl import load_workbook
    book = load_workbook(result['report_path'], read_only=True)
    assert 'Collection Issues' in book.sheetnames
    book.close()
