"""Offline tests for durable collection gap triage and conservative retry rules."""
import pytest

from su_crawler.collection_issues import sync_coverage_issues
from su_crawler.storage import Store


def _run(store, run_id, started):
    store.db.execute('''INSERT INTO runs(id,config_hash,status,started_at,demo)
                        VALUES(?,?,?,?,0)''', (run_id, 'config', 'completed', started))
    store.db.commit()


def test_issue_ledger_is_idempotent_and_resolution_requires_evidence(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        coverage = [{'url': 'https://example.test/item', 'source_id': 'source-a',
                     'adapter': 'example', 'status': 'visited',
                     'sections': [{'section': 'offers', 'status': 'partial',
                                   'reason': 'max_items reached', 'missing_locators': ['card-2'],
                                   'evidence_ref': 'capture.html'}]}]
        store.save_run_coverage('run-1', coverage)
        sync_coverage_issues(store, 'run-1', coverage)
        sync_coverage_issues(store, 'run-1', coverage)
        rows = store.collection_issues('run-1')
        assert len(rows) == 2
        assert {row['reason_code'] for row in rows} == {'exhausted_limit', 'unverified_option'}
        assert rows[0]['evidence']['evidence_ref'] == 'capture.html'
        with pytest.raises(ValueError, match='explicit evidence'):
            store.update_collection_issue(rows[0]['id'], status='resolved')
        with pytest.raises(ValueError, match='later run'):
            store.update_collection_issue(rows[0]['id'], status='resolved',
                                          evidence={'kind': 'revalidated_run', 'reference': 'run-1'})
        _run(store, 'run-2', '2026-01-02T00:00:00+00:00')
        store.save_run_coverage('run-2', [{**coverage[0], 'coverage_status': 'complete',
                                          'sections': [{'section': 'offers', 'status': 'complete',
                                                        'extracted_locators': ['card-2']}]}])
        resolved = store.update_collection_issue(rows[0]['id'], status='resolved',
                                                 evidence={'kind': 'revalidated_run', 'reference': 'run-2'})
        assert resolved['status'] == 'resolved'
        assert [(event['previous_status'], event['new_status']) for event in
                store.collection_issue_events(rows[0]['id'])] == [('open', 'resolved')]
        assert resolved['url'] == coverage[0]['url'] and resolved['section'] == 'offers'
        assert resolved['category'] == 'exhausted_limit' and resolved['attempt_count'] == 0
        assert store.run_coverage('run-1') == coverage
    finally:
        store.close()


def test_access_policy_and_unknown_have_distinct_actions(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        pages = [{'url': f'https://example.test/{code}', 'status': code, 'reason': reason}
                 for code, reason in [('policy_denied', 'robots disallowed'),
                                      ('needs_auth', 'login required'),
                                      ('blocked', 'interactive challenge'),
                                      ('no_data', 'No products found')]]
        sync_coverage_issues(store, 'run-1', pages)
        rows = store.collection_issues('run-1')
        assert {row['reason_code'] for row in rows} == {
            'policy_denied', 'needs_auth', 'browser_handoff', 'unknown_enumeration'}
        assert {row['next_action'] for row in rows} == {
            'review_policy', 'provide_authorized_access', 'submit_browser_evidence',
            'inspect_scope_and_recipe'}
    finally:
        store.close()


def test_prior_comparable_complete_section_only_signals_drift(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        _run(store, 'run-2', '2026-01-02T00:00:00+00:00')
        old = {'url': 'https://example.test/item', 'source_id': 'stable', 'adapter': 'site',
               'status': 'visited', 'sections': [{'section': 'offers', 'status': 'complete'}]}
        store.save_run_coverage('run-1', [old])
        current = {**old, 'sections': [{'section': 'page_scope', 'status': 'unknown'}]}
        sync_coverage_issues(store, 'run-2', [current])
        rows = store.collection_issues('run-2')
        assert {row['reason_code'] for row in rows} == {'structure_changed', 'unknown_enumeration'}
        assert next(row for row in rows if row['reason_code'] == 'structure_changed')['next_action'] == 'inspect_recipe'
    finally:
        store.close()


def test_unrelated_page_completion_and_manual_review_cannot_resolve(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        _run(store, 'run-2', '2026-01-02T00:00:00+00:00')
        issue = sync_coverage_issues(store, 'run-1', [{
            'url': 'https://example.test/item', 'source_id': 'source-a', 'adapter': 'site',
            'status': 'visited', 'sections': [{'section': 'offers', 'status': 'unknown'}]}])[0]
        store.save_run_coverage('run-2', [{
            'url': 'https://example.test/item', 'source_id': 'different', 'adapter': 'site',
            'coverage_status': 'complete', 'sections': [{'section': 'other', 'status': 'complete'}]}])
        with pytest.raises(ValueError, match='same URL, source and adapter'):
            store.update_collection_issue(issue['id'], status='resolved',
                                          evidence={'kind': 'revalidated_run', 'reference': 'run-2'})
        with pytest.raises(ValueError, match='only resolve'):
            store.update_collection_issue(issue['id'], status='resolved',
                                          evidence={'kind': 'scope_review', 'reference': 'note',
                                                    'decision': 'reviewed'})
        assert store.update_collection_issue(issue['id'], status='needs_review',
                                             evidence={'kind': 'scope_review', 'reference': 'note',
                                                       'decision': 'needs inspection'})['status'] == 'needs_review'
        assert store.collection_issue_events(issue['id'])[0]['evidence']['decision'] == 'needs inspection'
    finally:
        store.close()


def test_browser_environment_and_403_are_not_retryable(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        pages = [{'url': f'https://example.test/{index}', 'status': status, 'reason': reason}
                 for index, (status, reason) in enumerate([
                     ('blocked', 'HTTP 403 Forbidden'), ('tool_unavailable', 'browser_launch failed'),
                     ('failed', 'HTTP 429 rate limit')])]
        rows = sync_coverage_issues(store, 'run-1', pages)
        assert [row['category'] for row in rows] == [
            'browser_handoff', 'browser_environment', 'transient_failure']
    finally:
        store.close()


def test_enumeration_attempt_evidence_is_retained(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        enumeration = {'status': 'partial', 'skipped_options': ['size L'],
                       'unavailable_options': ['size XL']}
        issue = sync_coverage_issues(store, 'run-1', [{
            'url': 'https://example.test/item', 'status': 'visited',
            'sections': [{'section': 'variants', 'status': 'partial',
                          'reason': 'Some controls unavailable',
                          'enumeration': enumeration}]}])[0]
        assert issue['evidence']['enumeration'] == enumeration
    finally:
        store.close()


def test_issue_pagination_passes_500_without_losing_late_rows(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        pages = [{'url': f'https://example.test/item/{index:03d}',
                  'status': 'timeout', 'reason': 'temporary failure'}
                 for index in range(505)]
        sync_coverage_issues(store, 'run-1', pages)
        assert len(store.collection_issues('run-1', limit=500)) == 500
        later = store.collection_issues('run-1', offset=500, limit=500)
        assert len(later) == 5
        assert later[-1]['url'].endswith('/504')
        assert store.collection_issue_count(['run-1']) == 505
    finally:
        store.close()


def test_triage_refreshes_issue_sheet_and_retains_coverage_sheet(tmp_path, monkeypatch):
    import mimetypes
    monkeypatch.setattr(mimetypes.MimeTypes, 'read_windows_registry', lambda self: None)
    from openpyxl import Workbook, load_workbook
    from su_crawler import research_followup

    store = Store(tmp_path)
    report = tmp_path / 'prices-run-1.xlsx'
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        store.db.execute('UPDATE runs SET report_path=? WHERE id=?', (str(report), 'run-1'))
        store.db.commit()
        issue = sync_coverage_issues(store, 'run-1', [{
            'url': 'https://example.test/item', 'source_id': 'source-a',
            'status': 'visited', 'sections': [{'section': 'offers', 'status': 'unknown'}]}])[0]
        wb = Workbook()
        coverage = wb.active
        coverage.title = 'Catalog coverage'
        coverage.append(['immutable coverage marker'])
        wb.save(report)
        wb.close()
    finally:
        store.close()
    monkeypatch.setattr(research_followup.runtime, 'get_job',
                        lambda root, job_id: {'operation': 'research_plan'})
    monkeypatch.setattr(research_followup, 'completed',
                        lambda root, job_id: ({}, {'run_id': 'run-1'}, tmp_path))
    changed = research_followup.triage_collection_issue(
        tmp_path, 'job-1', issue['id'], status='needs_review')
    assert changed['status'] == 'needs_review'
    wb = load_workbook(report, read_only=True, data_only=True)
    try:
        assert wb['Catalog coverage']['A1'].value == 'immutable coverage marker'
        assert wb['Collection Issues']['D2'].value == 'needs_review'
    finally:
        wb.close()


def test_revalidation_accepts_reviewed_recipe_revision_only_for_same_scope(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'old', '2026-01-01T00:00:00+00:00')
        _run(store, 'new', '2026-01-02T00:00:00+00:00')
        url = 'https://example.test/item'
        issue = sync_coverage_issues(store, 'old', [{
            'url': url, 'source_id': 'source-a', 'adapter': 'site',
            'recipe_fingerprint': 'recipe-old', 'plan_fingerprint': 'plan-old',
            'status': 'visited', 'sections': [{'section': 'offers', 'status': 'unknown'}]}])[0]
        store.save_run_coverage('new', [{
            'url': url, 'source_id': 'source-a', 'adapter': 'site',
            'recipe_fingerprint': 'recipe-new', 'plan_fingerprint': 'plan-new',
            'coverage_status': 'complete',
            'sections': [{'section': 'offers', 'status': 'complete'}]}])
        base = {'kind': 'revalidated_run', 'reference': 'new'}
        with pytest.raises(ValueError, match='scope_revision'):
            store.update_collection_issue(issue['id'], status='resolved', evidence=base)
        revision = {'reason': 'Reviewed changed extraction recipe and confirmed the same offer section',
                    'changed': {'recipe_fingerprint': {'from': 'recipe-old', 'to': 'recipe-new'},
                                'plan_fingerprint': {'from': 'plan-old', 'to': 'plan-new'}}}
        assert store.update_collection_issue(issue['id'], status='resolved',
                                             evidence={**base, 'scope_revision': revision})['status'] == 'resolved'
    finally:
        store.close()


def test_triage_reports_xlsx_refresh_pending_after_committed_db_change(tmp_path, monkeypatch):
    import mimetypes
    monkeypatch.setattr(mimetypes.MimeTypes, 'read_windows_registry', lambda self: None)
    from openpyxl import Workbook
    from su_crawler import research_followup, report as report_module

    store = Store(tmp_path)
    report = tmp_path / 'prices-run-1.xlsx'
    try:
        _run(store, 'run-1', '2026-01-01T00:00:00+00:00')
        store.db.execute('UPDATE runs SET report_path=? WHERE id=?', (str(report), 'run-1'))
        store.db.commit()
        issue = sync_coverage_issues(store, 'run-1', [{
            'url': 'https://example.test/item', 'source_id': 'source-a',
            'status': 'visited', 'sections': [{'section': 'offers', 'status': 'unknown'}]}])[0]
        wb = Workbook(); wb.save(report); wb.close()
    finally:
        store.close()
    monkeypatch.setattr(research_followup.runtime, 'get_job',
                        lambda root, job_id: {'operation': 'research_plan'})
    monkeypatch.setattr(research_followup, 'completed',
                        lambda root, job_id: ({}, {'run_id': 'run-1'}, tmp_path))
    monkeypatch.setattr(report_module, 'append_collection_issues',
                        lambda *args: (_ for _ in ()).throw(PermissionError('workbook open')))
    changed = research_followup.triage_collection_issue(
        tmp_path, 'job-1', issue['id'], status='needs_review')
    assert changed['status'] == 'needs_review' and changed['report_refresh'] == 'pending'
    store = Store(tmp_path)
    try:
        assert store._collection_issue(issue['id'])['status'] == 'needs_review'
    finally:
        store.close()


def test_generic_selector_option_revalidation_without_adapter(tmp_path):
    store = Store(tmp_path)
    try:
        _run(store, 'old', '2026-01-01T00:00:00+00:00')
        _run(store, 'new', '2026-01-02T00:00:00+00:00')
        url = 'https://example.test/item'
        issue = next(row for row in sync_coverage_issues(store, 'old', [{
            'url': url, 'source_id': 'generic-source', 'adapter': None,
            'recipe_fingerprint': 'recipe-1', 'status': 'visited',
            'sections': [{'section': 'variants', 'status': 'partial',
                          'reason': 'unparsed option', 'missing_locators': ['size-m']}]}])
                     if row['option_key'] == 'size-m')
        store.save_run_coverage('new', [{
            'url': url, 'source_id': 'generic-source', 'adapter': None,
            'recipe_fingerprint': 'recipe-1', 'coverage_status': 'complete',
            'sections': [{'section': 'variants', 'status': 'complete',
                          'extracted_locators': ['size-m']}]}])
        assert store.update_collection_issue(issue['id'], status='resolved', evidence={
            'kind': 'revalidated_run', 'reference': 'new'})['status'] == 'resolved'
    finally:
        store.close()
