from __future__ import annotations

from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import subprocess
import sys
import threading

from openpyxl import load_workbook

from su_crawler.storage import Store


ROOT = Path(__file__).resolve().parents[1]


class _Site(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/robots.txt":
            body = b"User-agent: *\nAllow: /\n"
        elif self.path == "/good":
            body = (b'<article class="product"><span data-field="model">TEST-A</span>'
                    b'<span data-field="price">12000</span><span data-field="currency">KRW</span>'
                    b'<span data-field="unit">each</span><span data-field="pack_quantity">1</span>'
                    b'<span data-field="tax">included</span><span data-field="price_type">retail</span>'
                    b'<span data-field="price_basis">each</span></article>')
        elif self.path == "/missing":
            body = b'<article><span data-field="model">TEST-A</span><span>no price</span></article>'
        else:
            body = b"missing"
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


@contextmanager
def site():
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Site)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown(); server.server_close(); thread.join(timeout=5)


def cli(*args: str, expected=0) -> dict:
    result = subprocess.run([sys.executable, "-m", "su_crawler", *args], cwd=ROOT,
                            text=True, encoding="utf-8", capture_output=True, timeout=45)
    assert result.returncode == expected, result.stdout + result.stderr
    return json.loads(result.stdout)


def prepare(tmp_path: Path, origin: str, paths=("/good",)) -> tuple[Path, list[str]]:
    workspace = tmp_path / "research.json"
    cli("init", "--workspace", str(workspace), "--industry", "Parts", "--product", "Exact part", "--market", "Korea")
    cli("product-set", "--workspace", str(workspace), "--identifier", "model=TEST-A")
    ids = []
    for path in paths:
        value = cli("source-add", "--workspace", str(workspace), "--url", origin + path, "--scope", "internal")
        ids.append(value["sources"][-1]["id"])
    return workspace, ids


def test_agent_cli_collects_and_verifies_one_final_ledger(tmp_path):
    with site() as origin:
        workspace, ids = prepare(tmp_path, origin)
        state = cli("agent", "--workspace", str(workspace), "--run-dir", str(tmp_path / "run"),
                    "--source-id", ids[0], "--max-seconds", "30")
    assert state["status"] == "completed"
    assert state["tasks"][0]["status"] == "proposed"
    assert Path(state["collection"]["report_path"]).is_file()
    store = Store(tmp_path / "outputs")
    try:
        rows = store.observations(state["collection"]["run_id"])
    finally:
        store.close()
    assert len(rows) == 1
    assert rows[0]["amount"] == "12000"
    book = load_workbook(state["collection"]["report_path"], read_only=True)
    try:
        assert "Price Comparison" in book.sheetnames
    finally:
        book.close()


def test_agent_pause_resume_does_not_repeat_completed_source(tmp_path):
    with site() as origin:
        workspace, ids = prepare(tmp_path, origin, ("/good", "/missing"))
        initial = cli("agent", "--workspace", str(workspace), "--run-dir", str(tmp_path / "run"),
                      "--source-id", ids[0], "--source-id", ids[1], "--max-steps", "1", "--max-seconds", "30")
        resumed = cli("agent", "--workspace", str(workspace), "--run-dir", str(tmp_path / "run"), "--resume", expected=1)
    assert initial["status"] == "paused"
    # The first source reached a terminal proposal outcome. A missing or
    # temporarily unavailable browser can correctly leave static HTML in
    # needs_review; resume must still not repeat that completed attempt.
    assert initial["tasks"][0]["status"] in {"proposed", "verified", "needs_review"}
    assert resumed["tasks"][0]["attempt_dir"] == initial["tasks"][0]["attempt_dir"]
    assert resumed["tasks"][1]["status"] in {"needs_review", "failed"}
    assert resumed["status"] == "needs_review"


def test_agent_activation_with_known_sample(tmp_path):
    with site() as origin:
        workspace, ids = prepare(tmp_path, origin)
        samples = tmp_path / "samples.json"
        samples.write_text(json.dumps({"schema": "source-ledger/known-samples/v1", "samples": [{
            "source_id": ids[0], "product_id": "product", "price": "12000", "currency": "KRW"}]}), encoding="utf-8")
        state = cli("agent", "--workspace", str(workspace), "--run-dir", str(tmp_path / "run"),
                    "--source-id", ids[0], "--samples", str(samples), "--activate", "--max-seconds", "30")
    assert state["status"] == "completed"
    assert state["tasks"][0]["status"] == "proposed"
    assert state["collection"]["activation"]["status"] == "activated"
