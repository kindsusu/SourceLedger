from pathlib import Path

from openpyxl import load_workbook

from su_crawler.models import CollectionConfig, Product, Source
from su_crawler.report import SHEETS, append_collection_issues, append_section_coverage, export_report


def test_results_append_field_proof_and_comparison_status_columns(tmp_path):
    observation = {
        "id": "o", "source_url": "https://shop.example/p", "source_name": "Shop",
        "extraction_method": "adapter", "status": "review", "comparable": False,
        "amount": "7710000", "currency": "KRW",
        "raw_fields": {"name": "Car", "price": "7710000", "currency": "KRW"},
        "evidence": {"price": {"location": "quote", "raw": "7,710,000원",
                                "normalization": {"rule": "korean_won_amount"}}},
    }
    path = export_report(_config(tmp_path), {"id": "r"}, [], [observation], tmp_path / "proof.xlsx")
    wb = load_workbook(path, read_only=True)
    try:
        header, data = list(wb["Results"].values)
        assert header[-3:] == ("Price Field Evidence", "Declared Normalization", "Comparison Readiness")
        assert data[-3:] == ("matched", "matched", "not_ready")
        assert data[4] == "review"
    finally:
        wb.close()


def test_supplemental_tables_join_single_export_without_formula_execution(tmp_path):
    path = export_report(_config(tmp_path), {"id": "r"}, [], [], tmp_path / "all.xlsx",
                         supplemental_sheets=[{"name": "Catalog", "headers": ["Product", "Amount"],
                                               "rows": [["=HYPERLINK(\"unsafe\")", 1234], ["Unknown", None]]}])
    wb = load_workbook(path)
    assert "Results" in wb.sheetnames
    assert wb["Catalog"]["A2"].data_type == "s"
    assert wb["Catalog"]["B2"].value == 1234
    assert wb["Catalog"]["B3"].value is None
    assert wb["Catalog"].auto_filter.ref == "A1:B3"
    wb.close()


def test_catalog_coverage_sheet_replaces_rows_and_preserves_evidence(tmp_path):
    cfg = _config(tmp_path)
    path = export_report(cfg, {"id": "r"}, [], [], tmp_path / "sections.xlsx")
    coverage = [{"url": "https://example.test/rentpass", "sections": [{
        "section": "popular", "status": "partial", "observed_count": 8, "expected_count": 8,
        "extracted_count": 7, "price_count": 6, "unparsed_count": 1,
        "click_count": 2, "reason": "more control stalled", "evidence_ref": "evidence/popular.html",
    }], "collection_issues": [{"id": "issue-1", "url": "https://example.test/rentpass",
                              "section": "popular", "status": "open", "reason": "more control stalled",
                              "category": "browser_handoff", "next_action": "Open in browser",
                              "attempt_count": 2, "evidence_ref": "evidence/popular.html"}]}]
    assert append_section_coverage(path, coverage) == path
    append_section_coverage(path, coverage)
    wb = load_workbook(path, read_only=True)
    assert wb.sheetnames.count("Catalog coverage") == 1
    values = list(wb["Catalog coverage"].values)
    assert values[1] == ("https://example.test/rentpass", "popular", "partial", 8, 7, 6, 1, 2,
                         "more control stalled", "evidence/popular.html", None, 8)
    assert list(wb["Collection Issues"].values)[1] == ("issue-1", "https://example.test/rentpass",
        "popular", "open", "browser_handoff", "more control stalled", "Open in browser", 2,
        "evidence/popular.html")


def test_collection_issues_sheet_preserves_actions_and_evidence(tmp_path):
    path = export_report(_config(tmp_path), {"id": "r"}, [], [], tmp_path / "issues.xlsx")
    issue = {"id": "issue-1", "url": "https://example.test/list", "section": "More offers",
             "status": "open", "reason": '=HYPERLINK("bad")', "category": "browser_handoff",
             "next_action": "Open the page in the browser", "attempt_count": 2,
             "evidence_ref": "evidence/list.html"}
    append_collection_issues(path, [issue])
    append_collection_issues(path, [issue])
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        assert wb.sheetnames.count("Collection Issues") == 1
        row = list(wb["Collection Issues"].values)[1]
        assert row == ("issue-1", "https://example.test/list", "More offers", "open",
                       "browser_handoff", '=HYPERLINK("bad")', "Open the page in the browser",
                       2, "evidence/list.html")
        assert wb["Collection Issues"]["F2"].data_type == "s"
    finally:
        wb.close()


def test_pages_without_section_inventory_are_reported_as_unknown(tmp_path):
    cfg = _config(tmp_path)
    path = export_report(cfg, {"id": "r"}, [], [], tmp_path / "sections.xlsx")
    append_section_coverage(path, [{"url": "https://example.test/product", "evidence_path": "evidence/page.html"},
                                   {"url": "https://example.test/irrelevant", "status": "excluded"},
                                   {"url": "https://example.test/timeout", "status": "timeout"}])
    wb = load_workbook(path, read_only=True)
    try:
        rows = list(wb["Catalog coverage"].values)
        assert len(rows) == 4
        assert rows[1][:3] == ("https://example.test/product", "Page scope", "unknown")
        assert rows[1][3:8] == (None, None, None, None, None)
        assert rows[1][8] == "No section inventory recorded; completeness not checked."
        assert rows[1][10] == "evidence/page.html"
        assert rows[1][11] is None
        assert rows[2][:3] == ("https://example.test/irrelevant", "Page scope", "excluded")
        assert rows[2][3:8] == (None, None, None, None, None)
        assert rows[2][11] is None
        assert rows[3][:3] == ("https://example.test/timeout", "Page scope", "partial")
        assert rows[3][11] is None
    finally:
        wb.close()


def test_browser_only_size_variant_is_visible_in_results_and_evidence(tmp_path):
    cfg = _config(tmp_path)
    base = {"id": "base", "source_url": "https://shop.example/p", "source_name": "Shop",
            "extraction_method": "adapter", "status": "review", "raw_fields": {"name": "Shirt", "size": "M", "price": "20000"}}
    proof = {"id": "size-l", "source_url": "https://shop.example/p", "source_name": "Shop",
             "extraction_method": "host_browser_submission", "raw_fields": {"name": "Shirt", "size": "L", "price": "22000"}}
    path = export_report(cfg, {"id": "r"}, [], [base, proof], tmp_path / "variants.xlsx")
    wb = load_workbook(path, read_only=True)
    try:
        results = list(wb["Results"].values)
        assert len(results) == 3
        assert results[2][4] == "review"
        assert "Size: L" in results[2][5]
        assert "22000" in results[2][5]
        assert results[2][10] == "size-l"
        evidence = list(wb["Supporting Evidence"].values)
        assert evidence[1][0] == "size-l" and evidence[1][2] == "unmatched"
    finally:
        wb.close()


def test_browser_component_spec_and_quantity_tier_stay_visible_in_xlsx(tmp_path):
    cfg = _config(tmp_path)
    proof = {"id": "component", "source_url": "https://shop.example/p", "source_name": "Shop",
             "extraction_method": "host_browser_submission", "raw_fields": {
                 "name": "Resistor", "price": 0, "spec:resistance": "10kΩ", "spec:tolerance": "1%",
                 "pack_quantity": "10", "quantity_tier": "10-99"}}
    path = export_report(cfg, {"id": "r"}, [], [proof], tmp_path / "components.xlsx")
    wb = load_workbook(path, read_only=True)
    try:
        row = list(wb["Results"].values)[1]
        assert row[4] == "review" and row[2] is None
        assert "Displayed price text (review required): 0" in row[5]
        assert "spec:resistance: 10kΩ" in row[5]
        assert "Quantity tier: 10-99" in row[5]
        assert "Pack quantity: 10" in row[5]
    finally:
        wb.close()


def _config(tmp_path):
    return CollectionConfig("t", [Product("p1", "제품 1")], [Source("s1", "출처 1", "web", "https://shop.example/p", ["p1"])], str(tmp_path), str(tmp_path))


def test_evidence_first_workbook(tmp_path):
    cfg = _config(tmp_path)
    out = export_report(cfg, {"id": "r1", "status": "done", "started_at": "2026-01-01T00:00:00Z", "demo": True},
        [{"id": "t1", "product_id": "p1", "source_id": "s1", "status": "failed", "attempts": 0, "reason": "가격 없음", "updated_at": "2026-01-01T00:00:00Z"}],
        [{"id": "o0", "task_id": "t1", "product_id": "p1", "source_id": "s1", "source_name": "출처 1", "source_url": "https://shop.example/p", "collected_at": "2026-01-01T00:00:00Z", "status": "verified", "reason": "", "amount": "0", "currency": "KRW", "unit": "ea", "raw_fields": {"price": "0", "name": "=HYPERLINK(\"bad\")"}, "evidence": {"amount": {"location": "#price", "raw": "0"}}, "evidence_path": "evidence/a.html", "evidence_sha256": "abc", "locator": "#price", "extraction_method": "dom", "comparable": True, "comparison_key": "p1|KRW|ea", "freshness": "observed"},
         {"id": "o1", "task_id": "t1", "product_id": "p1", "source_id": "s1", "source_name": "출처 1", "source_url": "https://shop.example/p", "collected_at": "2026-01-01T00:00:00Z", "status": "review", "reason": "단위 미확인", "amount": None, "raw_fields": {"price": "문의"}, "evidence": {}, "evidence_path": "", "evidence_sha256": "", "locator": "", "extraction_method": "dom", "comparable": False, "freshness": "observed"}], out := tmp_path / "report.xlsx")
    assert out.exists()
    wb = load_workbook(out, read_only=True, data_only=False)
    assert wb.sheetnames == ["Results", *SHEETS]
    assert wb["Price Comparison"]["E3"].value == 0
    assert wb["Price Comparison"]["F3"].value == "0"
    assert wb["Observation History"]["H4"].value is None
    assert wb["Review Required"]["G3"].value is None
    assert wb["Review Required"]["H3"].value is None
    assert wb["Observation History"]["Q3"].value.startswith('{')  # formula-looking raw source stays text
    assert wb["Run Summary"]["A1"].value == "Demo data — not market prices"
    assert wb["Collection Status"].max_row == 3  # banner + header + planned task


def test_comparison_uses_latest_source_observation_and_group_statistics(tmp_path):
    cfg = CollectionConfig("t", [Product("p1", "제품 1")], [
        Source("s1", "출처 1", "web", "https://a.example", ["p1"]),
        Source("s2", "출처 2", "web", "https://b.example", ["p1"]),
    ], str(tmp_path), str(tmp_path))
    base = {"task_id": "t", "product_id": "p1", "status": "verified", "reason": "", "currency": "KRW", "unit": "ea", "raw_fields": {"price": "원문"}, "evidence": {}, "evidence_path": "", "evidence_sha256": "", "locator": "", "extraction_method": "dom", "comparable": True, "comparison_key": "same", "freshness": "observed"}
    observations = [
        dict(base, id="old", source_id="s1", source_name="출처 1", source_url="https://a.example", amount="10", normalized_amount="5", calculation="/2", collected_at="2026-01-01T00:00:00Z"),
        dict(base, id="new", source_id="s1", source_name="출처 1", source_url="https://a.example", amount="20", normalized_amount="10", calculation="/2", collected_at="2026-01-02T00:00:00Z"),
        dict(base, id="other", source_id="s2", source_name="출처 2", source_url="https://b.example", amount="40", normalized_amount="20", calculation="/2", collected_at="2026-01-01T00:00:00Z"),
        dict(base, id="bad", source_id="s2", source_name="출처 2", source_url="https://b.example", amount="NaN", collected_at="2026-01-03T00:00:00Z"),
    ]
    out = export_report(cfg, {"id": "r"}, [], observations, tmp_path / "stats.xlsx")
    ws = load_workbook(out, read_only=True, data_only=True)["Price Comparison"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert {r[4] for r in rows} == {20, 40}  # old source row and NaN excluded
    assert {r[13] for r in rows} == {2}
    assert {r[14] for r in rows} == {20}
    assert {r[15] for r in rows} == {40}
    assert {r[16] for r in rows} == {30}


def test_rental_report_separates_estimates_unknowns_and_deposit_kinds(tmp_path):
    cfg = _config(tmp_path)
    cfg.products[0].price_profile = "rental"
    base = {"task_id": "t", "product_id": "p1", "source_id": "s1", "source_name": "Synthetic",
            "source_url": "https://example.test/car", "status": "verified", "amount": "590000",
            "currency": "KRW", "price_profile": "rental", "value_origin": "observed", "source_visibility": "visible",
            "verification_level": "evidence_validated", "freshness": "observed", "comparable": True,
            "comparison_key": "synthetic", "collected_at": "2026-01-01T00:00:00Z",
            "raw_fields": {"price": "590000", "reference_price": "650000"}, "rental_conditions": {
                "price_basis": "monthly", "term_months": 36, "deposit_amount": "0", "advance_amount": "2000000",
                "advance_percent": "20",
                "upfront_deposit_amount": "1000000", "deposit_installment_amount": "3000000", "deposit_installment_months": 12}}
    observations = [dict(base, id="observed"),
        dict(base, id="estimate", value_origin="calculator_estimate", derived_amount="612345", status="review", amount=None,
             raw_fields={"displayed_deposit": "approximately 10 million"}, comparable=False, reason="calculated estimate"),
        # Even malformed flags cannot leak an estimate or hidden amount into comparison.
        dict(base, id="hidden", source_visibility="hidden"),
        dict(base, id="bad-estimate", value_origin="calculator_estimate", derived_amount="600000"),
        dict(base, id="unknown", status="review", comparable=False, amount=None, rental_conditions={"price_basis": "monthly"}),
        dict(base, id="daily", status="review", comparable=False, rental_conditions={"price_basis": "daily"}),
        dict(base, id="review", status="review", comparable=False, reason="conflicting terms"),
    ]
    out = export_report(cfg, {"id": "r"}, [], observations, tmp_path / "rental.xlsx")
    wb = load_workbook(out, data_only=False)
    assert wb.sheetnames == ["Results", *SHEETS, "Rental Quotes"]
    ws = wb["Rental Quotes"]
    data = [dict(zip([c.value for c in ws[1]], row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert data[0]["Observed Monthly Price"] == 590000
    assert data[0]["Estimated Monthly Price (not observed)"] is None
    assert data[0]["Deposit Amount"] == 0
    assert data[0]["Advance Amount"] == 2000000
    assert data[0]["Reference Price (observed)"] == "650000"
    assert data[0]["Advance Percent (observed)"] == "20"
    assert data[0]["Upfront Deposit Amount"] == 1000000
    assert data[0]["Deposit Installment Amount"] == 3000000
    assert data[0]["Deposit Installment Months"] == 12
    assert data[1]["Observed Monthly Price"] is None
    assert data[1]["Estimated Monthly Price (not observed)"] == 612345
    assert data[1]["Displayed Deposit (raw)"] == "approximately 10 million"
    assert data[2]["Observed Monthly Price"] is None
    assert data[3]["Observed Monthly Price"] is None
    assert data[4]["Term (months)"] is None and data[4]["Deposit Amount"] is None
    assert data[5]["Observed Monthly Price"] is None
    assert data[6]["Observed Monthly Price"] == 590000
    assert wb["Price Comparison"].max_row == 2
    assert wb["Observation History"]["V3"].value == 612345
    assert wb["Observation History"]["H3"].value is None
    assert wb["Observation History"]["Z8"].value == 590000
    assert wb["Observation History"]["H8"].value is None
    assert ws.freeze_panes == "C2"


def test_report_retains_capture_artifacts_and_excludes_unconfirmed_html(tmp_path):
    cfg = _config(tmp_path)
    base = {"task_id": "t", "product_id": "p1", "source_id": "s1", "status": "verified",
            "value_origin": "observed", "amount": "12", "currency": "USD",
            "comparable": True, "comparison_key": "same", "freshness": "observed",
            "collected_at": "2026-09-22T00:00:00Z", "raw_fields": {"price": "12"}}
    artifacts = {"receipt": {"path": str(tmp_path / "capture.json"), "sha256": "receipt-hash"},
                 "screenshot": {"path": str(tmp_path / "capture.png"), "sha256": "image-hash"}}
    rows = [dict(base, id="static", evidence_mode="static_html", source_visibility="unconfirmed"),
            dict(base, id="rendered", evidence_mode="rendered_dom", source_visibility="visible",
                 evidence_artifacts=artifacts),
            dict(base, id="unconfirmed", evidence_mode="rendered_dom", source_visibility="unconfirmed")]
    output = export_report(cfg, {"id": "r"}, [], rows, tmp_path / "artifacts.xlsx")
    wb = load_workbook(output, read_only=True, data_only=True)
    try:
        evidence_sheet = wb["Source Evidence"]
        headers = next(evidence_sheet.iter_rows(values_only=True))
        evidence = [dict(zip(headers, values)) for values in evidence_sheet.iter_rows(min_row=2, values_only=True)]
        assert evidence[0]["Screenshot File"] is None
        assert evidence[1]["Evidence Mode"] == "rendered_dom"
        assert evidence[1]["Capture Receipt"] == artifacts["receipt"]["path"]
        assert evidence[1]["Receipt SHA-256"] == "receipt-hash"
        assert evidence[1]["Screenshot File"] == artifacts["screenshot"]["path"]
        assert evidence[1]["Screenshot SHA-256"] == "image-hash"
        assert wb["Price Comparison"].max_row == 2

    finally:
        wb.close()


def test_hidden_price_field_evidence_is_left_blank_with_review_note(tmp_path):
    from su_crawler.report import HIDDEN_PRICE_NOTE
    cfg = _config(tmp_path)
    row = {"id": "hidden", "task_id": "t", "product_id": "p1", "source_id": "s1", "status": "verified",
           "value_origin": "observed", "source_visibility": "visible", "amount": "12", "currency": "USD",
           "comparable": True, "comparison_key": "k", "freshness": "observed", "evidence_mode": "rendered_dom",
           "collected_at": "2026-09-22T00:00:00Z", "source_url": "https://shop.example/p",
           "raw_fields": {"name": "Item", "price": "12"}, "reason": "",
           "evidence": {"price": {"raw": "12", "location": "span", "display_state": "hidden"}}}
    output = export_report(cfg, {"id": "r"}, [], [row], tmp_path / "hidden.xlsx")
    wb = load_workbook(output, read_only=True)
    try:
        results = list(wb["Results"].values)
        assert results[1][2] is None and HIDDEN_PRICE_NOTE in results[1][5]
        history = list(wb["Observation History"].values)
        header = history[0]
        assert history[1][header.index("Amount")] is None
        assert history[1][header.index("Raw Amount")] == "12"  # raw evidence stays visible as text
        assert HIDDEN_PRICE_NOTE in history[1][header.index("Reason")]
        assert wb["Price Comparison"].max_row == 1
    finally:
        wb.close()


def test_oversized_evidence_text_is_cut_with_explicit_marker(tmp_path):
    cfg = _config(tmp_path)
    card = "card text " * 5000  # 50,000 characters, beyond Excel's cell limit
    row = {"id": "big", "task_id": "t", "product_id": "p1", "source_id": "s1", "status": "review",
           "raw_fields": {"name": "Item", "price": "12"},
           "evidence": {"price": {"raw": "12", "location": "span", "card_text": card}}}
    output = export_report(cfg, {"id": "r"}, [], [row], tmp_path / "big.xlsx")
    wb = load_workbook(output, read_only=True)
    try:
        sheet = list(wb["Source Evidence"].values)
        cell = sheet[1][sheet[0].index("Field Evidence JSON")]
        assert len(cell) <= 32767
        assert cell.endswith("characters shown; the full value is kept in the SQLite ledger.]")
        assert "of 50,0" in cell  # total length of the JSON text is reported
    finally:
        wb.close()


def _no_openpyxl_fallback(monkeypatch):
    from su_crawler import xlsx_sheets

    def fail(*_args, **_kwargs):
        raise AssertionError("whole workbook was reloaded")
    monkeypatch.setattr(xlsx_sheets, "_replace_openpyxl", fail)


def test_append_sheets_replace_in_place_without_reloading_workbook(tmp_path, monkeypatch):
    from su_crawler.report import append_refresh_summary
    _no_openpyxl_fallback(monkeypatch)
    path = export_report(_config(tmp_path), {"id": "r", "demo": True}, [], [], tmp_path / "splice.xlsx")
    coverage = [{"url": "https://example.test/x", "sections": [{"section": "s", "status": "partial"}],
                 "collection_issues": [{"id": "i1", "reason": "=SUM(1)\x01"}]}]
    append_section_coverage(path, coverage)
    append_refresh_summary(path, {"inventory_completeness_verified": True, "changes": [{"url": "u", "added": 1}],
                                  "item_changes": [{"status": "baseline", "after": {"fields": {"price": 0}}}]})
    append_section_coverage(path, coverage)
    wb = load_workbook(path)
    try:
        # Replaced sheets move to the end, as with the former delete-and-append.
        assert wb.sheetnames[-4:] == ["Collection Changes", "Item Changes", "Catalog coverage", "Collection Issues"]
        assert wb.sheetnames.count("Catalog coverage") == 1
        issues = wb["Collection Issues"]
        assert issues["F2"].value == "=SUM(1)_x0001_" and issues["F2"].data_type == "s"
        assert issues["A1"].font.bold and issues["A1"].fill.fgColor.rgb == "001F4E78"
        assert issues.freeze_panes == "A2" and issues.auto_filter.ref == "A1:I2"
        assert wb["Collection Changes"]["F4"].value is True
        assert wb["Item Changes"]["E2"].value == 0
        assert wb["Item Changes"]["A1"].fill.fgColor.rgb == "0017365D"
        assert wb["Item Changes"]["B2"].alignment.wrap_text
        assert wb["Results"]["A1"].value == "Demo data — not market prices"
    finally:
        wb.close()


def test_append_falls_back_to_one_openpyxl_write_for_unsupported_packages(tmp_path, monkeypatch):
    from su_crawler import xlsx_sheets

    def unsupported(*_args, **_kwargs):
        raise xlsx_sheets.SpliceUnsupported("synthetic")
    monkeypatch.setattr(xlsx_sheets, "_splice", unsupported)
    path = export_report(_config(tmp_path), {"id": "r"}, [], [], tmp_path / "fallback.xlsx")
    append_section_coverage(path, [{"url": "https://example.test/x", "sections": [{"section": "s", "observed_count": 3}],
                                    "collection_issues": [{"id": "i1", "attempt_count": 2}]}])
    wb = load_workbook(path, read_only=True)
    try:
        assert list(wb["Catalog coverage"].values)[1][:4] == ("https://example.test/x", "s", "unknown", 3)
        assert list(wb["Collection Issues"].values)[1][7] == 2
    finally:
        wb.close()


def test_export_writes_audit_sheets_in_the_same_pass(tmp_path):
    path = export_report(_config(tmp_path), {"id": "r"}, [], [], tmp_path / "single.xlsx",
                         section_coverage=[{"url": "https://example.test/x",
                                            "sections": [{"section": "s", "status": "complete", "observed_count": 1}]}],
                         collection_issues=[{"id": "i1", "status": "open", "attempt_count": 1}])
    wb = load_workbook(path)
    try:
        assert wb.sheetnames[-2:] == ["Catalog coverage", "Collection Issues"]
        assert list(wb["Catalog coverage"].values)[1][:4] == ("https://example.test/x", "s", "complete", 1)
        assert list(wb["Collection Issues"].values)[1][:4] == ("i1", None, None, "open")
        assert wb["Catalog coverage"].auto_filter.ref == "A1:L2" and wb["Catalog coverage"].freeze_panes == "A2"
        assert wb["Catalog coverage"]["A1"].font.bold
    finally:
        wb.close()
