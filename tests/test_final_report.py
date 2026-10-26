from datetime import datetime
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from su_crawler.final_report import finalize_reports


def _source(path: Path, run: str, *, price=None, site="https://synthetic.example/item", optional=False, demo=False):
    wb = Workbook()
    summary = wb.active
    summary.title = "Run Summary"
    summary.append(["Field", "Value"])
    summary.append(["Run ID", run])
    summary.append(["Status", "complete"])
    summary.append(["Started at (UTC)", datetime(2025, 1, 1, 1, 2)])
    summary.append(["Products", 1])
    summary.append(["Demo data", "Yes" if demo else "No"])
    results = wb.create_sheet("Results")
    results.append(["Product", "Source", "Recorded Price (not a final quote)", "Source URL", "Category"])
    results.append(["Synthetic item", "Synthetic shop", price, site, None])
    results["A2"].font = Font(bold=True)
    results["D2"].hyperlink = site
    results["D2"].style = "Hyperlink"
    tasks = wb.create_sheet("Collection Status")
    tasks.append(["Task ID", "Status", "Updated at (UTC)"])
    tasks.append([f"task-{run}", "completed", datetime(2025, 1, 2, 3, 4)])
    evidence = wb.create_sheet("Source Evidence")
    evidence.append(["Observation ID", "Source URL", "Evidence File"])
    evidence.append([f"obs-{run}", site, f"evidence/{run}.html"])
    if optional:
        issues = wb.create_sheet("Collection Issues")
        issues.append(["Issue ID", "Reason"])
        issues.append([f"issue-{run}", '=HYPERLINK("bad")'])
        changes = wb.create_sheet("Collection Changes")
        changes.append(["Scope", f"scope-{run}"])
        changes.append(["Mode", f"mode-{run}", "Reason", f"reason-{run}"])
        changes.append(["Last full check (UTC)", datetime(2025, 1, 1, 1, 2),
                        "Next interval full check (incremental only, UTC)", None])
        changes.append(["Note", f"note-{run}", "Approved URLs rechecked", 1,
                        "Inventory completeness verified", False])
        changes.append(["Page URL", "Check status"])
        changes.append([site, "changed"])
    wb.save(path)
    wb.close()
    return path


def test_batch_merges_only_requested_runs_and_preserves_evidence(tmp_path):
    first = _source(tmp_path / "first.xlsx", "one", price=0, optional=True)
    second = _source(tmp_path / "second.xlsx", "two", price=None)
    omitted = _source(tmp_path / "omitted.xlsx", "three", price=999)
    before = {p: p.read_bytes() for p in (first, second, omitted)}
    outcome = finalize_reports([first, second], tmp_path / "final.xlsx")
    wb = load_workbook(outcome["report_path"])
    try:
        assert [row[0] for row in list(wb["Results"].values)[1:]] == ["Synthetic item"] * 2
        assert [row[2] for row in list(wb["Results"].values)[1:]] == [0, None]
        assert [row[-1] for row in list(wb["Results"].values)[1:]] == ["one", "two"]
        assert wb["Results"]["A2"].font.bold
        assert wb["Results"]["D2"].hyperlink.target == "https://synthetic.example/item"
        assert wb["Collection Status"]["C3"].value == datetime(2025, 1, 2, 3, 4)
        assert wb["Source Evidence"]["C3"].value == "evidence/two.html"
        assert wb["Collection Issues"]["B2"].value == '=HYPERLINK("bad")'
        assert wb["Collection Issues"]["B2"].data_type == "s"
        assert wb["Collection Changes"]["A5"].value == "Page URL"
        assert wb["Collection Changes"]["A6"].value == "https://synthetic.example/item"
        assert "three" not in str(list(wb["Source Runs"].values))
    finally:
        wb.close()
    assert all(p.read_bytes() == before[p] for p in before)
    assert outcome["observed_product_records"] == 2
    assert outcome["recorded_price_rows"] == 1
    assert outcome["collection_tasks"] == 2


def test_single_same_path_preserves_workbook_and_dashboard_escapes_source_data(tmp_path):
    source = _source(tmp_path / "final report.xlsx", "one", price=0,
                     site="https://synthetic.example/<script>alert(1)</script>")
    before = source.read_bytes()
    outcome = finalize_reports([source], source)
    html = Path(outcome["dashboard_path"]).read_text(encoding="utf-8")
    assert source.read_bytes() == before
    assert "href='final%20report.xlsx'" in html
    assert "download=" not in html
    assert "<script>" not in html
    assert "&lt;script&gt;" not in html
    assert "Not measured" in html
    assert "<div class='qvalue'>Unknown</div>" in html
    assert "측정되지 않음" in html and "未測定" in html
    assert "Category unavailable: no recorded category field." not in html  # synthetic fixture records it
    assert html.count("href='final%20report.xlsx'") == 3
    finalize_reports([source], source, metadata={"coverage": {"completed": 1, "total": 4}})
    assert "1 / 4 (25.0%)" in Path(outcome["dashboard_path"]).read_text(encoding="utf-8")


def test_batch_rejects_output_source_and_invalid_workbook_without_publishing(tmp_path):
    first = _source(tmp_path / "one.xlsx", "one")
    second = _source(tmp_path / "two.xlsx", "two")
    with pytest.raises(ValueError, match="source report"):
        finalize_reports([first, second], first)
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("invalid", encoding="utf-8")
    with pytest.raises(zipfile.BadZipFile):
        finalize_reports([first, invalid], tmp_path / "final.xlsx")
    assert not (tmp_path / "final.xlsx").exists()
    assert not (tmp_path / "final.html").exists()


def test_metadata_coverage_never_uses_task_success_as_completeness(tmp_path):
    source = _source(tmp_path / "run.xlsx", "one", price=None)
    output = tmp_path / "final.xlsx"
    result = finalize_reports([source], output, metadata={
        "coverage_summary": {"complete": 1, "partial": 1, "unknown": 2, "excluded": 0},
        "target_count": 5, "coverage_counts": {"visited": 4},
    })
    html = Path(result["dashboard_path"]).read_text(encoding="utf-8")
    assert "1 / 5 (20.0%) complete; 1 partial, 2 unknown, 0 excluded" in html
    assert result["recorded_price_rows"] == 0
    assert source.read_bytes() != b""
    result = finalize_reports([source], output, metadata={"coverage": [{"status": "visited"}]})
    assert "0 / 1 (0.0%) complete; 0 partial, 1 unknown" in Path(result["dashboard_path"]).read_text(encoding="utf-8")
    result = finalize_reports([source], output, metadata={
        "coverage": [{"status": "visited", "sections": [{"status": "complete"}]},
                     {"status": "visited"}, {"status": "timeout"}],
        "coverage_summary": {"complete": 1, "unknown": 0}, "target_count": 1,
    })
    assert "1 / 3 (33.3%) complete; 1 partial, 1 unknown" in Path(result["dashboard_path"]).read_text(encoding="utf-8")


def test_batch_keeps_every_run_summary_and_change_metadata(tmp_path):
    first = _source(tmp_path / "one.xlsx", "one", optional=True)
    second = _source(tmp_path / "two.xlsx", "two", optional=True)
    result = finalize_reports([first, second], tmp_path / "final.xlsx")
    wb = load_workbook(result["report_path"])
    try:
        details = list(wb["Run Details"].values)
        assert ("one", "Products", 1) in details
        assert ("two", "Products", 1) in details
        assert ("one", "Started at (UTC)", datetime(2025, 1, 1, 1, 2)) in details
        assert ("two", "Started at (UTC)", datetime(2025, 1, 1, 1, 2)) in details
        changes = list(wb["Change Summaries"].values)
        assert ("one", 2, "Mode", "mode-one", "Reason", "reason-one") == changes[2][:6]
        assert ("two", 2, "Mode", "mode-two", "Reason", "reason-two") == changes[6][:6]
        assert changes[4][7] is False
        assert wb["Collection Changes"]["B1"].value == "See Change Summaries; rows below combine source runs."
        assert list(wb["Collection Changes"].values)[5][-1] == "one"
        assert list(wb["Collection Changes"].values)[6][-1] == "two"
    finally:
        wb.close()


def test_samefile_does_not_copy_workbook(tmp_path, monkeypatch):
    source = _source(tmp_path / "run.xlsx", "one")
    import su_crawler.final_report as final_report

    def reject_copy(*_args, **_kwargs):
        raise AssertionError("same-file finalization copied the workbook")

    monkeypatch.setattr(final_report.shutil, "copyfile", reject_copy)
    result = finalize_reports([source], source)
    assert Path(result["dashboard_path"]).exists()


def test_missing_category_and_site_stay_unknown_without_invented_groups(tmp_path):
    source = _source(tmp_path / "run.xlsx", "one", site="", price=None)
    wb = load_workbook(source)
    results = wb["Results"]
    results.delete_cols(5)
    wb.save(source)
    wb.close()
    result = finalize_reports([source], source)
    html = Path(result["dashboard_path"]).read_text(encoding="utf-8")
    assert "Category unavailable: no recorded category field." in html
    assert "분류 정보 없음" in html and "分類情報なし" in html
    assert "<th scope='row'>알 수 없음</th>" in html
    assert "<th scope='row'>不明</th>" in html
    assert result["source_sites"] == 0


def test_dashboard_matches_offline_sample_design_and_uses_only_supplied_metadata(tmp_path):
    source = _source(tmp_path / "run.xlsx", "one", price=7)
    wb = load_workbook(source)
    wb["Results"].delete_cols(5)
    wb.save(source)
    wb.close()
    result = finalize_reports([source], source, metadata={
        "topic": {"industry": "<Synthetic>", "product": "Widget", "market": "Test"},
        "status": "needs_review", "started_at": "2025-01-01T00:00:00Z",
        "finished_at": "2025-01-02T00:00:00Z",
    })
    html = Path(result["dashboard_path"]).read_text(encoding="utf-8")
    assert "--green:#1d745b" in html and "--bg:#f3f5f1" in html
    assert "@media print" in html and "@media(max-width:760px)" in html
    assert "<script" not in html and "src=" not in html
    assert "Industry: &lt;Synthetic&gt;" in html
    assert "needs_review" in html
    assert "2025-01-01T00:00:00Z" in html
    assert "2025-01-02T00:00:00Z" in html
    assert "Included data" in html and "Source Evidence" in html
    assert "#en:focus-visible~main label[for=en]" in html
    assert html.count("href='run.xlsx'") == 3


def test_batch_requires_exact_headers_and_demo_banner_is_visible(tmp_path):
    first = _source(tmp_path / "one.xlsx", "one", demo=True)
    second = _source(tmp_path / "two.xlsx", "two")
    wb = load_workbook(second)
    wb["Results"].delete_cols(5)
    wb.save(second)
    wb.close()
    with pytest.raises(ValueError, match="incompatible sheet headers: Results"):
        finalize_reports([first, second], tmp_path / "final.xlsx")
    assert not (tmp_path / "final.xlsx").exists()
    result = finalize_reports([first], first)
    html = Path(result["dashboard_path"]).read_text(encoding="utf-8")
    assert "Demo data — not market prices" in html
    assert "데모 데이터 — 시장 가격이 아닙니다" in html
    assert "デモデータ — 市場価格ではありません" in html


def test_many_rows_keep_provenance_and_filter_after_cached_merge(tmp_path):
    sources = []
    for index in range(3):
        source = _source(tmp_path / f"run-{index}.xlsx", f"run-{index}", price=index)
        wb = load_workbook(source)
        for row in range(3, 53):
            wb["Results"].cell(row, 1, f"Synthetic item {row}")
            wb["Results"].cell(row, 3, index)
        wb["Results"].auto_filter.ref = "A1:E52"
        wb.save(source)
        wb.close()
        sources.append(source)
    result = finalize_reports(sources, tmp_path / "final.xlsx")
    wb = load_workbook(result["report_path"])
    try:
        rows = list(wb["Results"].values)
        assert len(rows) == 154
        assert [rows[n][-1] for n in (1, 51, 52, 102, 103, 153)] == [
            "run-0", "run-0", "run-1", "run-1", "run-2", "run-2"]
        assert wb["Results"].auto_filter.ref == "A1:F154"
        assert [rows[n][2] for n in (1, 52, 103)] == [0, 1, 2]
    finally:
        wb.close()


def _export_with_summary(tmp_path, demo=False):
    from su_crawler.models import CollectionConfig, Product, Source
    from su_crawler.report import _export
    config = CollectionConfig("t", [Product("p1", "P")], [Source("s1", "S", "web", "https://a.example/p", ["p1"])],
                              str(tmp_path), str(tmp_path))
    rows = [{"id": f"o{i}", "task_id": "t", "product_id": "p1", "source_id": "s1", "status": "verified",
             "amount": "10" if i % 2 else None, "source_url": url, "raw_fields": {"name": "P", "price": "10"},
             "evidence": {}} for i, url in enumerate(["https://a.example/p", "https://b.example/q", "not a url",
                                                      None, "https://a.example/r"])]
    path, summary = _export(config, {"id": "r", "demo": demo}, [{"id": "t", "product_id": "p1", "source_id": "s1"}],
                            rows, tmp_path / "summary.xlsx", section_coverage=[])
    stat = path.stat()
    summary["file"] = [stat.st_size, stat.st_mtime_ns]
    return path, summary


@pytest.mark.parametrize("demo", [False, True])
def test_exported_summary_matches_streamed_workbook_read(tmp_path, demo):
    from su_crawler.final_report import workbook_summary
    path, summary = _export_with_summary(tmp_path, demo)
    read = workbook_summary(path)
    assert {key: value for key, value in summary.items() if key != "file"} == read


def test_finalize_uses_current_summary_and_ignores_a_stale_one(tmp_path, monkeypatch):
    from su_crawler import final_report
    path, summary = _export_with_summary(tmp_path)
    expected = finalize_reports([path], path)
    monkeypatch.setattr(final_report, "load_workbook", lambda *a, **k: pytest.fail("workbook reloaded"))
    supplied = finalize_reports([path], path, metadata={"workbook_summary": summary})
    assert supplied == expected
    summary["file"] = [summary["file"][0], summary["file"][1] - 1]  # describes another file state
    monkeypatch.undo()
    calls = []
    original = final_report.workbook_summary
    monkeypatch.setattr(final_report, "workbook_summary", lambda p: calls.append(p) or original(p))
    assert finalize_reports([path], path, metadata={"workbook_summary": summary}) == expected
    assert calls


ESCAPE_LIKE = ["_x0041_ code", "line\r\nbreak", "keep x005F_ text", " spaced ", "plain"]


def _escape_like_reports(tmp_path):
    from su_crawler.models import CollectionConfig, Product, Source
    from su_crawler.report import export_report
    config = CollectionConfig("t", [Product("p1", "P")], [Source("s1", "S", "web", "https://a.example/p", ["p1"])],
                              str(tmp_path), str(tmp_path))
    paths = []
    for run in ("r1", "r2"):
        rows = [{"id": f"{run}-{i}", "task_id": "t", "product_id": "p1", "source_id": "s1", "status": "review",
                 "raw_fields": {"name": name, "price": "12"}} for i, name in enumerate(ESCAPE_LIKE)]
        paths.append(export_report(config, {"id": run}, [], rows, tmp_path / f"{run}.xlsx"))
    return paths


def test_batch_merge_keeps_text_exactly_as_excel_shows_the_sources(tmp_path):
    from su_crawler.xlsx_readback import read_values
    paths = _escape_like_reports(tmp_path)
    finalize_reports(paths, tmp_path / "final.xlsx")
    for path in paths:
        source = read_values(path)["Results"]
        assert [source[row][0] for row in sorted(source) if row] == ESCAPE_LIKE
    merged = read_values(tmp_path / "final.xlsx")["Results"]
    assert [merged[row][0] for row in sorted(merged) if row] == ESCAPE_LIKE * 2
    assert {merged[row][-1] for row in sorted(merged) if row} == {"r1", "r2"}


def test_batch_merge_that_does_not_read_back_publishes_nothing(tmp_path, monkeypatch):
    from su_crawler import final_report
    from su_crawler.xlsx_readback import WorkbookMismatch
    paths = _escape_like_reports(tmp_path)
    # A writer fault: text is stored unescaped, so Excel would decode "_x0041_".
    monkeypatch.setattr(final_report, "excel_literal", lambda text: text)
    with pytest.raises(WorkbookMismatch, match=r"wrote '_x0041_ code', read back 'A code'"):
        finalize_reports(paths, tmp_path / "final.xlsx")
    assert not list(tmp_path.glob("final*")) and not list(tmp_path.glob(".final.*"))
