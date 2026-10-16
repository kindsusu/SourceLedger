"""Offline checks for approved-plan refresh decisions and semantic comparison."""
from datetime import datetime, timedelta, timezone

import pytest

from su_crawler.catalog_refresh import (compare, decide, normalize_refresh_policy, scope_key,
                                       read_baseline, write_baseline, refresh_for_run, save_refresh_run, item_changes)
from su_crawler.catalog_collection import collect_plan
from su_crawler.models import FetchResult


def test_policy_defaults_and_bounds():
    assert normalize_refresh_policy(None) == {'mode': 'full', 'full_interval_days': 7}
    assert normalize_refresh_policy({})['mode'] == 'full'
    assert normalize_refresh_policy({'mode': 'incremental'})['mode'] == 'incremental'
    assert normalize_refresh_policy({'mode': 'full', 'full_interval_days': 30})['mode'] == 'full'
    for value in ({'mode': 'other'}, {'full_interval_days': True}, {'full_interval_days': 0},
                  {'full_interval_days': 366}, {'unknown': 1}):
        with pytest.raises(ValueError):
            normalize_refresh_policy(value)


def test_due_rule_change_and_failed_full_remains_due():
    now = datetime(2026, 10, 15, tzinfo=timezone.utc)
    policy = normalize_refresh_policy({'mode': 'incremental'})
    last = (now - timedelta(days=8)).isoformat()
    baseline = {'last_full_at': last, 'rules_version': 'v1', 'valid': True}
    assert decide(policy, baseline, 'v1', now=now)['reason'] == 'interval_elapsed'
    assert decide(policy, baseline, 'v2', now=now)['reason'] == 'rules_changed'
    assert decide(policy, baseline, 'v1', manual_full=True, now=now)['reason'] == 'manual_full'
    assert decide(policy, baseline, 'v1', now=now)['last_full_at'] == last


def test_scope_separates_selected_batches_and_ignores_order():
    snapshot = {'id': 'plan', 'revision': 1, 'topic': {'market': 'KR'}}
    assert scope_key(snapshot, ['https://a.example/a', 'https://a.example/b'], follow_links=False) == scope_key(
        snapshot, ['https://a.example/b', 'https://a.example/a'], follow_links=False)
    assert scope_key(snapshot, ['https://a.example/a'], follow_links=False) != scope_key(
        snapshot, ['https://a.example/b'], follow_links=False)
    newer = {**snapshot, 'revision': 2, 'refresh_policy': {'mode': 'full', 'full_interval_days': 1}}
    assert scope_key(snapshot, ['https://a.example/a'], follow_links=False) == scope_key(
        newer, ['https://a.example/a'], follow_links=False)


def test_semantic_multiset_preserves_duplicates_and_does_not_infer_removal():
    old = [{'fields': {'sku': '1', 'price': '10'}, 'specs': {'size': 'M'}}] * 2
    new = [old[0], {'fields': {'sku': '1', 'price': '11'}, 'specs': {'size': 'M'}}]
    assert compare(old, new) == {'added': 1, 'changed': 0, 'unchanged': 1,
                                 'removed_candidate': 0, 'unavailable': 1}
    assert compare(old, [old[0]])['unavailable'] == 1
    assert compare(old, [old[0]], inventory_complete=True)['removed_candidate'] == 1


def test_shared_sku_variants_are_not_paired_arbitrarily():
    old = [{'fields': {'sku': '1', 'price': '10'}, 'specs': {'size': size}} for size in ('M', 'L')]
    new = [{'fields': {'sku': '1', 'price': '11'}, 'specs': {'size': size}} for size in ('L', 'M')]
    assert compare(old, new)['changed'] == 2
    ambiguous = [{'fields': {'sku': '1', 'price': '12'}, 'specs': {'size': 'XL'}}]
    result = compare(old, ambiguous)
    assert result['changed'] == 0 and result['added'] == 1 and result['unavailable'] == 2


def _plan(url, policy=None):
    return {'id': 'plan', 'revision': 1, 'state': 'confirmed',
            'topic': {'industry': 'retail', 'product': 'item', 'market': 'public'},
            'candidates': [{'id': 'one', 'name': 'item', 'url': url, 'selected': True}],
            'refresh_policy': policy or normalize_refresh_policy({'mode': 'incremental'})}


def _html(price):
    return ('<script type="application/ld+json">{"@type":"Product","name":"Item",'
            '"sku":"SKU-1","offers":{"@type":"Offer","price":"%s",'
            '"priceCurrency":"USD"}}</script>' % price).encode()


def test_repeated_run_304_then_change_preserves_original_time_and_current_evidence(tmp_path):
    url = 'https://shop.example/item'
    calls = []
    state = {'price': '10'}

    def fetch(source, _base, backend, *, validators=None):
        calls.append((backend, bool(validators)))
        if backend != 'http':
            return FetchResult(source.id, 'tool_unavailable', backend)
        if validators and state['price'] == '10':
            return FetchResult(source.id, 'not_modified', backend, final_url=url,
                               http_metadata={'etag': 'v1'})
        return FetchResult(source.id, 'fetched', backend, content=_html(state['price']),
                           final_url=url, http_metadata={'etag': 'v1' if state['price'] == '10' else 'v2'})

    first = collect_plan(tmp_path, _plan(url), output_dir=tmp_path / 'out', collector=fetch,
                         max_pages=1, follow_links=False)
    second = collect_plan(tmp_path, _plan(url), output_dir=tmp_path / 'out', collector=fetch,
                          max_pages=1, follow_links=False)
    assert first['refresh']['effective_mode'] == 'full'
    assert second['refresh']['effective_mode'] == 'incremental'
    assert ('http', True) in calls
    assert second['refresh']['counts']['unchanged'] == 1
    assert second['refresh']['changes'][0]['original_fetched_at']
    assert any(item.get('code') == 'http_not_modified' for item in second['coverage'][0]['trace'])
    state['price'] = '11'
    third = collect_plan(tmp_path, _plan(url), output_dir=tmp_path / 'out', collector=fetch,
                         max_pages=1, follow_links=False)
    assert third['refresh']['counts']['changed'] == 1
    # Later report exports must retain the change sheet from the durable ledger.
    from openpyxl import load_workbook
    from su_crawler.config import load_config
    from su_crawler.pipeline import report_for_run
    from su_crawler.storage import Store
    store = Store(tmp_path / 'out')
    try:
        report_for_run(load_config(third['config_path']), store, third['run_id'])
    finally:
        store.close()
    book = load_workbook(third['report_path'])
    try:
        assert book['Collection Changes']['D6'].value == 1
    finally:
        book.close()


def test_explicit_full_forces_unconditional_fetch(tmp_path):
    url = 'https://shop.example/item'
    calls = []
    def fetch(source, _base, backend, *, validators=None):
        if backend == 'http':
            calls.append(validators)
            return FetchResult(source.id, 'fetched', backend, content=_html('10'), final_url=url,
                               http_metadata={'etag': 'v1'})
        return FetchResult(source.id, 'tool_unavailable', backend)
    policy = {'mode': 'full', 'full_interval_days': 1}
    for _ in range(2):
        result = collect_plan(tmp_path, _plan(url, policy), output_dir=tmp_path / 'out',
                              collector=fetch, max_pages=1, follow_links=False)
        assert result['refresh']['effective_mode'] == 'full'
    assert calls == [None, None]


def test_changed_source_signature_disables_validators(tmp_path):
    url = 'https://shop.example/item'
    calls = []
    def fetch(source, _base, backend, *, validators=None):
        if backend == 'http':
            calls.append(validators)
            return FetchResult(source.id, 'fetched', backend, content=_html('10'), final_url=url,
                               http_metadata={'etag': 'v1'})
        return FetchResult(source.id, 'tool_unavailable', backend)
    plan = _plan(url)
    collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch,
                 max_pages=1, follow_links=False)
    key = scope_key(plan, [url], follow_links=False)
    baseline = read_baseline(tmp_path / 'out', key)
    baseline['pages'][url]['source_signature'] = 'changed'
    write_baseline(tmp_path / 'out', key, baseline)
    collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch,
                 max_pages=1, follow_links=False)
    assert calls == [None, None]


def test_missing_previously_discovered_url_is_unavailable_and_preserved(tmp_path):
    root, child = 'https://shop.example/start', 'https://shop.example/child'
    state = {'link': True}
    def fetch(source, _base, backend, *, validators=None):
        if backend != 'http':
            return FetchResult(source.id, 'tool_unavailable', backend)
        content = _html('10') + (b'<a href="/child">Child</a>' if state['link'] and source.location == root else b'')
        return FetchResult(source.id, 'fetched', backend, content=content, final_url=source.location)
    plan = _plan(root)
    first = collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch, max_pages=2)
    assert any(page['url'] == child for page in first['coverage'])
    state['link'] = False
    second = collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch, max_pages=2)
    missing = next(row for row in second['refresh']['changes'] if row['url'] == child)
    assert missing['status'] == 'unavailable' and missing['removed_candidate'] == 0
    assert child in read_baseline(tmp_path / 'out', scope_key(plan, [root], follow_links=True))['pages']


def test_sqlite_retains_all_change_rows_beyond_bounded_summary(tmp_path):
    full = {'changes': [{'url': f'https://shop.example/{n}'} for n in range(51)]}
    save_refresh_run(tmp_path, 'run-1', full)
    assert len(refresh_for_run(tmp_path, 'run-1')['changes']) == 51


def test_due_full_failure_does_not_postpone_next_check(tmp_path):
    url = 'https://shop.example/item'
    state = {'fail': False}
    validators_seen = []
    def fetch(source, _base, backend, *, validators=None):
        validators_seen.append(validators)
        if state['fail']:
            return FetchResult(source.id, 'blocked', backend, message='access challenge')
        return FetchResult(source.id, 'fetched', backend, content=_html('10'),
                           final_url=url, http_metadata={'etag': 'v1'})
    plan = _plan(url)
    out = tmp_path / 'out'
    collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    key = scope_key(plan, [url], follow_links=False)
    baseline = read_baseline(out, key)
    old = (datetime.now(timezone.utc) - timedelta(days=8)).isoformat()
    baseline['last_full_at'] = old
    write_baseline(out, key, baseline)
    state['fail'] = True
    failed = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert failed['refresh']['reason'] == 'interval_elapsed'
    assert failed['refresh']['last_full_at'] == old
    assert failed['refresh']['counts']['unavailable'] == 1
    assert not failed['refresh']['full_recheck_succeeded']
    assert read_baseline(out, key)['last_full_at'] == old
    state['fail'] = False
    retry = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert retry['refresh']['effective_mode'] == 'full'
    assert retry['refresh']['last_full_at'] != old
    assert all(value is None for value in validators_seen)


def test_xlsx_preserves_all_changes_and_never_executes_source_text(tmp_path):
    from openpyxl import Workbook, load_workbook
    from su_crawler.report import append_refresh_summary
    path = tmp_path / 'report.xlsx'
    book = Workbook()
    book.save(path)
    book.close()
    append_refresh_summary(path, {'changes': [{'url': '=unsafe()', 'status': 'checked'} for _ in range(55)]})
    book = load_workbook(path)
    try:
        sheet = book['Collection Changes']
        assert sheet.max_row == 60
        assert sheet['A6'].value == '=unsafe()' and sheet['A6'].data_type == 's'
    finally:
        book.close()


@pytest.mark.parametrize('conditional', [True, False])
def test_identical_http_reuses_raw_extraction_and_reassesses_scope(monkeypatch, tmp_path, conditional):
    import su_crawler.catalog_collection as collection
    url = 'https://shop.example/item'
    extracted = []
    original = collection._jsonld_candidates
    content = _html('10') + _html('20').replace(b'Item', b'Other').replace(b'SKU-1', b'SKU-2')

    def tracked(result):
        extracted.append(result.content)
        return original(result)

    monkeypatch.setattr(collection, '_jsonld_candidates', tracked)
    calls = []

    def fetch(source, _base, backend, *, validators=None):
        calls.append((backend, validators))
        if backend != 'http':
            return FetchResult(source.id, 'tool_unavailable', backend)
        return FetchResult(source.id, 'not_modified' if conditional and validators else 'fetched',
                           backend, content=b'' if conditional and validators else content,
                           final_url=url, http_metadata={'etag': 'v1'})

    plan = _plan(url)
    plan['exclude_terms'] = ['Other']
    first = collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch,
                         max_pages=1, follow_links=False)
    second = collect_plan(tmp_path, plan, output_dir=tmp_path / 'out', collector=fetch,
                          max_pages=1, follow_links=False)
    assert len(extracted) == 1
    assert second['execution']['extraction_reused_responses'] == 1
    assert second['execution']['http_not_modified_responses'] == int(conditional)
    assert first['coverage'][0]['scope_assessments'] == second['coverage'][0]['scope_assessments']
    assert second['coverage'][0]['scope_counts']['excluded'] == 1
    assert second['refresh']['counts']['unchanged'] == 1
    # Reusing static extraction must not suppress required browser coverage.
    assert [backend for backend, _ in calls] == ['http', 'playwright', 'http', 'playwright']


def test_extraction_failure_preserves_baseline_and_full_timestamp_then_recovers(monkeypatch, tmp_path):
    import su_crawler.catalog_collection as collection
    url = 'https://shop.example/item'
    state = {'price': '10', 'fail': False}
    original = collection._jsonld_candidates

    def extraction(result):
        if state['fail']:
            raise RuntimeError('controlled failure')
        return original(result)

    monkeypatch.setattr(collection, '_jsonld_candidates', extraction)

    def fetch(source, _base, backend, *, validators=None):
        if backend != 'http':
            return FetchResult(source.id, 'tool_unavailable', backend)
        return FetchResult(source.id, 'fetched', backend, content=_html(state['price']), final_url=url,
                           http_metadata={'etag': state['price']})

    plan = _plan(url)
    out = tmp_path / 'out'
    collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    key = scope_key(plan, [url], follow_links=False)
    baseline = read_baseline(out, key)
    baseline['last_full_at'] = (datetime.now(timezone.utc) - timedelta(days=8)).isoformat()
    write_baseline(out, key, baseline)
    state.update(price='11', fail=True)
    failed = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert failed['coverage'][0]['status'] == 'failed'
    assert any(event.get('code') == 'extraction_error' for event in failed['coverage'][0]['trace'])
    assert failed['refresh']['counts']['unavailable'] == 1
    assert not failed['refresh']['full_recheck_succeeded']
    assert read_baseline(out, key) == baseline
    state['fail'] = False
    recovered = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert recovered['refresh']['counts']['changed'] == 1
    assert recovered['refresh']['counts']['added'] == 0
    assert recovered['refresh']['full_recheck_succeeded']
    assert recovered['refresh']['last_full_at'] != baseline['last_full_at']


def test_retained_reanalysis_rebuilds_cache_and_recovers_extraction_failure(monkeypatch):
    from dataclasses import asdict
    import su_crawler.catalog_collection as collection
    from su_crawler.models import Source
    url = 'https://shop.example/item'
    source = Source('one', 'Item', 'web', url, [], allowed_domains=['shop.example'])
    response = FetchResult(source.id, 'fetched', 'http', content=_html('10'), final_url=url,
                           trace=[{'code': 'extraction_reused'}])
    captures = {url: {'source': asdict(source), 'responses': [(asdict(response), [])],
                      'raw_http_candidates': [{'stale': True}]}}
    coverage = [{'url': url}]
    original = collection._jsonld_candidates

    def fail(_result):
        raise RuntimeError('controlled reanalysis failure')

    monkeypatch.setattr(collection, '_jsonld_candidates', fail)
    failed = collection._reanalyse(captures, coverage, [], [], [])
    assert coverage[0]['status'] == 'failed'
    assert failed[url]['raw_http_candidates'] is None
    assert not any(event.get('code') == 'extraction_reused' for event in coverage[0]['trace'])
    assert any(event.get('code') == 'extraction_error' for event in coverage[0]['trace'])
    monkeypatch.setattr(collection, '_jsonld_candidates', original)
    recovered = collection._reanalyse(failed, coverage, [], [], [])
    assert coverage[0]['status'] == 'visited'
    assert len(recovered[url]['raw_http_candidates']) == 1
    assert not any(event.get('code') == 'extraction_error' for event in coverage[0]['trace'])


def test_unsupported_empty_page_preserves_nonempty_baseline_until_recovery(tmp_path):
    url = 'https://shop.example/item'
    state = {'content': _html('10')}

    def fetch(source, _base, backend, *, validators=None):
        if backend != 'http':
            return FetchResult(source.id, 'tool_unavailable', backend)
        return FetchResult(source.id, 'fetched', backend, content=state['content'], final_url=url)

    plan = _plan(url)
    out = tmp_path / 'out'
    collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    key = scope_key(plan, [url], follow_links=False)
    baseline = read_baseline(out, key)
    assert not baseline['pages'][url]['inventory_complete']
    state['content'] = b'<html><body>New unsupported layout</body></html>'
    empty = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert empty['coverage'][0]['status'] == 'no_data'
    assert not empty['refresh']['approved_urls_rechecked']
    assert empty['refresh']['counts']['unavailable'] == 1
    assert read_baseline(out, key) == baseline
    state['content'] = _html('11')
    recovered = collect_plan(tmp_path, plan, output_dir=out, collector=fetch, max_pages=1, follow_links=False)
    assert recovered['refresh']['counts']['changed'] == 1
    assert recovered['refresh']['counts']['added'] == 0
