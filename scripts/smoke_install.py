"""Verify an installed distribution with local fixtures, never a live site.

Run with ``python -I scripts/smoke_install.py`` from the repository checkout.
The imports must resolve from the installed wheel, not the checkout.
"""
from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile


def main() -> int:
    import su_crawler
    from openpyxl import load_workbook

    root = Path(__file__).resolve().parents[1]
    package = Path(su_crawler.__file__).resolve()
    if package.is_relative_to(root / "su_crawler"):
        raise RuntimeError("Smoke check requires a non-editable wheel installation")
    with tempfile.TemporaryDirectory(prefix="sourceledger-install-") as folder:
        work = Path(folder)
        document = json.loads((root / "examples/verification.json").read_text(encoding="utf-8"))
        document["output_dir"] = "outputs"
        for source in document["sources"]:
            source["file_root"] = str(root / "examples")
        config = work / "verification.json"
        config.write_text(json.dumps(document), encoding="utf-8")
        receipt = work / "receipt.json"

        def cli(*args: str) -> dict:
            result = subprocess.run([sys.executable, "-I", "-m", "su_crawler", *args],
                                    cwd=work, capture_output=True, text=True,
                                    encoding="utf-8", timeout=60)
            if result.returncode:
                raise RuntimeError(result.stdout + result.stderr)
            return json.loads(result.stdout)

        verified = cli("verify", "--config", str(config), "--samples",
                       str(root / "examples/verification.samples.json"), "--receipt", str(receipt))
        assert verified["eligible"] is True
        activated = cli("activate", "--config", str(config), "--receipt", str(receipt),
                        "--output", str(work / "active.json"))
        assert activated["status"] == "activated"
        run = cli("run", "--config", str(work / "active.json"))
        book = load_workbook(run["report_path"], read_only=True)
        try:
            assert "Source Evidence" in book.sheetnames
            assert book["Price Comparison"].max_row > 1
        finally:
            book.close()
    print("Installed wheel: local verification, activation, SQLite and XLSX smoke check passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
