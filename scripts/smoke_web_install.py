"""Verify the installed wheel's browser assets, durable job, and XLSX offline."""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import shutil
import tempfile
import threading
import time


def main() -> int:
    import httpx
    import su_crawler
    from openpyxl import load_workbook
    from su_crawler.assistant_runtime import runtime_status, stop_worker
    from su_crawler.web_server import build_web_server

    checkout = Path(__file__).resolve().parents[1]
    if Path(su_crawler.__file__).resolve().is_relative_to(checkout / "su_crawler"):
        raise RuntimeError("Smoke check requires a non-editable wheel installation")
    mcp_checked = False

    with tempfile.TemporaryDirectory(prefix="sourceledger-web-") as folder:
        workspace = Path(folder) / "workspace 한글 space!"
        workspace.mkdir()
        fixtures = workspace / "fixtures"
        fixtures.mkdir()
        document = json.loads((checkout / "examples/verification.json").read_text(encoding="utf-8"))
        document["output_dir"] = "outputs"
        for source in document["sources"]:
            original = checkout / "examples" / source["location"]
            shutil.copyfile(original, fixtures / original.name)
        (workspace / "verification.json").write_text(json.dumps(document), encoding="utf-8")
        server = build_web_server(workspace, port=0)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        worker_started = False
        try:
            with httpx.Client(base_url=server.url, trust_env=False, timeout=20) as client:
                for asset, media_type in [("/", "text/html"), ("/app.js", "javascript"),
                                          ("/styles.css", "text/css"), ("/mark.svg", "image/svg+xml")]:
                    response = client.get(asset)
                    response.raise_for_status()
                    assert media_type in response.headers["content-type"]
                    assert response.content
                state = client.get("/api/bootstrap").json()
                assert state["status"] == "needs_setup"
                client.headers.update({"Origin": server.url, "X-SourceLedger-Token": state["csrf_token"]})

                def post(route: str, data: dict) -> dict:
                    result = client.post(route, json=data)
                    result.raise_for_status()
                    return result.json()

                post("/api/workspace", {"industry": "Synthetic components", "product": "Test part",
                                        "market": "Offline fixture"})
                post("/api/product", {"identifiers": {"model": "TEST-A"}})
                job = post("/api/jobs", {"operation": "run", "arguments": {"config_path": "verification.json"}})
                assert job["status"] == "queued"
                post("/api/worker", {"action": "start"})
                worker_started = True
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    job = client.get(f"/api/jobs/{job['id']}").json()
                    if job["status"] in {"succeeded", "failed", "interrupted"}:
                        break
                    time.sleep(0.1)
                assert job["status"] == "succeeded", job
                rows = client.get(f"/api/jobs/{job['id']}/observations?limit=1").json()
                assert rows["total"] >= 2 and len(rows["rows"]) == 1
                report = client.get(f"/api/jobs/{job['id']}/report")
                report.raise_for_status()
                assert "attachment" in report.headers["content-disposition"]
                workbook = load_workbook(BytesIO(report.content), read_only=True)
                try:
                    assert "Source Evidence" in workbook.sheetnames
                    assert workbook["Price Comparison"].max_row > 1
                finally:
                    workbook.close()
                if state["mcp_available"]:
                    # Exercise the real MCP tool boundary against the browser's workspace.
                    import asyncio
                    from su_crawler.assistant_server import build_assistant_server

                    mcp = build_assistant_server(workspace)
                    _, shared = asyncio.run(mcp.call_tool("get_job_status", {"job_id": job["id"]}))
                    assert shared["id"] == job["id"] and shared["status"] == "succeeded"
                    for name in ("codex", "claude-code", "claude-desktop"):
                        connection = post("/api/connections", {"client": name})
                        assert "serve-assistant" in connection["snippet"]
                        assert connection["workspace_root"] == str(workspace.resolve())
                    mcp_checked = True
        finally:
            if worker_started:
                stop_worker(workspace)
                deadline = time.monotonic() + 10
                while time.monotonic() < deadline and runtime_status(workspace)["status"] != "stopped":
                    time.sleep(0.1)
                assert runtime_status(workspace)["status"] == "stopped"
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)
    print("Installed wheel: browser assets, onboarding, durable job, observations and XLSX passed. "
          + ("Shared MCP and three client snippets passed." if mcp_checked else "Core-only UI passed without MCP installed."))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
