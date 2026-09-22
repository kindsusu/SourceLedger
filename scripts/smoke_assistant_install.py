"""Exercise an installed wheel's stdio MCP and independent worker offline.

Run with ``python -I scripts/smoke_assistant_install.py`` after installing the
wheel with its ``mcp`` extra. Uses synthetic files in a temporary workspace.
"""
from __future__ import annotations

import asyncio
import base64
from contextlib import asynccontextmanager
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time


def main() -> int:
    import su_crawler
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    from openpyxl import load_workbook

    checkout = Path(__file__).resolve().parents[1]
    if Path(su_crawler.__file__).resolve().is_relative_to(checkout / "su_crawler"):
        raise RuntimeError("Smoke check requires a non-editable wheel installation")

    with tempfile.TemporaryDirectory(prefix="sourceledger-assistant-") as folder:
        workspace = (Path(folder) / "research 한글 space!").resolve()
        workspace.mkdir()
        fixtures = workspace / "fixtures"
        fixtures.mkdir()
        document = json.loads((checkout / "examples/verification.json").read_text(encoding="utf-8"))
        document["output_dir"] = "outputs"
        for source in document["sources"]:
            original = checkout / "examples" / source["location"]
            shutil.copyfile(original, fixtures / original.name)
        (workspace / "verification.json").write_text(json.dumps(document), encoding="utf-8")

        def cli(*arguments: str) -> dict:
            result = subprocess.run([sys.executable, "-I", "-m", "su_crawler", *arguments],
                                    cwd=folder, capture_output=True, text=True, encoding="utf-8", timeout=30)
            if result.returncode:
                raise RuntimeError(result.stdout + result.stderr)
            return json.loads(result.stdout)

        connection = cli("connect", "--client", "claude-desktop", "--workspace-root", str(workspace))
        entry = json.loads(Path(connection["snippet_path"]).read_text(encoding="utf-8"))["mcpServers"]["sourceledger"]

        @asynccontextmanager
        async def session():
            parameters = StdioServerParameters(command=entry["command"], args=entry["args"],
                cwd=folder, env={**os.environ, **entry["env"]})
            async with stdio_client(parameters) as (read, write):
                async with ClientSession(read, write) as client:
                    await client.initialize()
                    yield client

        async def call(client, name, arguments=None):
            result = await client.call_tool(name, arguments or {})
            if result.isError:
                raise RuntimeError(f"{name}: {result.content}")
            return json.loads(result.content[0].text)

        async def enqueue():
            async with session() as client:
                status = await call(client, "get_workspace_status")
                assert status["status"] == "needs_setup"
                await call(client, "initialize_workspace", {
                    "industry": "Synthetic components", "product": "Test part", "market": "Offline fixture"})
                await call(client, "set_product_identity", {"identifiers": {"model": "TEST-A"}})
                queued = await call(client, "queue_collection", {"config_path": "verification.json"})
                assert queued["status"] == "queued"
                assert queued["worker"]["status"] == "stopped"
                return queued["id"]

        job_id = asyncio.run(enqueue())
        # The MCP client/server have both exited before independent execution.
        worker_args = ["--workspace-root", str(workspace)]
        try:
            status = cli("assistant", "start", *worker_args)
            assert status["status"] in {"idle", "running"}
            deadline = time.monotonic() + 30
            while True:
                job = cli("assistant", "job", *worker_args, "--job-id", job_id)
                if job["status"] in {"succeeded", "failed", "interrupted"}:
                    break
                if time.monotonic() >= deadline:
                    raise RuntimeError("Installed assistant job timed out")
                time.sleep(0.1)
            assert job["status"] == "succeeded", job

            async def read_results():
                async with session() as client:
                    status = await call(client, "get_job_status", {"job_id": job_id})
                    assert status["status"] == "succeeded"
                    observations = await call(client, "get_job_observations", {"job_id": job_id, "limit": 1})
                    assert observations["total"] >= 2 and len(observations["rows"]) == 1
                    metadata = await call(client, "get_job_report", {"job_id": job_id})
                    resource = await client.read_resource(metadata["resource_uri"])
                    content = base64.b64decode(resource.contents[0].blob)
                    assert len(content) == metadata["size"]
                    assert hashlib.sha256(content).hexdigest() == metadata["sha256"]
                    assert content == Path(metadata["path"]).read_bytes()
                    workbook = load_workbook(metadata["path"], read_only=True)
                    try:
                        assert "Source Evidence" in workbook.sheetnames
                        assert workbook["Price Comparison"].max_row > 1
                    finally:
                        workbook.close()

            asyncio.run(read_results())
        finally:
            cli("assistant", "stop", *worker_args)
            deadline = time.monotonic() + 10
            while cli("assistant", "status", *worker_args)["status"] != "stopped":
                if time.monotonic() >= deadline:
                    raise RuntimeError("Installed assistant worker did not stop")
                time.sleep(0.1)
    print("Installed wheel: stdio MCP onboarding, disconnected execution, observations and XLSX resource passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
